"""Two design workflows sharing the existing SOILFIRM PRO calculation engine."""
from copy import deepcopy
import json
import tkinter as tk
from ui_theme import UI_FONT, UI_FONT_MONO
from tkinter import ttk, filedialog, messagebox


SINGLE_GROUPS = (
    ('Thông tin và dữ liệu', (0,)),
    ('Thông số nền đắp', (1,)),
    ('Thông số địa chất', (2,)),
    ('Kiểm toán trước xử lý', (3,)),
    ('Thiết kế xử lý nền', (4, 5, 6)),
    ('Tổng hợp kết quả xử lý', (18,)),
)
MULTI_GROUPS = (
    ('Dự án và phân đoạn', (12,)),
    ('Địa chất và chỉ tiêu', (11, 10, 14)),
    ('Kiểm toán trước xử lý', (15,)),
    ('Thiết kế xử lý nền', (16, 4, 5, 6)),
    ('Tổng hợp và hồ sơ', (9,)),
)


# Display names are separate from existing navigation values and comparisons.
GROUP_DISPLAY_TITLES = {
    'Thông tin và dữ liệu': 'Dữ liệu dự án',
    'Thông số nền đắp': 'Hình học nền đắp',
    'Thông số địa chất': 'Địa tầng và chỉ tiêu đất',
    'Địa chất và chỉ tiêu': 'Địa tầng và chỉ tiêu đất',
    'Lựa chọn và hồ sơ': 'Lựa chọn phương án',
    'Tổng hợp và hồ sơ': 'Khối lượng xử lý nền',
}


def display_group_title(title):
    return GROUP_DISPLAY_TITLES.get(title, title)


def normalize_mode(mode):
    # Continue loading projects saved with the previous mode names.
    if str(mode).strip() in ('Tính nhiều đoạn', 'TÍNH NHIỀU ĐOẠN', 'TÍNH TOÀN TUYẾN', 'Tính toàn tuyến'):
        return 'TÍNH TOÀN TUYẾN'
    return 'TÍNH MỘT ĐOẠN'


def navigation_groups(mode):
    return MULTI_GROUPS if normalize_mode(mode) == 'TÍNH TOÀN TUYẾN' else SINGLE_GROUPS


def fmt(value):
    if value is None:
        return '—'
    return f'{value:.2f}' if isinstance(value, (float, int)) else str(value)


def table(parent, headings, widths):
    box = ttk.Frame(parent)
    box.pack(fill='both', expand=True, padx=8, pady=6)
    box.columnconfigure(0, weight=1)
    box.rowconfigure(0, weight=1)
    columns = tuple(f'c{i}' for i in range(len(headings)))
    tree = ttk.Treeview(box, columns=columns, show='headings')
    for key, title, width in zip(columns, headings, widths):
        tree.heading(key, text=title)
        tree.column(key, width=width, minwidth=70, stretch=False)
    sy = ttk.Scrollbar(box, command=tree.yview)
    sx = ttk.Scrollbar(box, orient='horizontal', command=tree.xview)
    tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
    tree.grid(row=0, column=0, sticky='nsew')
    sy.grid(row=0, column=1, sticky='ns')
    sx.grid(row=1, column=0, sticky='ew')
    tree.tag_configure('pass', foreground='#166534')
    tree.tag_configure('fail', foreground='#B91C1C')
    return tree


def section_signature(project):
    from result_summary import project_signature
    # Depth, elevation and borehole identity matter even if the layer values match.
    return project_signature(project) + json.dumps([
        project.borehole_name, project.ground_elevation, project.borehole_depth,
        project.cdm_inputs, project.alicc_inputs], sort_keys=True, default=str)


def option_time(attempt, record):
    """Display engine-produced schedule values without estimating a duration."""
    for value in attempt.get('design', {}).values():
        if value.get('unit') in ('ngày', 'days'):
            return fmt(value.get('value')) + ' ngày'
    payload = (record or {}).get('payload', {})
    if 'wait_days' in payload:
        return 'Chờ ' + fmt(payload['wait_days']) + ' ngày'
    return 'Theo lịch trình thiết kế' if record else '—'


class SingleSelection(ttk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app
        self.options = []
        self.note = tk.StringVar()
        self.tabs = ttk.Frame(self)
        self.tabs.pack(fill='both', expand=True)
        self.selection_detail = tk.Toplevel(self)
        self.selection_detail.withdraw()
        self.selection_detail.title('Chi tiết phương án đã chọn')
        self.selection_detail.geometry('820x580')
        self.selection_detail.transient(app)
        self.selection_detail.protocol('WM_DELETE_WINDOW', self.selection_detail.withdraw)
        pages = [ttk.Frame(self.tabs), ttk.Frame(self.selection_detail)]
        for page in pages:page.pack(fill='both', expand=True)
        bar = ttk.Frame(pages[0], padding=8)
        bar.pack(fill='x')
        ttk.Button(bar, text='Làm mới bảng', command=self.refresh).pack(side='left')
        ttk.Button(bar, text='Chọn phương án', command=self.choose, style='Accent.TButton').pack(side='left', padx=6)
        ttk.Label(pages[0], textvariable=self.note, wraplength=950, padding=8).pack(fill='x')
        self.tree = table(pages[0], ('Phương án', 'Thông số xử lý', 'Lún dư (cm)', 'Giới hạn (cm)',
            'Kiểm toán lún', 'Kiểm toán phương án', 'Thời gian'), (230, 350, 110, 110, 140, 160, 160))
        self.tree.bind('<Double-1>', lambda e: self.choose())
        self.detail = tk.Text(pages[1], wrap='word', font=(UI_FONT, 11), state='disabled')
        self.detail.pack(fill='both', expand=True, padx=12, pady=10)
        self.reason = tk.StringVar()
        ttk.Label(pages[1], text='Lý do lựa chọn', padding=8).pack(anchor='w')
        ttk.Entry(pages[1], textvariable=self.reason).pack(fill='x', padx=12, pady=6)
        ttk.Button(pages[1], text='Lưu lý do lựa chọn', command=self.save_reason).pack(anchor='w', padx=12, pady=6)
        ttk.Button(pages[1], text='Đóng', command=self.selection_detail.withdraw).pack(anchor='e', padx=12, pady=6)

    def refresh(self):
        from result_summary import collect_calculated_options
        self.options, missing = collect_calculated_options(self.app)
        self.tree.delete(*self.tree.get_children())
        for i, item in enumerate(self.options):
            self.tree.insert('', 'end', iid=str(i), values=(item.opt_name, item.opt_params_desc,
                fmt(item.residual_cm), fmt(item.limit_cm), 'ĐẠT' if item.residual_cm <= item.limit_cm else 'CHƯA ĐẠT',
                'ĐẠT' if item.is_pass else 'CHƯA ĐẠT', item.time_desc), tags=('pass' if item.is_pass else 'fail',))
        self.note.set('So sánh kết quả đã tính; không tự xếp phương án tối ưu. ' +
                      ('Chưa có kết quả: ' + '; '.join(missing) if missing else 'Đã có kết quả các nhóm.'))
        record = getattr(self.app, '_single_selected_option', None)
        self.reason.set((record or {}).get('reason', ''))
        self.show_record(record)

    def choose(self):
        if not self.tree.selection():
            messagebox.showinfo('Lựa chọn phương án', 'Chọn một phương án đã tính.', parent=self.app)
            return
        item = self.options[int(self.tree.selection()[0])]
        from result_summary import project_signature
        try:
            self.app.collect()
            if project_signature(self.app.project) != project_signature(item.source_project):
                raise ValueError('Đầu vào đã thay đổi. Tính lại và cập nhật bảng trước khi lựa chọn.')
            if not item.is_pass:
                raise ValueError('Phương án chưa đạt các kiểm toán áp dụng; điều chỉnh thiết kế trước khi lựa chọn.')
            method = item.apply_payload.get('method')
            special = deepcopy(getattr(self.app, '_cdm_reports', {}).get(method)) if method else None
            self.app._single_selected_option = {
                'name': item.opt_name, 'params': item.opt_params_desc, 'residual': item.residual_cm,
                'limit': item.limit_cm, 'pass_check': item.is_pass, 'notes': item.tech_notes,
                'time': item.time_desc, 'reason': self.reason.get(), 'payload': deepcopy(item.apply_payload),
                'project_snapshot': deepcopy(item.source_project), 'signature': section_signature(self.app.project),
                'special': special, 'method': method,
            }
            self.show_record(self.app._single_selected_option)
            self.selection_detail.deiconify()
            self.selection_detail.lift()
            board = getattr(self.app, 'single_treatment_summary_view', None)
            if board is not None:
                self.app.after_idle(board.refresh)
        except Exception as exc:
            messagebox.showerror('Lựa chọn phương án', str(exc), parent=self.app)

    def show_record(self, record):
        text = 'Chưa lựa chọn phương án.'
        if record:
            p = record['project_snapshot']
            stale = section_signature(self.app.project) != record['signature']
            text = (f'Dự án: {p.name}\nLý trình: {p.station_from} – {p.station_to}\nMặt cắt: {p.station}\n\n'
                    f'Phương án: {record["name"]}\nThông số: {record["params"]}\n'
                    f'Lún dư: {record["residual"]:.2f} cm ≤ {record["limit"]:.2f} cm\n'
                    f'Kiểm toán phương án: {"ĐẠT" if record["pass_check"] else "CHƯA ĐẠT"}\n'
                    f'Thời gian: {record["time"]}\n{record["notes"]}\n'
                    + ('\nĐầu vào đã thay đổi; cần tính và lựa chọn lại.' if stale else ''))
        self.detail.configure(state='normal')
        self.detail.delete('1.0', 'end')
        self.detail.insert('end', text)
        self.detail.configure(state='disabled')

    def save_reason(self):
        record = getattr(self.app, '_single_selected_option', None)
        if record:
            record['reason'] = self.reason.get()
            self.app.status_text.set('Đã cập nhật lý do lựa chọn; bấm Lưu dự án để lưu tệp.')

    def export_record(self):
        if not self.app.require_full_license():
            return None
        record = getattr(self.app, '_single_selected_option', None)
        if not record:
            messagebox.showinfo('Xuất hồ sơ', 'Lựa chọn phương án trước khi xuất.', parent=self.app)
            return None
        self.app.collect()
        if section_signature(self.app.project) != record['signature']:
            raise ValueError('Đầu vào đã thay đổi; cần lựa chọn lại phương án đã tính.')
        self.save_reason()
        return record

    def export_json(self):
        try:
            record = self.export_record()
            if not record: return
            dest = filedialog.asksaveasfilename(parent=self.app, defaultextension='.json', filetypes=[('JSON', '*.json')])
            if dest:
                from pathlib import Path
                Path(dest).write_text(json.dumps(self.app._pack_saved_result(record), ensure_ascii=False, indent=2), encoding='utf-8')
        except Exception as exc: messagebox.showerror('Xuất JSON', str(exc), parent=self.app)

    def export_excel(self):
        try:
            record = self.export_record()
            if not record: return
            dest = filedialog.asksaveasfilename(parent=self.app, defaultextension='.xlsx', filetypes=[('Excel', '*.xlsx')])
            if not dest: return
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook(); ws = wb.active; ws.title = 'Phuong an lua chon'
            p = record['project_snapshot']
            for row in [('Dự án', p.name), ('Lý trình đầu', p.station_from), ('Lý trình cuối', p.station_to),
                        ('Mặt cắt', p.station), ('Phương án', record['name']), ('Thông số', record['params']),
                        ('Lún dư (cm)', record['residual']), ('Giới hạn (cm)', record['limit']),
                        ('Kiểm toán', 'ĐẠT' if record['pass_check'] else 'CHƯA ĐẠT'),
                        ('Thời gian', record['time']), ('Lý do lựa chọn', record['reason']), ('Chi tiết kiểm toán', record['notes'])]:
                ws.append(row)
            ws.column_dimensions['A'].width = 25; ws.column_dimensions['B'].width = 100
            for cell in ws['A']: cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='E3EEF5')
            strata = wb.create_sheet('Dia tang dau vao')
            strata.append(['Lớp', 'Bề dày (m)', 'γ (T/m³)', 'e₀', 'Cc', 'Cs', 'Pc (T/m²)', 'Cv'])
            for s in p.soils: strata.append([s.name, s.thickness, s.gamma, s.e0, s.cc, s.cs, s.pc, str(s.cv)])
            wb.save(dest)
        except Exception as exc: messagebox.showerror('Xuất Excel', str(exc), parent=self.app)

    def export_pdf(self):
        try:
            record = self.export_record()
            if not record: return
            dest = filedialog.asksaveasfilename(parent=self.app, defaultextension='.pdf', filetypes=[('PDF', '*.pdf')])
            if not dest: return
            from report_pdf import create_report
            scope = 'cdm' if record['special'] else 'natural' if record['payload'].get('treatment_group') == 'natural' else 'treated'
            create_report(record['project_snapshot'], dest, scope=scope, special_method=record['method'],
                          special_data=record['special'], language='en' if self.app.ui_language.get() == 'English' else 'vi')
        except Exception as exc: messagebox.showerror('Xuất PDF', str(exc), parent=self.app)


def settlement_display_values(details=None, project=None, before=None):
    """Read calculated components at the governing location, in centimetres."""
    details = details or {}
    native = details.get('result', {}) or {}
    if 'sum_sc' in native or 'sum_si' in native:
        return native.get('sum_sc'), native.get('sum_si'), native.get('total_s')
    if 'S_total_cm' in native:
        # This report exposes S1/S2, without a verified Sc/Si decomposition.
        return None, None, native.get('S_total_cm')
    rows = [x.get('kết_quả', x) for x in details.get('output_locations', [])]
    rows = rows or details.get('locations', []) or details.get('rows', []) or (before or {}).get('locations', [])
    if project is not None and project.expansion_width > 0:
        rows = [r for r in rows if 'nền mở rộng' in str(r.get('vị_trí', '')).lower()]
    rows = [r for r in rows if isinstance(r, dict) and r.get('Sc_dư_cm') is not None]
    if not rows:
        return None, None, None
    row = max(rows, key=lambda r: r['Sc_dư_cm'])
    sc, si = row.get('Sc_cuối_cm'), row.get('Si_cm')
    total = row.get('St_cuối_cm')
    if total is None and sc is not None and si is not None:
        total = sc + si
    return sc, si, total


def treatment_display_details(attempt, option, project):
    """Format existing design inputs and calculated outputs without defaults."""
    details = attempt.get('calculation') or option.get('calculation_report') or option.get('details') or {}
    special = option.get('special') or {}
    params = details.get('params') or special.get('params') or {}
    payload = option.get('payload') or attempt.get('payload') or {}
    design = attempt.get('design') or {}
    group = payload.get('treatment_group') or getattr(project, 'treatment_group', '')
    name = str(attempt.get('option') or option.get('opt_name') or '')
    lines = []
    def add(label, value, unit=''):
        if value is None or value == '':return
        rendered = fmt(value) if isinstance(value, (int, float)) else str(value)
        line = label + ': ' + rendered + (' ' + unit if unit else '')
        if line not in lines:lines.append(line)
    if group == 'mechanical':
        add('Chiều sâu đào', payload.get('replacement_depth', getattr(project, 'replacement_depth', None)), 'm')
        for key, label, density in [('bamboo_depth', 'Cọc tre', 'bamboo_density'), ('cajuput_depth', 'Cừ tràm', 'cajuput_density')]:
            length = payload.get(key, getattr(project, key, None))
            if length is not None and length > 0:
                add(label + ' dài', length, 'm');add('Mật độ ' + label.lower(), payload.get(density), 'cọc/m²')
    elif group == 'cdm' or 'CDM' in name:
        for label, key, source in [('Chiều dài cọc L꜀', 'cdm_lc', 'Lc'), ('Đường kính cọc D', 'cdm_d', 'D'), ('Khoảng cách cọc s', 'cdm_s', 's')]:
            add(label, params.get(source, payload.get(key)), 'm')
        add('Loại cọc', params.get('pile_type'));add('Bố trí cọc', params.get('pattern'))
        native = details.get('result') or special.get('result') or {}
        add('Tỷ lệ diện tích cọc aₚ', native.get('ap'))
        geo = details.get('geo_params') or special.get('geo_params') or {}
        if payload.get('cdm_n_layer') or geo.get('n_layer'):
            add('Loại gia cường', payload.get('reinforcement_kind', geo.get('kind')))
            add('Số lớp gia cường', payload.get('cdm_n_layer', geo.get('n_layer')), 'lớp')
            add('Cường độ kéo T', geo.get('T_char_kn', geo.get('T_char')), 'kN/m')
    elif name != 'Không cần xử lý':
        treatment = getattr(project, 'treatment', '')
        if name.startswith(('PVD', 'SD')) or treatment.startswith(('PVD', 'SD')):
            add('Chiều sâu xử lý', payload.get('drain_length', getattr(project, 'drain_length', None)), 'm')
            add('Khoảng cách thoát nước s', payload.get('drain_spacing', getattr(project, 'drain_spacing', None)), 'm')
            add('Đường kính thoát nước d', getattr(project, 'drain_diameter', None), 'cm')
            add('Bố trí thoát nước', getattr(project, 'drain_pattern', None))
        add('Chiều cao gia tải', payload.get('surcharge_height'), 'm')
        add('Thời gian chờ', details.get('wait_days', payload.get('wait_days')), 'ngày')
        if 'chân không' in (name + ' ' + treatment).lower():
            add('Áp lực chân không', getattr(project, 'vacuum_pressure', None), 'T/m²')
    for key, item in design.items():
        add(item.get('label', key), item.get('value'), item.get('unit', ''))
    return lines


class SegmentBoard(ttk.Frame):
    def __init__(self, parent, app, after=False):
        super().__init__(parent)
        self.app, self.after_design = app, after
        self.rows = {}
        self.note = tk.StringVar()
        host = self
        if getattr(self, '_single_result_tabs', False):
            self.result_tabs = ttk.Notebook(self)
            self.result_tabs.pack(fill='both', expand=True)
            self.comparison_page = ttk.Frame(self.result_tabs)
            self.summary_page = ttk.Frame(self.result_tabs)
            self.result_tabs.add(self.comparison_page, text='So sánh phương án')
            self.result_tabs.add(self.summary_page, text='Bảng tổng hợp kết quả')
            host = self.comparison_page
        bar = ttk.Frame(host, padding=8); bar.pack(fill='x')
        if after:
            actions = [('Tính bằng AI', self.run_ai_design),
                       ('Xác nhận phương án', self.choose)]
        else:
            actions = [('Kiểm toán các đoạn chọn', lambda: self.calculate(False)),
                       ('Hỗ trợ kiểm toán', lambda: self.calculate(True)),
                       ('Chọn tất cả đoạn', lambda: self.tree.selection_set(self.tree.get_children()))]
        actions += [('Làm mới bảng', self.refresh)]
        for title, command in actions: ttk.Button(bar, text=title, command=command).pack(side='left', padx=3)
        note_row = ttk.Frame(host);note_row.pack(fill='x')
        if not after:
            export_bar = ttk.Frame(note_row)
            ttk.Label(export_bar, text='STT phân đoạn:').pack(side='left', padx=(0, 4))
            self.export_number = tk.StringVar()
            self.export_selector = ttk.Combobox(export_bar, textvariable=self.export_number, state='readonly', width=9)
            self.export_selector.pack(side='left', padx=(0, 6))

        ttk.Label(note_row, textvariable=self.note, padding=8, wraplength=950).pack(side='left', fill='x', expand=True)
        heads = ['STT', 'Lý trình', 'Chiều dài (m)', 'Mặt cắt', 'Hₜₖ (m)', 'Hₜₜ (m)', 'Lỗ khoan', 'Địa tầng / bề dày (m)']
        widths = [65, 190, 100, 160, 100, 100, 130, 320]
        heads += ['Phương án', 'Thông số xử lý'] if after else []
        widths += [220, 480] if after else []
        heads += ['S꜀ dư (cm)']
        widths += [130]
        if not after:
            heads += ['Lún cố kết S꜀ (cm)', 'Lún tức thời Sᵢ (cm)', 'Lún tổng S (cm)']
            widths += [165, 165, 150]
        heads += ['Cho phép [ΔS] (cm)', 'Kết quả kiểm toán', 'Trạng thái']
        widths += [160, 160, 240]
        self.tree = table(host, heads, widths)
        if after:
            ttk.Style(self).configure('DesignDetails.Treeview', rowheight=48, font=(UI_FONT, 10))
            self.tree.configure(style='DesignDetails.Treeview')
        detail_box = ttk.Frame(host); detail_box.pack(fill='x', padx=8, pady=6)
        self.detail = tk.Text(detail_box, height=8, wrap='word', state='disabled', font=(UI_FONT, 10))
        sy = ttk.Scrollbar(detail_box, command=self.detail.yview)
        self.detail.configure(yscrollcommand=sy.set)
        sy.pack(side='right', fill='y'); self.detail.pack(fill='x', expand=True)
        self.tree.bind('<<TreeviewSelect>>', self.show_detail)
        self.tree.bind('<Double-1>', lambda e: self.open_comparison() if after else self.show_detail())

    def current_project(self, section):
        return self.app._ai_analysis_workspace.build_project(section)

    def refresh(self):
        self.rows.clear()
        session = getattr(self.app, '_ai_analysis_state', {}) or {}
        if not self.after_design:
            available = [str(s['section_no']) for s in session.get('sections', [])]
            self.export_selector.configure(values=available)
            if self.export_number.get() not in available:self.export_number.set('')
        records = {int(r['section_no']): r for r in getattr(self.app, '_saved_sections_data', [])
                   if r.get('section_no') is not None}
        confirmed = dict(records)
        records = {**getattr(self.app, '_batch_candidates', {}), **records}
        selected = self.tree.selection()
        self.tree.delete(*self.tree.get_children())
        for s in session.get('sections', []):
            n = int(s['section_no']); p = None; error = ''
            try: p = self.current_project(s)
            except Exception as exc: error = str(exc)
            identity = [n, s.get('station_from', '') + ' – ' + s.get('station_to', ''), fmt(s.get('length')),
                        s.get('station', ''), fmt(p.h_design if p else s.get('h_design')), fmt(p.height if p else None), s.get('borehole_selected', ''),
                        '; '.join(f'{x.name}: {x.thickness:.2f}' for x in p.soils) if p else 'Chưa đủ địa tầng']
            pending = getattr(self.app, '_batch_runs', {}).get(n) or {}
            record = records.get(n) or {}
            if not pending and record.get('project_snapshot'):
                pending = {'project': record['project_snapshot'], 'before': record.get('before', {}),
                    'result': {'summary': {'attempts': [{'option': record.get('opt_name'),
                        'residual_cm': record.get('residual'), 'limit_cm': record.get('limit'),
                        'pass_check': record.get('status') == 'ĐẠT'}]},
                        'options': {record.get('opt_name'): record}}}
            if self.after_design:
                attempts = pending.get('result', {}).get('summary', {}).get('attempts', [])
                choices = pending.get('result', {}).get('options', {})
                if pending.get('natural'):
                    attempts = [{'option': 'Không cần xử lý', 'residual_cm': pending['before']['residual_cm'],
                                 'limit_cm': pending['before']['limit_cm'], 'pass_check': True}]
                stale = not p or not pending.get('project') or section_signature(p) != section_signature(pending['project'])
                for i, attempt in enumerate(attempts or [{}]):
                    iid = f'{n}:{i}'; option = choices.get(attempt.get('option'), {}) or {}
                    calculated_project = option.get('project_snapshot') or pending.get('project') or p
                    treatment_lines = treatment_display_details(attempt, option, calculated_project)
                    main = treatment_lines[:6]
                    params = ('; '.join(main[:3]) + ('\n' + '; '.join(main[3:]) if len(main) > 3 else '')) if main else option.get('params') or 'Chưa có dữ liệu'
                    status = error or attempt.get('error') or ('Đầu vào đã thay đổi' if attempts and stale else 'Đã chọn' if confirmed.get(n, {}).get('opt_name') == attempt.get('option') and confirmed.get(n, {}).get('payload') == option.get('payload') and confirmed.get(n, {}).get('status') == 'ĐẠT' else 'Chờ người dùng lựa chọn' if attempts else 'Chưa tính')
                    checked = 'CẦN TÍNH LẠI' if stale and attempts else 'ĐẠT' if attempt.get('pass_check') else 'CHƯA ĐẠT' if attempts and not attempt.get('error') else 'CHƯA TÍNH'
                    self.rows[iid] = {'section': s, 'project': p, 'pending': pending, 'attempt': attempt, 'option': option, 'status': status}
                    details = attempt.get('calculation') or option.get('calculation_report') or option.get('details') or {}
                    if not details and option.get('special'):
                        details = {'result': option['special'].get('result', {})}
                    components = settlement_display_values(details, pending.get('project') or p,
                        pending.get('before') if pending.get('natural') or option.get('payload', {}).get('treatment_group') == 'natural' else None)
                    row_identity = list(identity)
                    calculated_project = option.get('project_snapshot') or pending.get('project')
                    if calculated_project is not None:
                        row_identity[4:6] = [fmt(calculated_project.h_design), fmt(calculated_project.height)]
                    native = details.get('result', {}) or {}
                    if native.get('H_tk') is not None:row_identity[4] = fmt(native['H_tk'])
                    if native.get('H_tt') is not None:row_identity[5] = fmt(native['H_tt'])
                    self.tree.insert('', 'end', iid=iid, values=row_identity + [attempt.get('option', '—'), params,
                        fmt(attempt.get('residual_cm')), fmt(attempt.get('limit_cm', s.get('limit_cm'))), checked, status],
                        tags=('pass' if checked == 'ĐẠT' else 'fail',))
            else:
                result = getattr(self.app, '_before_section_results', {}).get(n) or {}
                before = result.get('before') or pending.get('before') or record.get('before') or {}
                src = result.get('project') or pending.get('project')
                stale = not p or not src or section_signature(p) != section_signature(src)
                status = error or result.get('error') or ('Đầu vào đã thay đổi' if before and stale else 'Đã tính' if before else 'Chưa tính')
                check = 'CẦN TÍNH LẠI' if before and stale else 'ĐẠT' if before.get('pass_check') else 'CHƯA ĐẠT' if before else 'CHƯA TÍNH'
                iid = str(n); self.rows[iid] = {'section': s, 'project': p, 'before': before, 'status': status}
                self.tree.insert('', 'end', iid=iid, values=identity + [fmt(before.get('residual_cm')),
                    *[fmt(v) for v in settlement_display_values(project=src or p, before=before)],
                    fmt(before.get('limit_cm', s.get('limit_cm'))), check, status], tags=('pass' if check == 'ĐẠT' else 'fail',))
        for iid in selected:
            if self.tree.exists(iid): self.tree.selection_add(iid)
        self.note.set('Mỗi phương án một hàng; chọn hàng để xem chi tiết. Bấm Xác nhận phương án để tính khối lượng và xuất hồ sơ tại mục Khối lượng xử lý nền.' if self.after_design
                      else 'Kiểm toán riêng trước xử lý. Chọn đoạn cần tính hoặc để trống để tính tất cả; thiếu dữ liệu được ghi theo từng đoạn.')
        self.show_detail()

    def export_before_report(self):
        if not self.app.require_full_license():return
        try:
            raw = self.export_number.get().strip()
            if not raw:raise ValueError('Chọn STT phân đoạn cần xuất báo cáo trước xử lý.')
            row = self.rows.get(str(int(raw)))
            if not row or not row.get('before'):raise ValueError('Phân đoạn chưa có kết quả kiểm toán trước xử lý. Hãy tính trước khi xuất báo cáo.')
            n = int(raw)
            saved = getattr(self.app, '_before_section_results', {}).get(n) or {}
            pending = getattr(self.app, '_batch_runs', {}).get(n) or {}
            record = next((r for r in getattr(self.app, '_saved_sections_data', []) if r.get('section_no') == n), {})
            source = saved.get('project') or pending.get('project') or record.get('before_project')
            current = self.current_project(row['section'])
            if source is None or section_signature(current) != section_signature(source):
                raise ValueError('Dữ liệu phân đoạn đã thay đổi. Kiểm toán lại trước xử lý rồi xuất báo cáo.')
            dest = filedialog.asksaveasfilename(parent=self.app, title=f'Xuất báo cáo trước xử lý · STT {n}',
                initialfile=f'Truoc_xu_ly_STT_{n}.pdf', defaultextension='.pdf', filetypes=[('PDF', '*.pdf')])
            if not dest:return
            show_logo = messagebox.askyesno('Logo PDF', 'Hiển thị logo SOILFIRM PRO trên PDF?', parent=self.app)
            import os, tempfile
            from report_pdf import create_report
            fd, temporary = tempfile.mkstemp(suffix='.pdf', dir=os.path.dirname(os.path.abspath(dest)))
            os.close(fd)
            try:
                project = deepcopy(source)
                project.station_from = row['section'].get('station_from', '')
                project.station_to = row['section'].get('station_to', '')
                create_report(project, temporary, scope='natural', show_logo=show_logo,
                    language='en' if self.app.ui_language.get() == 'English' else 'vi')
                os.replace(temporary, dest)
            finally:
                if os.path.exists(temporary):os.remove(temporary)
            messagebox.showinfo('Xuất báo cáo', f'Đã xuất báo cáo trước xử lý STT {n}:\n{dest}', parent=self.app)
        except Exception as exc:
            messagebox.showerror('Xuất báo cáo trước xử lý', str(exc), parent=self.app)

    def show_detail(self, event=None):
        ids = self.tree.selection(); text = 'Chọn một hàng để xem đầu vào, địa tầng và kết quả đầy đủ.'
        if ids and ids[0] in self.rows:
            row = self.rows[ids[0]]; p = row['project']; s = row['section']
            text = f'STT {s["section_no"]} · {row["status"]}\n'
            if p:
                text += f'Dự án: {p.name} · Mặt cắt: {p.station} · Lỗ khoan: {p.borehole_name}\n'
                top = 0
                for soil in p.soils:
                    text += (f'{soil.name}: sâu {top:.2f}–{top + soil.thickness:.2f} m; h={soil.thickness:.2f} m; '
                             f'γ={soil.gamma:.2f}; e₀={soil.e0:.3f}; Cc={soil.cc:.3f}; Cs={soil.cs:.3f}; Pc={soil.pc:.2f}; Cv={soil.cv}\n')
                    top += soil.thickness
            if not self.after_design and row.get('before'):
                before = row['before']
                text += f'\nThời điểm đánh giá: {fmt(before.get("evaluation_day"))} ngày\n'
                for location in before.get('locations', []):
                    residual = location.get('Sc_dư_cm'); limit = before.get('limit_cm')
                    text += (f'{location.get("vị_trí", "")}: Sc cuối={fmt(location.get("Sc_cuối_cm"))} cm; '
                             f'Sc dư={fmt(residual)} cm; U={fmt(location.get("U_%"))}%\n'
                             f'Sc dư={fmt(residual)} cm ≤ ΔS={fmt(limit)} cm: '
                             f'{"ĐẠT" if residual is not None and limit is not None and residual <= limit else "CHƯA ĐẠT"}\n')
                    for layer in location.get('phân_tố', []):
                        text += ' · '.join(f'{k}: {fmt(v) if not isinstance(v, (list, dict)) else str(v)}' for k, v in layer.items()) + '\n'
            if self.after_design and row.get('option', {}).get('manual'):
                text += '\n' + row['option'].get('params', '') + '\n' + row['option'].get('notes', '')
                if row['option'].get('special'):
                    text += '\n' + json.dumps(row['option']['special'].get('result', {}), ensure_ascii=False, indent=2, default=str)
            elif self.after_design and row.get('option'):
                from soilfirm_ai_engine import present_result
                view = present_result(row['pending']['result'])
                name = row['attempt'].get('option')
                text += '\nĐiều kiện kiểm toán và bảng tính:\n'
                for key in ('parameter_rows', 'check_rows', 'detail_rows'):
                    for values in view.get(key, []):
                        if values and values[0] == name:
                            text += ' · '.join(str(v) for v in values) + '\n'
            if self.after_design:
                option = row.get('option') or {};attempt = row.get('attempt') or {}
                calculated_project = option.get('project_snapshot') or row.get('pending', {}).get('project') or p
                lines = treatment_display_details(attempt, option, calculated_project)
                text += '\nChi tiết xử lý:\n' + ('\n'.join(lines) if lines else 'Chưa có dữ liệu')
                residual, limit = attempt.get('residual_cm'), attempt.get('limit_cm')
                if residual is not None and limit is not None:
                    text += f'\nKiểm toán lún: S꜀ dư = {fmt(residual)} cm ≤ [ΔS] = {fmt(limit)} cm: ' + ('ĐẠT' if residual <= limit else 'CHƯA ĐẠT')
                report = attempt.get('calculation') or option.get('calculation_report') or option.get('details') or {}
                if not report and option.get('special'):report = option['special']
                native = report.get('result') or {}
                if native.get('stress'):
                    for label, key, bound in [('Ứng suất đầu cọc', 'sigma_p', 'qu_allow'), ('Đầu cọc, trạng thái giới hạn 1', 'qu_tt1', 'qu_allow'), ('Ứng suất đất nền', 'sigma_s', 'Rtc')]:
                        value, allowed = native['stress'].get(key), native['stress'].get(bound)
                        if value is not None and allowed is not None:
                            text += f'\n{label}: {fmt(value)} ≤ {fmt(allowed)} T/m²: ' + ('ĐẠT' if value <= allowed else 'CHƯA ĐẠT')
        self.detail.configure(state='normal'); self.detail.delete('1.0', 'end'); self.detail.insert('end', text); self.detail.configure(state='disabled')

    def open_comparison(self):
        ids = self.tree.selection()
        if not ids: return
        n = int(ids[0].split(':')[0]); workspace = self.app._ai_analysis_workspace
        pending = getattr(self.app, '_batch_runs', {}).get(n)
        if not pending: return
        workspace.state['index'] = next(i for i, s in enumerate(workspace.state['sections']) if int(s['section_no']) == n)
        workspace.state['pending'] = deepcopy(pending); workspace.refresh(); workspace.show_pending()
        self.show_detail()

    def choose(self):
        ids = self.tree.selection()
        if not ids or ids[0] not in self.rows:
            messagebox.showinfo('Chọn phương án', 'Chọn một hàng phương án đã tính.', parent=self.app)
            return
        row = self.rows[ids[0]]
        try:
            pending, option = row['pending'], row['option']
            if not pending or not row['project'] or section_signature(row['project']) != section_signature(pending['project']):
                raise ValueError('Đầu vào đã thay đổi hoặc chưa tính. Tính lại trước khi chọn.')
            if not pending.get('natural') and (not option or not row['attempt'].get('pass_check')):
                raise ValueError('Chỉ chọn phương án đạt các kiểm toán.')
            name = 'Không cần xử lý' if pending.get('natural') else row['attempt']['option']
            if not messagebox.askyesno('Xác nhận phương án',
                    f'STT {row["section"]["section_no"]}: chọn {name} để tổng hợp và xuất hồ sơ?', parent=self.app):
                return
            workspace = self.app._ai_analysis_workspace
            workspace.state['index'] = next(i for i, s in enumerate(workspace.state['sections'])
                if int(s['section_no']) == int(row['section']['section_no']))
            from ai_analysis_data import make_record
            record = make_record(workspace.state, row['section'], pending['project'], pending.get('before', {}),
                                 None if pending.get('natural') else option)
            if option:
                record['params'] = option.get('params', record['params'])
                record['notes'] = option.get('notes', record['notes'])
            record.pop('needs_review', None)
            workspace.save_record(record)
            getattr(self.app, '_batch_candidates', {}).pop(int(record['section_no']), None)
            self.refresh()
            self.app.status_text.set(f'Đã chọn {name} cho STT {record["section_no"]}.')
        except Exception as exc:
            messagebox.showerror('Chọn phương án', str(exc), parent=self.app)

    def run_ai_design(self):
        self.app._summary_selected_numbers = sorted({int(i.split(':')[0]) for i in self.tree.selection()})
        self.app._summary_run_ai(self.app._summary_selected_numbers)

    def calculate(self, use_ai, selected_numbers=None):
        if not self.app.require_full_license(): return
        workspace = self.app._ai_analysis_workspace
        if workspace.busy: return
        if not all(workspace.state.get(k) for k in ('geology_approved', 'boreholes_approved', 'sections_approved')):
            messagebox.showinfo('Đầu vào', 'Xác nhận phân đoạn, địa tầng và chỉ tiêu trước khi kiểm toán.', parent=self.app)
            return
        if use_ai and (not self.app.current_username or not self.app.current_login_key):
            messagebox.showinfo('AI', 'Đăng nhập trước khi gọi AI.', parent=self.app); return
        sections = deepcopy(workspace.state['sections']); selected = ({int(i) for i in selected_numbers}
            if selected_numbers is not None else {int(i) for i in self.tree.selection()})
        if selected: sections = [s for s in sections if int(s['section_no']) in selected]
        captured = deepcopy(workspace.state); provider = self.app.ai_provider_var.get()
        def action(progress):
            from ai_analysis_data import project_for_section, validate_analysis_project
            from soilfirm_ai_engine import before_treatment
            output = {}
            for i, s in enumerate(sections, 1):
                if workspace.cancel_event.is_set(): break
                n = int(s['section_no']); result = {}
                progress(f'Kiểm toán trước xử lý {i}/{len(sections)} · STT {n}')
                try:
                    hole = next(h for h in captured['boreholes'] if h['name'] == s['borehole_selected'])
                    p = project_for_section(captured['template'], s, captured['materials'], hole)
                    validate_analysis_project(p)
                    result['project'] = p
                    if use_ai:
                        from batch_hub import consult_batch_ai
                        result['ai_review'] = consult_batch_ai(self.app, provider,
                            'Rà dữ liệu kiểm toán trước xử lý. Trả JSON {"review":"nhận xét ngắn", "missing":[]}. '
                            'Chỉ chỉ ra số liệu thiếu hoặc bất nhất; không tự điền giá trị hoặc sửa công thức.',
                            {'section': s, 'borehole': hole, 'materials': captured['materials']}, workspace.cancel_event)
                        if result['ai_review'].get('missing'):
                            raise ValueError('AI yêu cầu đối chiếu: ' + str(result['ai_review']['missing']))
                    result['before'] = before_treatment(p)
                except InterruptedError: break
                except Exception as exc: result['error'] = str(exc)
                output[n] = result
            return output
        def done(output):
            self.app._before_section_results = {**getattr(self.app, '_before_section_results', {}), **output}
            self.refresh()
        workspace.run_job('AI rà dữ liệu và kiểm toán trước xử lý…' if use_ai else 'Kiểm toán trước xử lý hàng loạt…', action, done, keep_partial=True)


class RouteSectionSelector(ttk.Frame):
    """Load one section and retain manual designs independently by section/group."""
    def __init__(self, parent, app):
        super().__init__(parent, padding=8)
        self.app = app
        self.number = tk.StringVar()
        self.notice = tk.StringVar(value='Chọn STT đoạn tính trước khi tính.')
        self.active = None
        self.group = None
        self.base = None
        ttk.Label(self, text='STT đoạn tính').pack(side='left')
        self.combo = ttk.Combobox(self, textvariable=self.number, width=12)
        self.combo.pack(side='left', padx=8)
        self.combo.bind('<<ComboboxSelected>>', self.load)
        self.combo.bind('<Return>', self.load)
        ttk.Button(self, text='Nạp số liệu đoạn', command=self.load).pack(side='left')
        ttk.Label(self, textvariable=self.notice).pack(side='left', padx=10)

    def leave(self):
        """Save the departing editor before the shared widgets change groups."""
        self.stash()
        self._visible_group = None

    def refresh(self, group):
        self.combo.configure(values=[str(s['section_no']) for s in self.app._ai_analysis_workspace.state['sections']])
        if getattr(self, '_visible_group', None) == group:
            return
        number = self.active
        self.group = group
        self._visible_group = group
        if number is None:
            self.number.set('')
            self.notice.set('Chọn STT đoạn tính trước khi tính.')
        else:
            self.activate(number)

    def stash(self):
        if self.active is None or self.group is None:
            return
        app = self.app
        # Retain unfinished text as well as the last valid project values.
        try:
            app.collect()
        except (ValueError, TypeError):
            pass
        snapshot = capture_workspace(app, include_cdm=self.group == 6)
        snapshot['base'] = deepcopy(self.base)
        saved = getattr(app, '_route_manual_designs', {})
        saved.setdefault(self.active, {})[self.group] = snapshot
        app._route_manual_designs = saved

    def activate(self, number):
        app = self.app
        section = next((s for s in app._ai_analysis_workspace.state['sections'] if int(s['section_no']) == number), None)
        if section is None:
            self.active = None
            self.base = None
            self.number.set('')
            self.notice.set('STT đoạn tính không còn trong danh sách phân đoạn.')
            return False
        base = app._ai_analysis_workspace.build_project(section)
        saved = getattr(app, '_route_manual_designs', {}).get(number, {}).get(self.group, {})
        if saved and section_signature(saved['base']) != section_signature(base):
            saved = {}
        switching = getattr(app, '_switching_step', False)
        app._switching_step = True
        try:
            app._step_result_cache = {}
            app._choice_group_results = {}
            app.clear_cdm_results()
            if saved:
                restore_workspace(app, saved)
            else:
                app.project = deepcopy(base)
                app.populate()
                if self.group == 5:
                    app.treatment_vars['treatment'].set('PVD')
            if self.group == 6:
                app.refresh_cdm_inputs()
                if saved.get('widgets'):
                    app._cdm_saved_widgets(saved['widgets'])
            if self.group in (4,5):
                app.treatment_group.set('mechanical' if self.group == 4 else 'drainage')
                app.update_treatment_visibility()
            self.active, self.base = number, deepcopy(base)
            app._active_section_no = number
            app._active_section_length = section['length']
            self.number.set(str(number))
            self.notice.set(f'{base.station_from} – {base.station_to} · {base.station} · {base.borehole_name}')
        finally:
            app._switching_step = switching
        return True

    def load(self, event=None):
        app = self.app
        try:
            number = int(self.number.get().strip())
            if not any(int(s['section_no']) == number for s in app._ai_analysis_workspace.state['sections']):
                raise ValueError('STT đoạn tính không có trong danh sách phân đoạn.')
            self.stash()
            if self.activate(number):
                app.switch_step(self.group)
        except Exception as exc:
            messagebox.showerror('Nạp số liệu đoạn', str(exc), parent=app)

    def require_section(self):
        if self.active is None or self.group != self.app.current_step or self.number.get().strip() != str(self.active):
            raise ValueError('Chọn STT đoạn tính và nạp mặt cắt trước khi tính.')
        section = next((s for s in self.app._ai_analysis_workspace.state['sections'] if int(s['section_no']) == self.active), None)
        if section is None or section_signature(self.base) != section_signature(self.app._ai_analysis_workspace.build_project(section)):
            raise ValueError('Dữ liệu tuyến của đoạn đã thay đổi. Nạp lại mặt cắt trước khi tính.')

    def calculated(self, method=None):
        app = self.app
        if app.design_mode.get() != 'TÍNH TOÀN TUYẾN' or self.active is None:
            return
        from result_summary import collect_calculated_options, project_signature
        enabled = (False, self.group == 4, self.group == 5, self.group == 6)
        options, _ = collect_calculated_options(app, enabled)
        if method:
            options = [o for o in options if o.apply_payload.get('method') == method]
        n = self.active
        section = next(s for s in app._ai_analysis_workspace.state['sections'] if int(s['section_no']) == n)
        app._batch_runs = getattr(app, '_batch_runs', {})
        pending = deepcopy(app._batch_runs.get(n))
        if not pending or section_signature(pending['project']) != section_signature(self.base):
            pending = {'project': deepcopy(self.base), 'section': deepcopy(section),
                       'before': deepcopy(getattr(app, '_before_section_results', {}).get(n, {}).get('before', {})),
                       'natural': False, 'result': {'summary': {'attempts': []}, 'options': {}}}
        if pending.get('natural'):
            pending['result'] = {'summary': {'attempts': []}, 'options': {}}
        pending['natural'] = False
        result = pending.setdefault('result', {'summary': {'attempts': []}, 'options': {}})
        for o in options:
            special = deepcopy(app._cdm_reports.get(o.apply_payload.get('method')))
            record = {'opt_name': o.opt_name, 'params': o.opt_params_desc, 'payload': deepcopy(o.apply_payload),
                      'project_snapshot': deepcopy(o.source_project), 'residual': o.residual_cm, 'limit': o.limit_cm,
                      'status': 'ĐẠT' if o.is_pass else 'CHƯA ĐẠT', 'notes': o.tech_notes, 'special': special,
                      'method': o.apply_payload.get('method'), 'manual': True}
            if special:
                record['calculation_report'] = {'result': deepcopy(special.get('result', {}))}
            else:
                group = o.apply_payload.get('treatment_group')
                cached = getattr(app, '_choice_group_results', {}).get(group, {})
                if cached and project_signature(cached['project']) == project_signature(o.source_project):
                    record['calculation_report'] = {'locations': deepcopy(cached.get('results', []))}
            result['options'][o.opt_name] = record
            attempts = result['summary']['attempts']
            attempts[:] = [a for a in attempts if a.get('option') != o.opt_name]
            attempts.append({'option': o.opt_name, 'residual_cm': o.residual_cm,
                             'limit_cm': o.limit_cm, 'pass_check': o.is_pass})
        app._batch_runs[n] = pending
        self.stash()
        app.switch_step(16)
        app.status_text.set(f'Đã tính STT {n}; chọn hàng phương án và xác nhận để tổng hợp hồ sơ.')
        for iid in app.batch_design_view.tree.get_children():
            if iid.startswith(str(n) + ':'):
                app.batch_design_view.tree.selection_set(iid)
                app.batch_design_view.tree.see(iid)
                break


def capture_workspace(app, include_cdm=True):
    return {
        'project': deepcopy(app.project), 'steps': deepcopy(app._step_result_cache),
        'choice': deepcopy(app._choice_group_results), 'reports': deepcopy(app._cdm_reports),
        'widgets': app._cdm_saved_widgets() if include_cdm else [], 'method': app._cdm_active_method,
        'section_excel_path': app._section_excel_path,
        'active_section_no': app._active_section_no, 'active_section_length': app._active_section_length,
        'last_rows': deepcopy(getattr(app, 'last_rows', [])),
        'last_result_view': deepcopy(getattr(app, 'last_result_view', None)),
        'chart': deepcopy(app.chart_data),
        'vars': {k: v.get() for k, v in app.vars.items()},
        'treatment_vars': {k: v.get() for k, v in app.treatment_vars.items()},
        'stages': [[v.get() for v in row] for row in app.stage_vars],
        'stage_flags': {k: v.get() for k,v in app.stage_flags.items()},
        'input_flags': {name: getattr(app, name).get() for name in (
            'has_cw', 'has_expansion', 'expansion_side', 'main_treatment_var',
            'settlement_factor_var', 'mechanical_wait', 'mechanical_surcharge',
            'surcharge_enabled', 'vacuum_enabled', 'cdm_scope_var') if hasattr(app, name)},
        'flag_maps': {name: {key: value.get() for key, value in getattr(app, name, {}).items()}
                      for name in ('mechanical_enabled', 'expansion_use_main', 'treatment_flags')},
    }


def restore_workspace(app, saved):
    app.project = deepcopy(saved['project'])
    app.populate()
    for name in ('vars', 'treatment_vars'):
        for key, value in saved.get(name, {}).items():
            variable = getattr(app, name).get(key)
            if variable is not None: variable.set(value)
    for name, value in saved.get('input_flags', {}).items():
        if hasattr(app, name): getattr(app, name).set(value)
    for name, values in saved.get('flag_maps', {}).items():
        for key, value in values.items():
            variable = getattr(app, name, {}).get(key)
            if variable is not None: variable.set(value)
    for row, values in zip(app.stage_vars, saved.get('stages', [])):
        for variable, value in zip(row, values): variable.set(value)
    for key, value in saved.get('stage_flags', {}).items():
        if int(key) in app.stage_flags: app.stage_flags[int(key)].set(value)
    app._step_result_cache = {int(k): deepcopy(v) for k,v in saved.get('steps', {}).items()}
    app._choice_group_results = deepcopy(saved.get('choice', {}))
    app._cdm_reports = deepcopy(saved.get('reports', {}))
    app._cdm_active_method = saved.get('method', 'standard')
    app._section_excel_path = saved.get('section_excel_path')
    app._active_section_no = saved.get('active_section_no')
    app._active_section_length = saved.get('active_section_length', 0)
    app.last_rows = deepcopy(saved.get('last_rows', []))
    app.last_result_view = deepcopy(saved.get('last_result_view'))
    app.chart_data = deepcopy(saved.get('chart', {}))
    app._cdm_saved_widgets(saved.get('widgets', []))


class SingleTreatmentBoard(SegmentBoard):
    """Reuse the treatment table layout with single-section calculation records."""
    def __init__(self, parent, app):
        self._single_result_tabs = True
        super().__init__(parent, app, after=True)
        bar = self.comparison_page.winfo_children()[0]
        for child in bar.winfo_children():
            child.destroy()
        template_bar = ttk.Frame(bar)
        app.project_template_bar = template_bar
        template_bar.pack(side='left', padx=(0, 8))
        ttk.Button(template_bar, text='Hướng dẫn nhập dữ liệu', command=app.show_import_guide).pack(side='left', padx=6)
        ttk.Button(bar, text='Nhập dữ liệu phân đoạn', command=app.open_section_excel).pack(side='left', padx=(0, 8))
        ttk.Label(bar, text='STT phân đoạn:').pack(side='left')
        app.excel_section_combo = ttk.Combobox(bar, textvariable=app.excel_section_no, width=9, state='readonly')
        app.excel_section_combo.pack(side='left', padx=5)
        ttk.Button(bar, text='Nạp đoạn', command=app.load_excel_section).pack(side='left')
        ttk.Button(bar, text='Tính toán', command=app.open_batch_calculation).pack(side='left', padx=(8, 0))
        ttk.Button(bar, text='Tính bằng AI', command=app.open_ai_batch_calculation).pack(side='left', padx=(6, 0))
        self.comparison_page.winfo_children()[1].pack_forget()
        ttk.Label(self.comparison_page, textvariable=app.excel_section_notice, padding=(8, 2)).pack(fill='x', after=bar)
        ttk.Button(bar, text='Chọn phương án', command=self.choose).pack(side='left', padx=6)
        self.summary_view = SingleCombinedResults(self.summary_page, app)
        self.summary_view.pack(fill='both', expand=True)
        self.result_tabs.bind('<<NotebookTabChanged>>', lambda event: self.summary_view.refresh())
        self.tree.bind('<Double-1>', self.show_detail)
        self.result_menu = tk.Menu(self, tearoff=False)
        self.result_menu.add_command(label='Chọn phương án', command=self.choose)
        self.tree.bind('<Button-3>', self.open_result_menu)

    def refresh(self):
        config = self.app.calculation_settings()
        if config['summary']:
            self.result_tabs.tab(self.summary_page, state='normal')
        else:
            self.result_tabs.tab(self.comparison_page, state='normal')
            self.result_tabs.select(self.comparison_page)
            self.result_tabs.hide(self.summary_page)
        if config['after'] or not config['summary']:
            self.result_tabs.tab(self.comparison_page, state='normal',
                                 text='So sánh phương án' if config['after'] else 'Tính toán')
        else:
            self.result_tabs.select(self.summary_page)
            self.result_tabs.hide(self.comparison_page)
        from result_summary import collect_calculated_options
        options, missing = collect_calculated_options(self.app)
        selected = self.tree.selection()
        self.rows.clear()
        self.tree.delete(*self.tree.get_children())
        record = getattr(self.app, '_single_selected_option', None) or {}
        for i, item in enumerate(options):
            p = item.source_project
            iid = str(i)
            is_selected = (record.get('signature') == section_signature(p)
                           and record.get('name') == item.opt_name
                           and record.get('payload') == item.apply_payload)
            status = 'Đã chọn' if is_selected else 'Chờ lựa chọn'
            checked = 'ĐẠT' if item.is_pass else 'CHƯA ĐẠT'
            self.rows[iid] = {'item': item, 'status': status}
            self.tree.insert('', 'end', iid=iid, values=[
                getattr(self.app, '_active_section_no', None) or 1, str(p.station_from or '') + ' – ' + str(p.station_to or ''),
                fmt(getattr(self.app, '_active_section_length', None)), p.station, fmt(p.h_design), fmt(p.height), p.borehole_name,
                '; '.join(f'{soil.name}: {soil.thickness:.2f}' for soil in p.soils),
                item.opt_name, item.opt_params_desc, fmt(item.residual_cm),
                fmt(item.limit_cm), checked, status], tags=('pass' if item.is_pass else 'fail',))
        for record in getattr(self.app, '_single_batch_records', []):
            p = record.get('project_snapshot')
            if p is None:continue
            n = record.get('section_no'); iid = f'batch:{n}'
            self.rows[iid] = {'record': record}
            checked = record.get('status', 'CHƯA TÍNH')
            params = record.get('params', '')
            lines = treatment_display_details({}, record, p)
            if lines:params = '; '.join(lines[:3]) + ('\n' + '; '.join(lines[3:6]) if len(lines) > 3 else '')
            self.tree.insert('', 'end', iid=iid, values=[n,
                str(p.station_from or '') + ' – ' + str(p.station_to or ''),
                fmt(record.get('length')), p.station, fmt(p.h_design), fmt(p.height),
                p.borehole_name, '; '.join(f'{soil.name}: {soil.thickness:.2f}' for soil in p.soils),
                record.get('opt_name', ''), params, fmt(record.get('residual')),
                fmt(record.get('limit')), checked,
                'Đã chọn theo ưu tiên' if checked == 'ĐẠT' else 'Chưa có phương án đạt'],
                tags=('pass' if checked == 'ĐẠT' else 'fail',))
        for i, failure in enumerate(getattr(self.app, '_single_batch_failures', [])):
            if self.tree.exists(f'batch:{failure.get("section_no")}'):continue
            iid = f'error:{i}';self.rows[iid] = {'failure': failure}
            self.tree.insert('', 'end', iid=iid, values=[failure.get('section_no'),
                *['—'] * 11, 'LỖI', failure.get('reason', '')], tags=('fail',))
        for iid in selected:
            if self.tree.exists(iid):self.tree.selection_add(iid)
        self.note.set('Kết quả đã tính của một đoạn; chọn hàng để xem chi tiết. '
                      + ('Chưa có kết quả: ' + '; '.join(missing) if missing else ''))
        self.show_detail()
        self.summary_view.refresh()

    def show_detail(self, event=None):
        ids = self.tree.selection()
        text = 'Chọn một hàng để xem số liệu và kết quả xử lý.'
        if ids and ids[0] in self.rows:
            row = self.rows[ids[0]]
            if 'failure' in row:
                failure = row['failure']
                text = f'STT {failure.get("section_no")}: {failure.get("reason", "")}'
                self._set_detail_text(text)
                return
            if 'record' in row:
                record = row['record'];p = record['project_snapshot']
                lines = treatment_display_details({}, record, p)
                text = (f'STT {record.get("section_no")} · {record.get("status", "")}\n'
                        f'Dự án: {p.name} · Mặt cắt: {p.station} · Lỗ khoan: {p.borehole_name}\n'
                        f'Phương án: {record.get("opt_name", "")}\n'
                        f'S꜀ dư = {fmt(record.get("residual"))} cm; [ΔS] = {fmt(record.get("limit"))} cm\n'
                        + '\n'.join(lines) + '\n' + record.get('notes', ''))
                if record.get('calculation_report'):
                    text += '\n\nChi tiết tính toán:\n' + json.dumps(record['calculation_report'], ensure_ascii=False, indent=2, default=str)
                self._set_detail_text(text)
                return
            item = row['item']; p = item.source_project
            text = (f'Dự án: {p.name} · Mặt cắt: {p.station} · Lỗ khoan: {p.borehole_name}\n'
                    f'Hₜₖ = {fmt(p.h_design)} m; Hₜₜ = {fmt(p.height)} m\n'
                    f'Phương án: {item.opt_name} · {row["status"]}\n'
                    f'Thông số xử lý: {item.opt_params_desc}\n'
                    f'S꜀ dư = {fmt(item.residual_cm)} cm; [ΔS] = {fmt(item.limit_cm)} cm\n'
                    f'Kiểm toán: {"ĐẠT" if item.is_pass else "CHƯA ĐẠT"}\n'
                    f'Thời gian: {item.time_desc}\n{item.tech_notes}\n\nĐịa tầng:\n')
            for soil in p.soils:
                text += f'{soil.name}: h = {soil.thickness:.2f} m\n'
        self.detail.configure(state='normal')
        self.detail.delete('1.0', 'end'); self.detail.insert('end', text)
        self.detail.configure(state='disabled')

    def _set_detail_text(self, text):
        self.detail.configure(state='normal')
        self.detail.delete('1.0', 'end');self.detail.insert('end', text)
        self.detail.configure(state='disabled')

    def open_result_menu(self, event):
        iid = self.tree.identify_row(event.y)
        if not iid or 'item' not in self.rows.get(iid, {}):return
        self.tree.selection_set(iid)
        self.result_menu.tk_popup(event.x_root, event.y_root)

    def choose(self):
        ids = self.tree.selection()
        if not ids or ids[0] not in self.rows:
            messagebox.showinfo('Chọn phương án', 'Chọn một hàng phương án đã tính.', parent=self.app)
            return
        if 'item' not in self.rows[ids[0]]:
            self.show_detail()
            return
        item = self.rows[ids[0]]['item']
        view = self.app.single_selection_view
        view.refresh()
        target = next((str(i) for i, candidate in enumerate(view.options)
                       if candidate.opt_name == item.opt_name
                       and candidate.apply_payload == item.apply_payload
                       and section_signature(candidate.source_project) == section_signature(item.source_project)), None)
        if target is None:
            messagebox.showinfo('Chọn phương án', 'Kết quả đã thay đổi. Làm mới bảng trước khi chọn.', parent=self.app)
            return
        view.tree.selection_set(target)
        view.choose()
        self.refresh()


class SingleCombinedResults(ttk.Frame):
    """Một dòng mỗi phân đoạn, chỉ lấy kết quả đã lưu trong luồng Phân đoạn."""
    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app
        bar = ttk.Frame(self, padding=8)
        bar.pack(fill='x')
        ttk.Label(bar, text='Một phân đoạn một dòng · Trước xử lý và sau xử lý đặt cạnh nhau').pack(side='left')
        ttk.Button(bar, text='Làm mới bảng', command=self.refresh).pack(side='right')
        identity = ['STT', 'Lý trình', 'Chiều dài (m)', 'Mặt cắt', 'Hₜₖ (m)', 'Hₜₜ (m)', 'Lỗ khoan']
        before = ['Trước · Sᵢ (cm)', 'Trước · S꜀ (cm)', 'Trước · S (cm)',
                  'Trước · S꜀ dư (cm)', 'Trước · [ΔS] (cm)', 'Trước · Đánh giá']
        after = ['Sau · Phương án chọn', 'Sau · Thông số xử lý', 'Sau · Sᵢ (cm)',
                 'Sau · S꜀ (cm)', 'Sau · S (cm)', 'Sau · S꜀ dư (cm)',
                 'Sau · [ΔS] (cm)', 'Sau · Đánh giá', 'Trạng thái']
        self.tree = table(self, identity+before+after,
            [65,190,100,150,100,100,130]+[145]*5+[160]+[240,400]+[145]*5+[160,250])
        self.tree.tag_configure('odd', background='#EDF5FA')
        self.tree.tag_configure('even', background='#FFFFFF')

    def refresh(self):
        from result_summary import project_signature
        app = self.app
        if app.design_mode.get() != 'TÍNH MỘT ĐOẠN':
            return
        config = app.calculation_settings()
        indices = list(range(7))
        if config['before']: indices.extend(range(7, 13))
        if config['after']: indices.extend(range(13, 22))
        self.tree.configure(displaycolumns=tuple(f'c{i}' for i in indices))
        old_order = list(self.tree.get_children())
        selected_rows = self.tree.selection()
        self.tree.delete(*self.tree.get_children())
        records = {int(r['section_no']): r for r in getattr(app, '_single_batch_records', [])
                   if r.get('section_no') is not None}
        failures = {int(r['section_no']): r.get('reason', '')
                    for r in getattr(app, '_single_batch_failures', []) if r.get('section_no') is not None}
        current = int(getattr(app, '_active_section_no', None) or 1)
        numbers = list(records)
        for n in failures:
            if n not in numbers:numbers.append(n)
        combo = getattr(app, 'excel_section_combo', None)
        for value in combo.cget('values') if combo is not None else ():
            try:n = int(value)
            except (TypeError, ValueError):continue
            if n not in numbers:numbers.append(n)
        if current not in numbers:numbers.append(current)
        chosen = getattr(app, '_single_selected_option', None) or {}
        natural = getattr(app, '_choice_group_results', {}).get('natural') or {}
        for index, n in enumerate(numbers):
            rec = records.get(n, {})
            p = rec.get('before_project') or rec.get('project_snapshot')
            before = rec.get('before') or {}
            active = n == current
            stale = False
            if active:
                p = app.project
                if rec.get('before_project') is not None:
                    stale = project_signature(p) != project_signature(rec['before_project'])
                if natural.get('project') is not None and project_signature(natural['project']) == project_signature(p):
                    rows = natural.get('results') or []
                    candidates = [r for r in rows if r.get('Sc_dư_cm') is not None]
                    if candidates:
                        row = max(candidates, key=lambda r:r['Sc_dư_cm'])
                        before = {'sc_cm':row.get('Sc_cuối_cm'), 'si_cm':row.get('Si_cm'),
                                  'total_cm':row.get('St_cuối_cm'), 'residual_cm':row['Sc_dư_cm']}
                        if before['total_cm'] is None and before['sc_cm'] is not None and before['si_cm'] is not None:
                            before['total_cm'] = before['sc_cm']+before['si_cm']
                if chosen and chosen.get('signature') == section_signature(p):
                    rec = dict(chosen, opt_name=chosen.get('name'), status='ĐẠT' if chosen.get('pass_check') else 'CHƯA ĐẠT')
                elif chosen:
                    stale = True
            limit = rec.get('limit') if rec.get('limit') is not None else getattr(p, 'residual_limit_cm', None)
            residual_before = before.get('residual_cm', before.get('total_cm'))
            sc, si, total = settlement_display_values(rec.get('calculation_report') or rec.get('special'),
                                                       rec.get('project_snapshot') or p)
            if sc is None and total is None and active and rec:
                group = (rec.get('payload') or {}).get('treatment_group')
                cached = getattr(app, '_choice_group_results', {}).get(group) or {}
                src = cached.get('project')
                if src is not None and project_signature(src) == project_signature(p):
                    sc, si, total = settlement_display_values({'rows':cached.get('results', [])}, src)
            status = failures.get(n) or ('Đầu vào đã thay đổi — cần tính lại' if stale else
                'Đã chọn phương án' if rec.get('status') == 'ĐẠT' else 'Chưa chọn phương án')
            before_check = ('CẦN TÍNH LẠI' if stale and before else
                'ĐẠT' if residual_before is not None and limit is not None and residual_before <= limit else
                'CHƯA ĐẠT' if residual_before is not None and limit is not None else 'CHƯA TÍNH')
            after_check = 'CẦN TÍNH LẠI' if stale and rec else rec.get('status', 'CHƯA CHỌN')
            params = rec.get('params', '')
            if rec and p is not None:
                lines = treatment_display_details({}, rec, rec.get('project_snapshot') or p)
                if lines:params = '; '.join(lines)
            identity = [n, (str(p.station_from or '')+' – '+str(p.station_to or '')) if p else '—',
                fmt(rec.get('length', getattr(app, '_active_section_length', None) if active else None)),
                p.station if p else '—', fmt(p.h_design if p else None), fmt(p.height if p else None),
                p.borehole_name if p else '—']
            values = identity + [fmt(before.get('si_cm')), fmt(before.get('sc_cm')),
                fmt(before.get('total_cm')), fmt(residual_before), fmt(limit), before_check,
                rec.get('opt_name') or 'Chưa chọn phương án', params or '—', fmt(si), fmt(sc),
                fmt(total), fmt(rec.get('residual')), fmt(limit), after_check, status]
            self.tree.insert('', 'end', iid=str(n), values=values, tags=('odd' if index%2 else 'even',))
        position = 0
        for iid in old_order:
            if self.tree.exists(iid):
                self.tree.move(iid, '', position);position += 1
        for iid in selected_rows:
            if self.tree.exists(iid):self.tree.selection_add(iid)


def workflow_pdf_records(app):
    """Capture only records belonging to the active workflow, in table order."""
    from result_summary import project_signature
    route = app.design_mode.get() == 'TÍNH TOÀN TUYẾN'
    board = app.batch_design_view if route else app.single_treatment_summary_view
    previous_order = list(board.tree.get_children())
    board.refresh()
    position = 0
    for iid in previous_order:
        if board.tree.exists(iid):
            board.tree.move(iid, '', position)
            position += 1
    order = []
    for iid in board.tree.get_children():
        values = board.tree.item(iid, 'values')
        if values:
            try:n = int(values[0])
            except (TypeError, ValueError):continue
            if n not in order:order.append(n)
    records = {int(r['section_no']): deepcopy(r) for r in
               getattr(app, '_saved_sections_data' if route else '_single_batch_records', [])
               if r.get('section_no') is not None}
    entries = []
    if route:
        sections = {int(s['section_no']): s for s in app._ai_analysis_workspace.state.get('sections', [])}
        for n in order:
            rec = records.get(n) or {}
            section = sections.get(n)
            try:current = board.current_project(section) if section else None
            except Exception:current = None
            pending = getattr(app, '_batch_runs', {}).get(n) or {}
            before_row = getattr(app, '_before_section_results', {}).get(n) or {}
            before = before_row.get('before') or pending.get('before') or rec.get('before')
            before_project = before_row.get('project') or pending.get('project') or rec.get('before_project')
            # A saved candidate is not a user-selected treatment record.
            entries.append({'number': n, 'before': before, 'before_project': before_project,
                            'record': rec, 'current': current})
    else:
        current_number = getattr(app, '_active_section_no', None) or 1
        selected = getattr(app, '_single_selected_option', None) or {}
        natural = getattr(app, '_choice_group_results', {}).get('natural') or {}
        for n in order:
            rec = records.get(n) or {}
            current = None
            before = rec.get('before');before_project = rec.get('before_project')
            if n == current_number:
                current = app.project
                if natural and project_signature(natural['project']) == project_signature(current):
                    before = {'calculated': True, 'evaluation_day': natural.get('days')}
                    before_project = natural['project']
                if selected and selected.get('signature') == section_signature(current):
                    rec = deepcopy(selected)
                    rec['opt_name'] = rec.get('name')
                    rec['status'] = 'ĐẠT' if rec.get('pass_check') else 'CHƯA ĐẠT'
                    rec['calculation_report'] = rec.get('special')
            entries.append({'number': n, 'before': before, 'before_project': before_project,
                            'record': rec, 'current': current})
    return entries


def export_workflow_pdf(app, whole=False):
    """Choose stage(s) and merge PDF reports using saved calculation snapshots."""
    if not app.require_full_license():return
    if getattr(app, '_single_batch_busy', False) or getattr(app._ai_analysis_workspace, 'busy', False):
        messagebox.showinfo('Xuất PDF', 'Chờ tính toán hoàn thành trước khi xuất báo cáo.', parent=app)
        return
    try:entries = workflow_pdf_records(app)
    except Exception as exc:
        messagebox.showerror('Xuất PDF', str(exc), parent=app);return
    if not entries:
        messagebox.showinfo('Xuất PDF', 'Chưa có phân đoạn trong bảng tổng hợp của luồng đang chọn.', parent=app)
        return
    mode = app.design_mode.get()
    popup = tk.Toplevel(app)
    popup.title('Xuất PDF toàn tuyến' if whole else 'Xuất PDF theo đoạn')
    popup.transient(app);popup.resizable(False, False)
    body = ttk.Frame(popup, padding=16);body.pack(fill='both', expand=True)
    ttk.Label(body, text='Dữ liệu: ' + ('Tính toàn tuyến' if mode == 'TÍNH TOÀN TUYẾN' else 'Tính một đoạn')).pack(anchor='w', pady=(0, 8))
    selected = tk.StringVar(value=str(entries[0]['number']))
    if not whole:
        line = ttk.Frame(body);line.pack(fill='x', pady=4)
        ttk.Label(line, text='STT phân đoạn:').pack(side='left', padx=(0, 8))
        ttk.Combobox(line, textvariable=selected, values=[str(e['number']) for e in entries],
                     state='readonly', width=12).pack(side='left')
    before = tk.BooleanVar(value=True);after = tk.BooleanVar(value=True);logo = tk.BooleanVar(value=True)
    ttk.Checkbutton(body, text='Trước xử lý', variable=before).pack(anchor='w', pady=3)
    ttk.Checkbutton(body, text='Sau xử lý', variable=after).pack(anchor='w', pady=3)
    ttk.Checkbutton(body, text='Hiển thị logo', variable=logo).pack(anchor='w', pady=3)
    ttk.Label(body, text='Thứ tự trong bảng tổng hợp; mỗi phân đoạn: trước xử lý → sau xử lý.',
              wraplength=440).pack(anchor='w', pady=8)

    def publish():
        if app.design_mode.get() != mode:
            messagebox.showinfo('Xuất PDF', 'Luồng tính đã thay đổi. Mở lại lệnh xuất PDF.', parent=popup);return
        if not before.get() and not after.get():
            messagebox.showinfo('Xuất PDF', 'Chọn trước xử lý hoặc sau xử lý.', parent=popup);return
        try:
            fresh = workflow_pdf_records(app)
            chosen = fresh if whole else [e for e in fresh if str(e['number']) == selected.get()]
            if not chosen:raise ValueError('Phân đoạn không còn trong bảng tổng hợp.')
            jobs = [];missing = []
            from result_summary import project_signature
            for entry in chosen:
                n = entry['number'];rec = entry['record'];current = entry['current']
                if before.get():
                    source = entry['before_project']
                    if not entry['before'] or source is None:
                        missing.append(f'STT {n}: chưa có kết quả trước xử lý.')
                    elif current is not None and project_signature(current) != project_signature(source):
                        missing.append(f'STT {n}: dữ liệu trước xử lý đã thay đổi; cần tính lại.')
                    else:
                        jobs.append((n, 'Trước xử lý', deepcopy(source), 'natural', None,
                                     (entry['before'] or {}).get('evaluation_day')))
                if after.get():
                    source = rec.get('project_snapshot');payload = rec.get('payload') or {}
                    group = payload.get('treatment_group')
                    if not rec or source is None or rec.get('status') != 'ĐẠT':
                        missing.append(f'STT {n}: chưa có phương án sau xử lý đã chọn và đạt yêu cầu.')
                    elif group in ('none', 'natural') or rec.get('opt_name') in ('Không cần xử lý', 'Lún khi chưa xử lý'):
                        missing.append(f'STT {n}: không cần xử lý; chỉ có báo cáo trước xử lý.')
                    elif current is not None and project_signature(current) != project_signature(source):
                        missing.append(f'STT {n}: dữ liệu sau xử lý đã thay đổi; cần tính và chọn lại.')
                    else:
                        special = rec.get('calculation_report') or rec.get('special')
                        scope = 'cdm' if group == 'cdm' else 'treated'
                        if scope == 'cdm' and not special:
                            missing.append(f'STT {n}: thiếu kết quả trộn sâu đã tính.')
                        else:
                            jobs.append((n, 'Sau xử lý', deepcopy(source), scope,
                                         {'method': payload.get('method', rec.get('method', 'standard')), 'data': special}, None))
            if missing:
                messagebox.showwarning('Chưa đủ kết quả xuất PDF', '\n'.join(missing), parent=popup);return
            if not jobs:raise ValueError('Không có kết quả phù hợp để xuất.')
            dest = filedialog.asksaveasfilename(parent=popup, title=popup.title(), defaultextension='.pdf',
                    initialfile='Bao_cao_toan_tuyen.pdf' if whole else f'Bao_cao_STT_{selected.get()}.pdf',
                    filetypes=[('PDF', '*.pdf')])
            if not dest:return
            import os, tempfile
            from pathlib import Path
            from pypdf import PdfWriter
            from report_pdf import create_report
            writer = PdfWriter()
            try:
                with tempfile.TemporaryDirectory(prefix='soilfirm_pdf_', dir=os.path.dirname(os.path.abspath(dest))) as folder:
                    for i, (n, stage, project, scope, special, days) in enumerate(jobs):
                        project.calculation_results = {}
                        part = Path(folder) / f'{i}.pdf'
                        create_report(project, str(part), scope=scope, evaluation_days=days,
                            special_method=special['method'] if scope == 'cdm' else None,
                            special_data=special['data'] if scope == 'cdm' else None,
                            show_logo=logo.get(), language='en' if app.ui_language.get() == 'English' else 'vi')
                        writer.append(str(part), outline_item=f'STT {n} · {stage}')
                    combined = Path(folder) / 'combined.pdf'
                    with combined.open('wb') as stream:writer.write(stream)
                    writer.close()
                    os.replace(combined, dest)
            finally:writer.close()
            popup.destroy()
            messagebox.showinfo('Xuất PDF', f'Đã xuất {len(chosen)} phân đoạn theo thứ tự bảng tổng hợp:\n{dest}', parent=app)
        except Exception as exc:messagebox.showerror('Xuất PDF', str(exc), parent=popup)
    controls = ttk.Frame(body);controls.pack(fill='x', pady=(8, 0))
    ttk.Button(controls, text='Đóng', command=popup.destroy).pack(side='right')
    ttk.Button(controls, text='Xuất PDF', command=publish).pack(side='right', padx=8)
