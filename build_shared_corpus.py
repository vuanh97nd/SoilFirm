"""Đọc bản sao nguồn, tạo kiến thức cấu trúc. Không lưu trị số thí nghiệm vào RAG.

Chạy: python build_shared_corpus.py --inventory source_inventory.json --output .
File inventory gồm relative, id, status và path bản sao tải từ Drive.
"""
import argparse, hashlib, json, math, re, sys
from pathlib import Path
from collections import Counter
ROOT=Path(__file__).parent
sys.path.insert(0,str(ROOT.parent/'server_shared_memory'))
import ai_analysis_data as d
from geotech_memory import packed, digest, normalized
from openpyxl.utils import column_index_from_string

HEADER=re.compile(r'dung trong|khoi luong the tich|he so rong|chi so nen|chi so no|co ket|suc khang cat|cat canh|luc dinh|goc ma sat|lo khoan|ho khoan|so hieu mau|ma lop|lop dat|n.?spt|\bspt\b|\bsu\b|\bcv\b|\bcc\b|\bcs\b|\bpc\b|\be0\b|gamma|void ratio|consolidation|shear|friction|borehole|sample no|undrained|\bphi\b|ap luc|e.?log.?p|cv.?log.?p|cao do|chieu sau|do sau|be day|chieu day|ly trinh|mat cat|phan doan|he so tham|mo dun|\bset\b|\bcat\b|\bbun\b|\bsoi\b|\bsan\b|clay|sand|silt|gravel|trang thai')
CAUTION='Tham khảo nhãn nguồn, không phải ánh xạ đã xác nhận. Python đọc số; không suy đơn vị, không tạo số liệu; kiểm tra nhóm thí nghiệm trong file hiện tại.'

def compact_text(text):
    return re.sub(r'\s+',' ',str(text)).strip()

def safe_text(text):
    """Chỉ giữ tên/ký hiệu; che số đo lẫn trong đoạn văn. Giữ đơn vị và cấp áp lực nhãn đường cong."""
    text=compact_text(text);saved=[]
    pattern=r'(?:\b(?:e|Cv)\s*[_₀]?\s*\d+(?:[.,]\d+)?(?:\s*[-–]\s*Cv\s*\d+(?:[.,]\d+)?)?|\b10\s*(?:[eE]|\^)?\s*[-−]\s*\d+|(?:10\s*\^?\s*[-−]\s*\d+\s*)?(?:kgf?|kG|g|T|kN|cm|m)\s*[/²³^0-9−-]+(?:cm|m|s|sec)?[²³0-9]*(?:/s)?)'
    def protect(match):
        saved.append(match[0]);return '§'+chr(65+len(saved)-1)+'§'
    text=re.sub(pattern,protect,text,flags=re.I)
    text=re.sub(r'(?<![\w])[-+]?\d+(?:[.,]\d+)?','[số]',text)
    for i,token in enumerate(saved):text=text.replace('§'+chr(65+i)+'§',token)
    return text[:500]

def numeric(value):
    if isinstance(value,bool):return False
    if isinstance(value,(int,float)):return math.isfinite(value)
    return isinstance(value,str) and bool(re.fullmatch(r'[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?',value.strip()))

def main(inventory, output):
    output.mkdir(parents=True,exist_ok=True)
    records=json.loads(inventory.read_text(encoding='utf-8'))
    registry=json.loads((ROOT/'parameter_registry.json').read_text(encoding='utf-8'))
    catalog={'version':1,'origin':'My Drive/Dev/Dao tao AI','purpose':CAUTION,'templates':[],'files':[],'unavailable':[]}
    knowledge=[];report=[]
    def add_reference(record, section, labels, topic):
        # Mỗi mảnh nhỏ để vừa ngân sách RAG 3000 ký tự; không cắt mất JSON.
        batch=[]
        def flush():
            if not batch:return
            payload={'topic':topic,'title':record['title']+' / '+section+' / '+str(len(knowledge)+1),
                     'body':CAUTION+'\n'+packed(batch),
                     'source':record['relative']+'; '+section+'; Drive '+record['id']+'; SHA-256 '+record['sha256']}
            if len(packed(payload))>2700:raise ValueError('Mảnh kiến thức quá dài.')
            knowledge.append({'id':digest(['knowledge',payload]),'kind':'knowledge','payload':payload,'state':'pending'})
            batch.clear()
        for label in labels:
            if len(packed(batch+[label]))>850:flush()
            batch.append(label)
        flush()
    for record in records:
        result={'file':record['relative'],'id':record['id'],'status':record['status'],'sections':[],'warnings':[]}
        if record['status']!='downloaded':
            result['warnings'].append(record.get('error','Không tải được'));catalog['unavailable'].append(result);report.append(result);continue
        path=Path(record['path']);record['sha256']=hashlib.sha256(path.read_bytes()).hexdigest();result['sha256']=record['sha256']
        ext=Path(record['title']).suffix.lower()
        try:
            if ext in ('.xlsx','.xlsm','.xls'):
                book=d.open_source_excel(path)
                try:
                    result['warnings'].extend(getattr(book,'_soilfirm_read_warnings',[]))
                    for sheet in book:
                        counts=Counter();errors=[];source_labels=[];seen_labels=set()
                        for row in sheet.iter_rows():
                            for cell in row:
                                value=cell.value
                                if value is None:continue
                                counts['nonempty']+=1
                                if cell.data_type=='e':
                                    counts['errors']+=1
                                    if len(errors)<25:errors.append({'cell':cell.coordinate,'error':str(value)})
                                elif numeric(value):counts['numeric']+=1
                                elif isinstance(value,str):
                                    counts['text']+=1;text=compact_text(value)
                                    if len(text)<=500 and HEADER.search(normalized(text)):
                                        label=safe_text(text)
                                        if label not in seen_labels:
                                            seen_labels.add(label);source_labels.append({'label':label,'cell':cell.coordinate})
                        context=''
                        try:context=d._excel_header_context(sheet)
                        except ValueError as exc:result['warnings'].append(str(exc))
                        columns,_,_=d._strict_mapping_columns({'header_context':context,'excel_rows':[]})
                        candidates={column_index_from_string(x['column']):x['field'] for x in d._excel_review_column_evidence(context) if x['field'] in registry}
                        evidence=[]
                        for col in columns:
                            # Hàm header của bộ đọc đã dừng trước dòng mẫu. Giữ
                            # cấp áp lực và số mũ đơn vị trong header xác định.
                            text=d._unit_label(col['label'])
                            pressure=bool(re.search(r'void ratio.*pressure|he so rong.*ap luc|cv.*(?:ap luc|pressure|co ket)',text))
                            if not pressure:
                                col['label']=' / '.join(part for part in col['label'].split(' / ') if not re.fullmatch(r'[+-]?\d+(?:[.,]\d+)?',part.strip()))
                            if col['cot_id'] in candidates:col['candidate_id']=candidates[col['cot_id']]
                            matches=[x for x in source_labels if column_index_from_string(re.match(r'[A-Z]+',x['cell'])[0])==col['cot_id']]
                            evidence.append({'label':col['label'],'unit':col.get('unit',''),'group':col.get('group',''),
                                             'candidate_id':col.get('candidate_id','unknown'),
                                             'cells':[x['cell'] for x in matches[:6]],'column':col['cot_id']})
                        if columns:catalog['templates'].append({'file':record['title'],'source_path':record['relative'],'sheet':sheet.title,'sha256':record['sha256'],'columns':columns})
                        add_reference(record,'sheet '+sheet.title,evidence or source_labels,'Biểu mẫu Excel địa kỹ thuật')
                        # Giữ các nhãn phát hiện ngoài vùng header đầu; không bỏ bảng con/cấp áp lực.
                        header_parts={normalized(part.strip()) for c in columns for part in c['label'].split(' / ')}
                        extras=[x for x in source_labels if normalized(x['label']) not in header_parts]
                        if evidence and extras:add_reference(record,'sheet '+sheet.title+' / nhãn bổ sung',extras,'Nhãn bảng con và thí nghiệm')
                        result['sections'].append({'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'counts':dict(counts),'errors':errors,'header_columns':len(columns),'source_labels':len(source_labels)})
                finally:book.close()
                result['status']='read_excel'
            elif ext=='.json':
                obj=json.loads(path.read_text(encoding='utf-8-sig'));types={};nodes=0
                def walk(value, key='$'):
                    nonlocal nodes
                    nodes+=1;types.setdefault(key,set()).add(type(value).__name__)
                    if isinstance(value,dict):
                        for k,v in value.items():walk(v,key+'.'+str(k))
                    elif isinstance(value,list):
                        for child in value:walk(child,key+'[]')
                walk(obj)
                labels=[{'path':k,'types':sorted(v)} for k,v in sorted(types.items())]
                add_reference(record,'cấu trúc JSON',labels,'Cấu trúc dự án SoilFirm')
                result['sections']=[{'nodes_scanned':nodes,'key_paths':len(types)}];result['status']='read_json'
            elif ext=='.dxf':
                import ezdxf
                doc=ezdxf.readfile(path);counts=Counter();labels=[];seen=set();entities_seen=set()
                # Cả model/layout và block, không chỉ những text nhìn thấy trên model.
                spaces=list(doc.layouts)+list(doc.blocks)
                for space in spaces:
                    for entity in space:
                        handle=entity.dxf.get('handle','')
                        if handle and handle in entities_seen:continue
                        if handle:entities_seen.add(handle)
                        counts[entity.dxftype()]+=1
                        if entity.dxftype() not in ('TEXT','MTEXT','ATTRIB','ATTDEF'):continue
                        text=entity.plain_text() if entity.dxftype()=='MTEXT' else entity.dxf.get('text','')
                        text=compact_text(text)
                        if HEADER.search(normalized(text)) and len(text)<=500:
                            text=safe_text(text)
                            if text not in seen:
                                seen.add(text);labels.append({'label':text,'handle':entity.dxf.get('handle',''),'space':space.name})
                add_reference(record,'text DXF',labels,'Cấu trúc hình trụ và địa tầng CAD')
                result['sections']=[{'entities':dict(counts),'source_labels':len(labels)}];result['status']='read_dxf'
            elif ext=='.pdf':
                import fitz
                doc=fitz.open(path)
                try:
                    for i,page in enumerate(doc):
                        text=page.get_text();method='text';labels=[]
                        if not text.strip():
                            # OCR là gợi ý; không coi số/đơn vị OCR là dữ liệu đã xác nhận.
                            import subprocess,tempfile
                            with tempfile.TemporaryDirectory() as tmp:
                                image=Path(tmp)/'page.png';page.get_pixmap(matrix=fitz.Matrix(2,2)).save(str(image))
                                r=subprocess.run(['tesseract',str(image),'stdout','-l','eng'],capture_output=True,text=True,timeout=60)
                                if r.returncode:raise ValueError('Không OCR được trang '+str(i+1))
                                text=r.stdout;method='OCR eng, cần xác minh'
                        for line_no,line in enumerate(text.splitlines(),1):
                            if HEADER.search(normalized(line)):
                                labels.append({'label':safe_text(line),'line':line_no,'method':method})
                        add_reference(record,'trang '+str(i+1),labels,'Tài liệu thí nghiệm PDF')
                        result['sections'].append({'page':i+1,'method':method,'characters':len(text),'source_labels':len(labels)})
                    result['status']='read_pdf'
                finally:doc.close()
            else:
                result['status']='unsupported';result['warnings'].append('DWG chưa có bộ giải mã; không suy nội dung từ file DXF cạnh nó.' if ext=='.dwg' else 'Chưa có bộ đọc định dạng này.')
        except Exception as exc:
            result['status']='read_error';result['warnings'].append(str(exc))
        report.append(result);catalog['files'].append({k:record[k] for k in ('title','relative','id','sha256')})
        print(result['status'],record['relative'],len(result['sections']),flush=True)
    dictionary={'title':'parameter_registry.json','relative':'SoilFirm_Pro_2026.11/parameter_registry.json',
                'id':'software-registry','sha256':hashlib.sha256((ROOT/'parameter_registry.json').read_bytes()).hexdigest()}
    add_reference(dictionary,'từ điển thông số phần mềm',
        [{'id':key,'label':spec['label'],'unit':spec['unit'],'aliases':spec.get('alias',[]),'groups':spec.get('groups',[])}
         for key,spec in registry.items()], 'Ký hiệu và đơn vị của phần mềm')
    for name,value in [('geotech_template_catalog.json',catalog),('shared_memory_seed.json',{'version':1,'origin':'My Drive/Dev/Dao tao AI','rule':CAUTION,'items':knowledge}),('source_read_report.json',report)]:
        (output/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    summary={'files':len(report),'status':dict(Counter(r['status'] for r in report)),'projects':sorted({r['file'].split('/')[1] for r in report if r['file'].count('/')>1}),'templates':len(catalog['templates']),'knowledge_chunks':len(knowledge)}
    (output/'corpus_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(summary)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--inventory',type=Path,default=ROOT/'source_inventory_portable.json' if (ROOT/'source_inventory_portable.json').is_file() else ROOT/'source_inventory.json');parser.add_argument('--output',type=Path,default=ROOT)
    parser.add_argument('--source-dir',type=Path,help='Thư mục Dao tao AI trên máy; kiểm tra tồn tại từng file trước khi đọc.')
    args=parser.parse_args()
    if args.source_dir:
        import tempfile
        rows=json.loads(args.inventory.read_text(encoding='utf-8'))
        for row in rows:
            local=args.source_dir/row['relative'].split('/',1)[1];row['path']=str(local)
            row['status']='downloaded' if local.is_file() else 'unavailable'
            if not local.is_file():row['error']='File chưa có tại thư mục nguồn trên máy.'
        with tempfile.TemporaryDirectory() as tmp:
            inventory=Path(tmp)/'inventory.json';inventory.write_text(json.dumps(rows,ensure_ascii=False),encoding='utf-8');main(inventory,args.output)
    else:main(args.inventory,args.output)
