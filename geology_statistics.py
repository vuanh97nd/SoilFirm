"""Geology sample statistics and reviewed scalar input values. No AI-generated formula execution."""
from __future__ import annotations
import csv
import json
import math
import statistics
import unicodedata
import uuid
from copy import deepcopy
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter import font as tkfont
from ui_theme import soil_parameter_keys, COLORS, UI_FONT, configure_treeview_style, apply_row_stripes

PARAMS={'gamma':('γ','T/m³'),'e0':('e₀','—'),'cc':('Cc','—'),'cs':('Cs','—'),'pc':('Pc','T/m²'),
        'cv':('Cv','cm²/s'),'co':('Co','T/m²'),'cohesion_c':('c','T/m²'),
        'phi_cu_effective':('φ′ CU','độ'),'friction_phi':('φ cắt','độ'),'spt_n':('SPT N','búa')}
MODES=('Nhập trực tiếp','Trung bình','Hàm tại độ sâu đại diện')
REFS=('Từ mặt đất','Từ đỉnh lớp','Cao độ mẫu')
KINDS=('Tuyến tính','Bậc hai','Nội suy từng đoạn')

def numeric(value):
    if value is None or str(value).strip() in ('','-','—'):return None
    if isinstance(value,bool):raise ValueError('Giá trị số không được là đúng/sai.')
    result=float(str(value).strip().replace(',','.'))
    if not math.isfinite(result):raise ValueError('Giá trị phải là số hữu hạn.')
    return result

def parameter_text(value,key):
    if value is None:return ""
    if not isinstance(value,(int,float)):return str(value)
    if key=="cv":return f"{value:.3e}"
    return f"{value:.{0 if key=="spt_n" else 3 if key in ("e0","cc","cs") else 2}f}"

def layer_order(code):
    import re
    return tuple((0,int(part)) if part.isdigit() else (1,part.casefold()) for part in re.split(r"(\d+)",str(code)) if part)

def fmt(value):
    if value is None:return '—'
    return f'{value:.5g}' if isinstance(value,(int,float)) else str(value)

def chart_number(value,key=None):
    if value is None:return '—'
    scale=1000 if key=='cv' else 1
    decimals=3 if key in ('e0','cc','cs','cv') else 2
    return f'{value*scale:.{decimals}f}'

def smooth_pressure_pairs(pairs,log_y=False,steps=12):
    """Shape-preserving cubic display interpolation; measured values stay intact."""
    if len(pairs)<3:return pairs
    xs=[math.log10(p) for p,v in pairs]
    ys=[math.log10(v) if log_y else v for p,v in pairs]
    widths=[b-a for a,b in zip(xs,xs[1:])]
    if any(h<=0 for h in widths):return pairs
    secants=[(ys[i+1]-ys[i])/widths[i] for i in range(len(widths))]
    slopes=[0.]*len(xs)
    for i in range(1,len(xs)-1):
        before,after=secants[i-1],secants[i]
        if before*after>0:
            w1=2*widths[i]+widths[i-1];w2=widths[i]+2*widths[i-1]
            slopes[i]=(w1+w2)/(w1/before+w2/after)
    def endpoint(h0,h1,d0,d1):
        slope=((2*h0+h1)*d0-h0*d1)/(h0+h1)
        if slope*d0<=0:return 0.
        if d0*d1<=0 and abs(slope)>3*abs(d0):return 3*d0
        return slope
    slopes[0]=endpoint(widths[0],widths[1],secants[0],secants[1])
    slopes[-1]=endpoint(widths[-1],widths[-2],secants[-1],secants[-2])
    result=[pairs[0]]
    for i,h in enumerate(widths):
        for j in range(1,steps+1):
            t=j/steps;t2=t*t;t3=t2*t
            y=(2*t3-3*t2+1)*ys[i]+(t3-2*t2+t)*h*slopes[i]+(-2*t3+3*t2)*ys[i+1]+(t3-t2)*h*slopes[i+1]
            y=min(max(y,min(ys[i],ys[i+1])),max(ys[i],ys[i+1]))
            result.append((10.**(xs[i]+h*t),10.**y if log_y else y))
    return result


def merge_sample_observations(existing,incoming):
    """One observation per physical sample/indicator; never average duplicates."""
    import re
    from ai_analysis_data import borehole_key
    def key(row):
        hole=re.sub(r'[\s_-]+','',str(row.get('hole') or '')).casefold()
        start,end=row.get('from'),row.get('to')
        sample=re.sub(r'\s+','',str(row.get('sample') or '')).casefold()
        sample=re.sub(r'\d+',lambda m:str(int(m.group())),sample)
        point_kind=row.get('depth_point_field')
        if point_kind:sample='point:'+point_kind
        if not hole or start is None or end is None or not sample:return None
        return (row['layer'],hole,round(float(start),8),round(float(end),8),sample)
    report={'added_rows':0,'added_values':0,'duplicates':0,'conflicts':[]}
    lookup={key(row):row for row in existing if key(row) is not None}
    for raw in incoming:
        row=deepcopy(raw)
        present=[k for k in PARAMS if row.get(k) is not None]
        if not present and not (row.get('hole') and row.get('sample')) and not any(row.get('original',{}).get('values',{}).get(k) for k in ('e','cv')):continue
        ident=key(row);old=lookup.get(ident) if ident is not None else None
        # Missing identifying fields are not proof of duplication between files.
        if old is None:
            same=next((item for item in existing if row.get('ai_identity') and item.get('ai_identity')==row['ai_identity']),None)
            if same is not None:old=same
        if old is None and row.get('hole') and row.get('sample') and not row.get('depth_point_field'):
            # A known sample can have curves on another sheet without depth.
            # Attach only when identity is unique within this borehole and layer.
            candidates=[]
            for item in existing:
                if item.get('layer')!=row.get('layer') or borehole_key(item.get('hole'))!=borehole_key(row.get('hole')):continue
                old_sample=re.sub(r'\d+',lambda match:str(int(match[0])),re.sub(r'\s+','',str(item.get('sample') or '')).casefold())
                new_sample=re.sub(r'\d+',lambda match:str(int(match[0])),re.sub(r'\s+','',str(row.get('sample') or '')).casefold())
                if not old_sample or old_sample!=new_sample:continue
                if any(row.get(field) is not None and item.get(field) is not None and abs(float(row[field])-float(item[field]))>1e-6 for field in ('from','to')):continue
                candidates.append(item)
            if len(candidates)==1:old=candidates[0]
        if old is None:
            existing.append(row);report['added_rows']+=1;report['added_values']+=len(present)
            if ident is not None:lookup[ident]=row
            continue
        incoming_original=row.get('original',{});original=old.setdefault('original',{})
        original_values=original.setdefault('values',{})
        category=row.get('category') or incoming_original.get('values',{}).get('category')
        if category in ('Đất dính','Đất rời') and not old.get('category'):
            old['category']=category;original_values['category']=category;report['added_values']+=1
            if incoming_original.get('category_evidence'):original['category_evidence']=deepcopy(incoming_original['category_evidence'])
            if not original.get('description'):original['description']=incoming_original.get('description','')
        for field in ('from','to'):
            if old.get(field) is None and row.get(field) is not None:old[field]=row[field]
        for grid,parameter in (('ep','e'),('cvp','cv'),('mvp','mv')):
            new_values=incoming_original.get('values',{})
            observations=new_values.get(parameter,[]);pressures=new_values.get(grid,[])
            if not observations or len(pressures)!=len(observations):continue
            existing_pairs=dict(zip(original_values.get(grid,[]),original_values.get(parameter,[])))
            for pressure,value in zip(pressures,observations):
                if pressure not in existing_pairs:
                    existing_pairs[pressure]=value;report['added_values']+=1
                elif math.isclose(existing_pairs[pressure],value,rel_tol=1e-9,abs_tol=1e-10):report['duplicates']+=1
                else:
                    conflict=f"{row.get('hole','')} / {row.get('sample','')} / {parameter} tại P={pressure}: đã có {existing_pairs[pressure]}, nguồn mới {value}; giữ số cũ để đối chiếu."
                    report['conflicts'].append(conflict)
                    if conflict not in old.setdefault('import_conflicts',[]):old['import_conflicts'].append(conflict)
            ordered=sorted(existing_pairs)
            original_values[grid]=ordered;original_values[parameter]=[existing_pairs[p] for p in ordered]
        if row.get('source') and row['source'] not in str(old.get('source','')):
            old['source']=str(old.get('source',''))+'; '+row['source']
        for parameter in present:
            value=row[parameter];previous=old.get(parameter)
            if previous is None:
                old[parameter]=value;report['added_values']+=1
                old.setdefault('parameter_sources',{})[parameter]=row.get('source','')
                if row.get('source') and row['source'] not in str(old.get('source','')):old['source']=str(old.get('source',''))+'; '+row['source']
            elif math.isclose(float(previous),float(value),rel_tol=1e-9,abs_tol=1e-10):report['duplicates']+=1
            else:
                conflict=f"{row.get('hole','')} / {row.get('sample','')} / {parameter}: đã có {previous}, nguồn mới {value}; giữ số cũ, cần đối chiếu {row.get('source','')}"
                report['conflicts'].append(conflict)
                if conflict not in old.setdefault('import_conflicts',[]):old['import_conflicts'].append(conflict)
    return report


def import_report_text(report):
    return (f"Thêm {report['added_rows']} mẫu, {report['added_values']} giá trị mới; "
            f"bỏ qua {report['duplicates']} giá trị trùng; {len(report['conflicts'])} giá trị khác nhau cần đối chiếu.")


def validate_statistical_value(key,value):
    value=numeric(value)
    if value is None:raise ValueError('Chưa có giá trị hợp lệ.')
    if key=='spt_n' and (value<0 or not float(value).is_integer()):raise ValueError('SPT N phải là số nguyên không âm; dùng Nhập trực tiếp để chọn giá trị.')
    if key not in ('phi_cu_effective',) and value<0:raise ValueError('Giá trị chọn phải không âm.')
    if key=='phi_cu_effective' and not 0<=value<90:raise ValueError('φ′ CU phải từ 0 đến dưới 90 độ.')
    return value


def mean_value_choices(samples,allowed,fingerprint,kind=KINDS[0],reference=REFS[0]):
    """Create separate reviewed means; missing/invalid fields never become zero."""
    choices={};skipped={}
    for key in PARAMS:
        if key not in allowed:continue
        used=[sample for sample in samples if sample.get('use',True) and sample.get(key) is not None]
        try:
            if not used:raise ValueError('Không có mẫu hợp lệ.')
            value=validate_statistical_value(key,statistics.mean(numeric(sample[key]) for sample in used))
            choices[key]={'mode':MODES[1],'kind':kind,'reference':reference,'at':None,
                          'value':value,'curve':None,'fingerprint':fingerprint,
                          'sample_ids':[sample['id'] for sample in used]}
        except (ValueError,TypeError) as exc:skipped[key]=str(exc)
    return choices,skipped


def apply_scalar(soil,key,value):
    """Convert displayed Cv (cm²/s) to the engine's 10^-3 cm²/s."""
    if key=='cv':
        soil.cv_constant=value*1000
    else:setattr(soil,key,value)
    if key=='phi_cu_effective':soil.strength_m=math.tan(math.radians(value))

def depth(sample,reference):
    if sample.get('from') is None or sample.get('to') is None:return None
    z=(sample['from']+sample['to'])/2
    if reference==REFS[1]:return z-sample['top'] if sample.get('top') is not None else None
    if reference==REFS[2]:return sample['ground']-z if sample.get('ground') is not None else None
    return z

def fit_curve(points,kind):
    if kind==KINDS[2]:
        grouped={}
        for x,y in points:grouped.setdefault(x,[]).append(y)
        xy=sorted((x,statistics.mean(values)) for x,values in grouped.items())
        if len(xy)<2:raise ValueError('Nội suy cần ít nhất 2 độ sâu khác nhau.')
        return {'kind':kind,'points':xy,'min':xy[0][0],'max':xy[-1][0],'formula':'Nội suy tuyến tính giữa các độ sâu; giá trị trùng độ sâu lấy trung bình.'}
    degree=2 if kind==KINDS[1] else 1
    if len({x for x,y in points})<degree+1:raise ValueError('Chưa đủ độ sâu khác nhau để khớp hàm.')
    center=statistics.mean(x for x,y in points);scale=max(abs(x-center) for x,y in points)
    if scale==0:raise ValueError('Các mẫu cùng độ sâu.')
    xy=[((x-center)/scale,y) for x,y in points];size=degree+1
    matrix=[[sum(x**(i+j) for x,y in xy) for j in range(size)]+[sum(y*x**i for x,y in xy)] for i in range(size)]
    for i in range(size):
        pivot=max(range(i,size),key=lambda r:abs(matrix[r][i]));matrix[i],matrix[pivot]=matrix[pivot],matrix[i]
        if abs(matrix[i][i])<1e-12:raise ValueError('Dữ liệu chưa đủ để xác định hàm ổn định.')
        divisor=matrix[i][i];matrix[i]=[v/divisor for v in matrix[i]]
        for r in range(size):
            if r!=i:
                factor=matrix[r][i];matrix[r]=[a-factor*b for a,b in zip(matrix[r],matrix[i])]
    coefficients=[row[-1] for row in matrix]
    curve={'kind':kind,'coefficients':coefficients,'center':center,'scale':scale,'min':min(x for x,y in points),'max':max(x for x,y in points)}
    average=statistics.mean(y for x,y in points);ss=sum((y-average)**2 for x,y in points);err=sum((y-evaluate(curve,x))**2 for x,y in points)
    curve['r2']=1-err/ss if ss>0 else None
    a=coefficients[2]/scale**2 if degree==2 else 0
    b=coefficients[1]/scale-2*a*center
    c=coefficients[0]-coefficients[1]*center/scale+a*center**2
    curve['formula']=(f'y = {a:.6g} x² + {b:.6g} x + {c:.6g}' if degree==2 else f'y = {b:.6g} x + {c:.6g}')
    return curve

def evaluate(curve,x):
    if curve['kind']==KINDS[2]:
        xy=curve['points']
        for (a,b),(c,d) in zip(xy,xy[1:]):
            if a<=x<=c:return b+(d-b)*(x-a)/(c-a)
        raise ValueError('Độ sâu ngoài phạm vi nội suy.')
    t=(x-curve['center'])/curve['scale']
    return sum(value*t**i for i,value in enumerate(curve['coefficients']))

def normal(text):
    text=unicodedata.normalize('NFD',str(text or '').lower()).replace('đ','d').replace('γ','gamma').replace('φ','phi')
    return ''.join(c for c in text if c.isascii() and c.isalnum())

class GeologyStatistics(ttk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent);self.app=app;self.layer_map={};self.computed={};self.visible=[];self.dirty=False
        self._chart_axes={}
        self.layer=tk.StringVar();self.hole=tk.StringVar(value='Tất cả');self.indicator=tk.StringVar(value='γ')
        self.reference=tk.StringVar(value=REFS[0]);self.summary=tk.StringVar();self.detail=tk.StringVar();self.status=tk.StringVar()
        self.borehole_import_notice=tk.StringVar()
        self.soil_category=tk.StringVar()
        self.columnconfigure(0,weight=1);self.rowconfigure(0,weight=1)
        self.mode=tk.StringVar(value=MODES[1]);self.kind=tk.StringVar(value=KINDS[0]);self.manual=tk.StringVar();self.at=tk.StringVar()
        # Hai cột độc lập: biểu đồ bắt đầu từ đỉnh, không bị đẩy xuống
        # bởi số hàng của thanh công cụ/bảng số liệu bên trái.
        panes=ttk.Panedwindow(self,orient='horizontal');panes.grid(row=0,column=0,sticky='nsew')
        left_area=ttk.Frame(panes);right=ttk.Frame(panes)
        panes.add(left_area,weight=4);panes.add(right,weight=1)
        left_area.columnconfigure(0,weight=1);left_area.rowconfigure(4,weight=1)
        top=ttk.Frame(left_area,padding=(6,6));top.grid(row=0,column=0,sticky='ew')
        reads=ttk.Frame(top);reads.pack(side='left',fill='x',expand=True)
        read_buttons=[]
        for column,(label,kind) in enumerate((('📋 Chỉ tiêu đất','geology'),('🔬 Sức kháng','strength'),('📊 Xuyên tiêu chuẩn','spt'),('🔩 Cố kết','curves'))):
            button=ttk.Button(reads,text=label,command=lambda selected=kind:self.read_ai_table(selected))
            read_buttons.append(button)
            workspace=getattr(self.app,'_ai_analysis_workspace',None)
            if workspace is not None:workspace.buttons.append(button)
        approve_button=ttk.Button(reads,text='Xác nhận chỉ tiêu',style='Accent.TButton',command=lambda:self.layer_action('approve'))
        read_buttons.append(approve_button)
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        if workspace is not None:workspace.buttons.append(approve_button)
        self.wrap_controls(reads,read_buttons)
        toolbar=ttk.Frame(left_area,padding=(6,3));toolbar.grid(row=1,column=0,sticky='ew')
        for label,var,values,width in [('Lớp đất',self.layer,(),14),('Hố khoan',self.hole,('Tất cả',),12),('Chỉ tiêu',self.indicator,(*tuple(v[0] for v in PARAMS.values()),'e-logP','Cv-logP'),10)]:
            ttk.Label(toolbar,text=label).pack(side='left',padx=(5,3))
            widget=ttk.Combobox(toolbar,textvariable=var,values=values,state='readonly',width=width);widget.pack(side='left')
            widget.bind('<<ComboboxSelected>>',lambda e,selected=var:self.change_table_selection(selected))
            if var is self.layer:self.layer_combo=widget
            if var is self.hole:self.hole_combo=widget
        ttk.Button(toolbar,text='Lọc',command=self.refresh,width=5).pack(side='left',padx=4)
        for label,action in [('Thêm lớp','add'),('Sửa lớp','edit'),('↑','up'),('↓','down')]:
            button=ttk.Button(toolbar,text=label,width=3 if action in ('up','down') else 8,command=lambda selected=action:self.layer_action(selected))
            button.pack(side='left',padx=2)
            workspace=getattr(self.app,'_ai_analysis_workspace',None)
            if workspace is not None:workspace.buttons.append(button)
        sep=ttk.Separator(toolbar,orient='vertical')
        sep.pack(side='left',fill='y',padx=6,pady=2)
        for label,action in [('Xóa lớp','delete'),('Xóa dữ liệu bảng','reset'),('Xóa dữ liệu tuyến','reset_all')]:
            button=ttk.Button(toolbar,text=label,style='Danger.TButton',command=lambda selected=action:self.layer_action(selected))
            button.pack(side='left',padx=2)
            workspace=getattr(self.app,'_ai_analysis_workspace',None)
            if workspace is not None:workspace.buttons.append(button)
            if action == 'reset_all': self.reset_input_button = button
        toolbar_widgets=list(toolbar.winfo_children())
        for widget in toolbar_widgets:widget.pack_forget()
        self.wrap_controls(toolbar,toolbar_widgets)
        import_label=ttk.Label(left_area,textvariable=self.borehole_import_notice,padding=(8,3),anchor='w',wraplength=1000)
        import_label.grid(row=2,column=0,sticky='ew')
        import_label.bind('<Configure>',lambda event:import_label.configure(wraplength=max(100,event.width-16)))
        ttk.Label(left_area,textvariable=self.summary,padding=(8,3)).grid(row=3,column=0,sticky='ew')
        self.table_split=ttk.Panedwindow(left_area,orient='vertical')
        self.table_split.grid(row=4,column=0,sticky='nsew')
        upper=ttk.Frame(self.table_split);lower=ttk.Frame(self.table_split)
        self.table_split.add(upper,weight=3);self.table_split.add(lower,weight=2)
        left=upper
        self.chart_split=panes;self.chart_panel=right;self._chart_visible=True
        self._layout_initialized=False
        self.table_split.bind('<Configure>',self.initialize_splits,add='+')
        self.chart_split.bind('<Map>',lambda e:self.after_idle(self.position_chart) if e.widget is self.chart_split else None,add='+')
        # Căn lại khi vùng chứa đổi chiều rộng; kéo vạch chia vẫn được giữ
        # khi kích thước vùng chứa không đổi.
        self._chart_container_width = None
        def resize_chart_panel(event):
            if event.widget is self.chart_split and event.width > 200:
                if self._chart_container_width != event.width:
                    self._chart_container_width = event.width
                    self.after_idle(self.position_chart)
        self.chart_split.bind('<Configure>',resize_chart_panel,add='+')
        self.sample_tree=self.make_sample_table(left)
        self.sample_tree.bind('<Double-1>',self.edit_selected);self.sample_tree.bind('<<TreeviewSelect>>',lambda e:self.draw_chart(),add='+')
        chart_stack=ttk.Panedwindow(right,orient='vertical');chart_stack.pack(fill='both',expand=True)
        depth_panel=ttk.Frame(chart_stack);pressure_panel=ttk.Frame(chart_stack)
        chart_stack.add(depth_panel,weight=1);chart_stack.add(pressure_panel,weight=1)
        self._chart_pair_initialized=False
        def balance_chart_panels(event):
            if event.widget is chart_stack and event.height>200 and not self._chart_pair_initialized:
                self._chart_pair_initialized=True
                chart_stack.after_idle(lambda:chart_stack.sashpos(0,chart_stack.winfo_height()//2))
        chart_stack.bind('<Configure>',balance_chart_panels,add='+')
        self.pressure_kind=tk.StringVar(value='e-logP');self.pressure_detail=tk.StringVar()
        self._last_depth_param='gamma'
        self.canvas=tk.Canvas(depth_panel,bg='white',highlightthickness=0,width=280,height=260);self.canvas.pack(fill='both',expand=True,padx=8,pady=5);self.canvas.bind('<Configure>',lambda e:self.draw_chart());self.canvas.bind('<Motion>',self.hover_point)
        chart_controls=ttk.Frame(depth_panel,padding=(8,3));chart_controls.pack(fill='x',before=self.canvas)
        chart_controls.columnconfigure(0,weight=1)
        self.chart_heading_text=tk.StringVar(value='Giá trị theo độ sâu')
        ttk.Label(chart_controls,textvariable=self.chart_heading_text,font=(UI_FONT,10,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,4))
        self.reference.set(REFS[0])
        self.plot_kind=tk.StringVar(value='Hằng số')
        self.plot_kind_selector=ttk.Combobox(chart_controls,textvariable=self.plot_kind,values=('Hằng số','Hàm số'),state='readonly',width=10)
        self.plot_kind_selector.grid(row=1,column=0,sticky='ew',padx=(0,4))
        ttk.Button(chart_controls,text='Chỉnh sửa',width=9,command=self.edit_chart_axes).grid(row=1,column=1,padx=(0,4))
        ttk.Button(chart_controls,text='Copy',width=7,command=self.copy_chart_image).grid(row=1,column=2)
        self.plot_kind.trace_add('write',lambda *_:self.draw_chart())
        pressure_bar=ttk.Frame(pressure_panel,padding=(8,3));pressure_bar.pack(fill='x')
        pressure_bar.columnconfigure(0,weight=1)
        ttk.Label(pressure_bar,text='Đường cong cố kết',font=(UI_FONT,10,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,4))
        pressure_selector=ttk.Combobox(pressure_bar,textvariable=self.pressure_kind,values=('e-logP','Cv-logP'),state='readonly',width=10)
        pressure_selector.grid(row=1,column=0,sticky='ew',padx=(0,4))
        pressure_selector.bind('<<ComboboxSelected>>',lambda event:self.draw_pressure_chart())
        ttk.Button(pressure_bar,text='Chỉnh sửa',width=9,command=lambda:self.edit_chart_axes(pressure=True)).grid(row=1,column=1,padx=(0,4))
        ttk.Button(pressure_bar,text='Copy',width=7,command=self.copy_pressure_chart).grid(row=1,column=2)
        pressure_plot=ttk.Frame(pressure_panel);pressure_plot.pack(fill='both',expand=True,padx=8,pady=5)
        self.pressure_canvas=tk.Canvas(pressure_plot,bg='white',highlightthickness=0,width=280,height=260)
        self.pressure_canvas.pack(side='left',fill='both',expand=True)
        pressure_scroll=ttk.Scrollbar(pressure_plot,orient='vertical',command=self.pressure_canvas.yview)
        pressure_scroll.pack(side='right',fill='y')
        self.pressure_canvas.configure(yscrollcommand=pressure_scroll.set)
        self.pressure_canvas.bind('<Configure>',lambda event:self.draw_pressure_chart())
        self.pressure_canvas.bind('<Motion>',self.hover_pressure_point)
        ttk.Label(lower,text='Thống kê từng chỉ tiêu của lớp đang chọn',padding=(8,4),font=(UI_FONT,10,'bold')).pack(fill='x')
        choice_bar=ttk.Frame(lower,padding=(8,4));choice_bar.pack(fill='x')
        ttk.Button(choice_bar,text='Chọn giá trị',command=self.open_value_choice).pack(side='left')
        ttk.Button(choice_bar,text='TB tất cả',command=self.choose_all_means).pack(side='left',padx=6)
        ttk.Button(choice_bar,text='Áp dụng',command=self.apply_values).pack(side='left',padx=6)
        ttk.Button(choice_bar,text='Xuất thống kê',command=self.export_calculation_values).pack(side='left',padx=6)
        ttk.Button(choice_bar,text='Thống kê e–P',command=self.open_lab_statistics).pack(side='left',padx=6)
        choice_widgets=list(choice_bar.winfo_children())
        for widget in choice_widgets:widget.pack_forget()
        self.wrap_controls(choice_bar,choice_widgets)
        self.stats_tree=self.make_tree(lower,('metric',*PARAMS),('Đại lượng',*(v[0] for v in PARAMS.values())),height=7)
        self.stats_tree.bind('<ButtonRelease-1>',self.select_indicator)
        self.stats_tree.bind('<Double-1>',self.open_value_choice)
        status_label=ttk.Label(left_area,textvariable=self.status,padding=(6,2),anchor='w');status_label.grid(row=8,column=0,sticky='ew')
        left_area.bind('<Configure>',lambda e:status_label.configure(wraplength=max(200,e.width-16)) if e.widget is left_area else None,add='+')

    @staticmethod
    def wrap_controls(parent,widgets):
        def arrange(event=None):
            width=parent.winfo_width()
            if width<20:return
            x=y=row_height=0
            for widget in widgets:
                size=widget.winfo_reqwidth()+6;height=widget.winfo_reqheight()+4
                if x and x+size>width-12:y+=row_height;x=row_height=0
                widget.place(x=x+3,y=y+2)
                x+=size;row_height=max(row_height,height)
            wanted=y+row_height+4
            if int(parent.cget('height'))!=wanted:parent.configure(height=wanted)
        parent.bind('<Configure>',arrange,add='+');parent.after_idle(arrange)

    def initialize_splits(self,event=None):
        if self._layout_initialized or self.table_split.winfo_height()<180:return
        self._layout_initialized=True
        def position():
            if not self.winfo_exists():return
            self.table_split.sashpos(0,int(self.table_split.winfo_height()*.62))
            self.position_chart()
        self.after_idle(position)

    def position_chart(self):
        if not self.winfo_exists() or not self._chart_visible:return
        width=self.chart_split.winfo_width()
        if width<200:return  # Tab is not yet mapped; don't hide its chart.
        chart_width=min(340,max(280,int(width*.24)))
        self.chart_split.sashpos(0,max(100,width-chart_width))
        self.after_idle(self.draw_chart)

    def toggle_chart(self):
        if self._chart_visible:
            self.chart_split.forget(self.chart_panel)
            self._chart_visible=False
        else:
            self.chart_split.add(self.chart_panel,weight=1)
            self._chart_visible=True
            self.after_idle(self.position_chart)
        self.chart_toggle_label.set('Ẩn biểu đồ' if self._chart_visible else 'Hiện biểu đồ')

    def open_lab_statistics(self):
        from lab_statistics import show_lab_statistics
        show_lab_statistics(self)
    def copy_chart_image(self):
        from lab_statistics import copy_canvas_image
        try:
            copy_canvas_image(self.canvas)
            self.status.set('Đã copy ảnh biểu đồ. Dán bằng Ctrl+V.')
        except Exception as exc:messagebox.showerror('Copy ảnh biểu đồ',str(exc),parent=self)
    def make_sample_table(self,parent):
        box=ttk.Frame(parent);box.pack(fill='both',expand=True);box.rowconfigure(0,weight=1);box.columnconfigure(1,weight=1)
        fixed=ttk.Treeview(box,columns=('use','hole','sample','from','to'),show='headings',height=10)
        for key,label,width in [('use','Dùng',42),('hole','Hố khoan',135),('sample','Mã mẫu',90),('from','Từ (m)',65),('to','Đến (m)',65)]:fixed.heading(key,text=label);fixed.column(key,width=width,minwidth=width,stretch=False)
        moving=ttk.Treeview(box,columns=('category',*PARAMS,'e_logp','cv_logp','source'),show='headings',height=10)
        for key in ('category',*PARAMS,'e_logp','cv_logp','source'):
            label=PARAMS[key][0]+' ('+PARAMS[key][1]+')' if key in PARAMS else {'category':'Loại đất (sét/cát)','e_logp':'e-logP (P:e)','cv_logp':'Cv-logP (P:Cv)','source':'Nguồn'}[key]
            moving.heading(key,text=label);moving.column(key,width=95 if key in PARAMS else 145 if key=='category' else 260,minwidth=65,stretch=False)
        self.value_tree=moving;self._sample_scroll_sync=False;self._sample_selection_sync=False
        def scroll(*args):fixed.yview(*args);moving.yview(*args)
        y=ttk.Scrollbar(box,command=scroll);x=ttk.Scrollbar(box,orient='horizontal',command=moving.xview);moving.configure(xscrollcommand=x.set)
        def moved(origin,other,first,last):
            y.set(first,last)
            if not self._sample_scroll_sync:
                self._sample_scroll_sync=True
                try:
                    if abs(float(other.yview()[0])-float(first))>1e-6:other.yview_moveto(first)
                finally:self._sample_scroll_sync=False
        fixed.configure(yscrollcommand=lambda a,b:moved(fixed,moving,a,b));moving.configure(yscrollcommand=lambda a,b:moved(moving,fixed,a,b))
        def selected(origin,other):
            if self._sample_selection_sync:return
            ids=origin.selection()
            if other.selection()==ids:return
            self._sample_selection_sync=True
            try:other.selection_set(ids)
            finally:self._sample_selection_sync=False
            self.draw_chart()
        fixed.bind('<<TreeviewSelect>>',lambda e:selected(fixed,moving),add='+');moving.bind('<<TreeviewSelect>>',lambda e:selected(moving,fixed));moving.bind('<Double-1>',self.edit_selected)
        fixed.grid(row=0,column=0,sticky='ns');moving.grid(row=0,column=1,sticky='nsew');y.grid(row=0,column=2,sticky='ns');x.grid(row=1,column=1,sticky='ew')
        for tree in (fixed,moving):tree.tag_configure('excluded',foreground='#94A3B8')
        return fixed
    @staticmethod
    def make_tree(parent,columns,headers,height=10):
        box=ttk.Frame(parent);box.pack(fill='both',expand=True);box.rowconfigure(0,weight=1);box.columnconfigure(0,weight=1)
        tree=ttk.Treeview(box,columns=columns,show='headings',height=height)
        for i,(key,label) in enumerate(zip(columns,headers)):tree.heading(key,text=label);tree.column(key,width=125 if i<3 else 90,minwidth=65,anchor='w' if i<3 else 'center',stretch=False)
        x=ttk.Scrollbar(box,orient='horizontal',command=tree.xview);y=ttk.Scrollbar(box,command=tree.yview);tree.configure(xscrollcommand=x.set,yscrollcommand=y.set)
        tree.grid(row=0,column=0,sticky='nsew');x.grid(row=1,column=0,sticky='ew');y.grid(row=0,column=1,sticky='ns')
        tree.tag_configure('excluded',foreground='#94A3B8');tree.tag_configure('missing',foreground='#B45309');return tree
    def data_project(self):
        if self.app.design_mode.get() == 'TÍNH TOÀN TUYẾN':
            return self.app._ai_analysis_workspace.state['template']
        return self.app.project

    def store(self):
        data=self.data_project().geology_statistics
        data.setdefault('version',1);data.setdefault('samples',[]);data.setdefault('choices',{});data.setdefault('revision',0)
        return data
    def layer_action(self, action):
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        if workspace is None or workspace.busy:return
        if action=='reset':workspace.reset_section_data('statistics');return
        if action=='reset_all':workspace.reset_imported_data();return
        if action=='approve':workspace.approve_geology();return
        if action=='add':workspace.edit_material(new=True);return
        uid=self.layer_id()
        code=next((entry['code'] for entry in self.store().get('layer_catalog',{}).values() if entry['id']==uid),None)
        if code is None:
            soil=next((soil for soil in self.data_project().soils if soil.statistics_id==uid),None)
            code=soil.name if soil is not None else None
        index=next((i for i,material in enumerate(workspace.state['materials']) if code and material['code'].casefold()==code.casefold()),None)
        if index is None:
            messagebox.showinfo('Chọn lớp','Chọn lớp có trong bảng tổng hợp trước.',parent=self);return
        workspace.material_tree.selection_set(str(index))
        if action=='edit':workspace.edit_material()
        elif action in ('up','down'):workspace.move_material(-1 if action=='up' else 1)
        elif action=='delete':
            workspace.remove_row('materials',workspace.material_tree)
            if not any(material['code'].casefold()==code.casefold() for material in workspace.state['materials']):
                self.store()['samples']=[sample for sample in self.store()['samples'] if sample['layer']!=uid]
                self.store().get('layer_catalog',{}).pop(code.casefold(),None)
                self.store()['choices'].pop(uid,None)
        self.refresh()

    def read_ai_table(self, kind='geology'):
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        if workspace is None:
            messagebox.showinfo('Đọc BTH','Chưa khởi tạo bàn làm việc dữ liệu.',parent=self)
            return
        workspace.read_source(kind)

    def import_ai_samples(self, materials=None):
        materials=materials if materials is not None else getattr(self.app,'_ai_analysis_state',{}).get('materials',[])
        data=self.store();catalog=data.setdefault('layer_catalog',{})
        before_catalog=deepcopy(catalog);before_count=len(data['samples'])
        incoming=[]
        for material in materials:
            code=material['code'];normalized=code.casefold()
            soil=next((soil for soil in self.data_project().soils if soil.name.casefold()==normalized),None)
            if soil is not None:
                if not soil.statistics_id:soil.statistics_id=uuid.uuid4().hex
                uid=soil.statistics_id
            else:uid=catalog.get(normalized,{}).get('id') or uuid.uuid4().hex
            catalog[normalized]={'id':uid,'code':code}
            chosen_category=data.get('layer_categories',{}).get(uid)
            if chosen_category in ('Đất dính','Đất rời'):
                material['values']['category']=chosen_category
                material['category_source']='Người dùng chọn trong Thống kê'
            for sample in material.get('samples') or [material]:
                values=sample.get('values',sample)
                identity=json.dumps([uid,sample.get('borehole_name'),sample.get('sample_id'),
                    sample.get('depth_from'),sample.get('depth_to'),sample.get('test_depth'),sample.get('source')],ensure_ascii=False,default=str)
                row={'id':uuid.uuid4().hex,'ai_identity':identity,'layer':uid,'layer_code':code,
                    'hole':sample.get('borehole_name',''),'sample':sample.get('sample_id',''),
                    'from':sample.get('depth_from') if sample.get('depth_from') is not None else sample.get('test_depth'),
                    'to':sample.get('depth_to') if sample.get('depth_to') is not None else sample.get('test_depth'),
                    'ground':None,'top':None,'category':chosen_category or values.get('category'),'use':True,'source':sample.get('source',''),
                    'original':deepcopy(sample),'reason':'','depth_point_field':sample.get('depth_point_field')}
                for key in PARAMS:
                    row[key]=values.get(key) if key!='cv' else (
                        numeric(sample.get('cv_constant',values.get('cv_constant')))*.001
                        if sample.get('cv_constant',values.get('cv_constant')) is not None else None)
                incoming.append(row)
        report=merge_sample_observations(data['samples'],incoming)
        self._last_import_report=report
        order=[material['code'].casefold() for material in materials]
        data['layer_catalog']={key:catalog[key] for key in dict.fromkeys(order+list(catalog))}
        if report['added_values'] or len(data['samples'])!=before_count or list(data['layer_catalog'])!=list(before_catalog) or catalog!=before_catalog:
            data['revision']+=1
            self.update_ai_summary()
        # Nếu bộ lọc cũ không có số liệu, mở lớp vừa nhập có chỉ tiêu thật.
        # Không thay lớp người dùng đang xem khi lớp đó vẫn có số liệu.
        def readable(row):
            if any(row.get(key) is not None for key in PARAMS):return True
            from ai_analysis_data import has_measured_indicators
            return has_measured_indicators(row.get('original',{}))
        current=self.layer_id()
        current_rows=[row for row in data['samples'] if row['layer']==current and
                      (self.hole.get()=='Tất cả' or row.get('hole','')==self.hole.get())]
        if not any(readable(row) for row in current_rows):
            target=next((row['layer'] for row in incoming if readable(row)),None)
            if target is not None:
                self.refresh()  # Cập nhật danh sách lớp trước khi chọn nhãn.
                label=next((label for label,uid in self.layer_map.items() if uid==target),None)
                if label is not None:self.layer.set(label)
                self.hole.set('Tất cả')
        # Phiếu Su/SPT hoặc đường cong không có gamma: chọn chỉ tiêu có
        # số liệu thay vì giữ biểu đồ gamma trắng sau lần nhập mới.
        selected=[row for row in data['samples'] if row['layer']==self.layer_id() and
                  (self.hole.get()=='Tất cả' or row.get('hole','')==self.hole.get())]
        if selected and hasattr(self,'indicator'):
            key=next((k for k,v in PARAMS.items() if v[0]==self.indicator.get()),None)
            curve=self.indicator.get() in ('e-logP','Cv-logP')
            if key and not any(row.get(key) is not None for row in selected) and not curve:
                available=next((k for k in PARAMS if any(row.get(k) is not None for row in selected)),None)
                if available is not None:self.indicator.set(PARAMS[available][0])
                elif any(row.get('original',{}).get('values',{}).get('e') for row in selected):
                    self.indicator.set('e-logP')
                    if hasattr(self,'pressure_kind'):self.pressure_kind.set('e-logP')
                elif any(row.get('original',{}).get('values',{}).get('cv') for row in selected):
                    self.indicator.set('Cv-logP')
                    if hasattr(self,'pressure_kind'):self.pressure_kind.set('Cv-logP')
        self.refresh()
        if hasattr(self,'status'):self.status.set(import_report_text(report))

    def show_spt_results(self,source_holes):
        return GeologyStatistics.show_test_results(self,"spt_n",source_holes)

    def show_strength_results(self,source_holes):
        return GeologyStatistics.show_test_results(self,"co",source_holes)

    def show_test_results(self,parameter,source_holes):
        """Focus an assigned test row belonging to the imported borehole."""
        from ai_analysis_data import borehole_key,strength_borehole_key
        if parameter=='co':borehole_key=strength_borehole_key
        allowed={borehole_key(name) for name in source_holes if name}
        points=[sample for sample in self.store()['samples']
                if sample.get('depth_point_field')==parameter and sample.get(parameter) is not None
                and (not allowed or borehole_key(sample.get('hole')) in allowed)]
        if not points:return False
        target=next((sample for sample in points if sample.get('hole')==self.hole.get()),points[0])
        self.refresh()
        label=next((label for label,uid in self.layer_map.items() if uid==target['layer']),None)
        if label is None:return False
        self.layer.set(label);self.hole.set(target.get('hole') or 'Tất cả');self.indicator.set(PARAMS[parameter][0])
        self.refresh();self.sample_tree.see(target['id']);self.value_tree.see(target['id'])
        return True

    def update_ai_summary(self):
        state=getattr(self.app,'_ai_analysis_state',{})
        data=self.store();catalog=data.get('layer_catalog',{})
        for material in state.get('materials',[]):
            if material.get('average_edited'):continue
            info=catalog.get(material['code'].casefold())
            if not info:continue
            samples=[sample for sample in data['samples'] if sample['layer']==info['id'] and sample.get('use',True)]
            for key in PARAMS:
                available=[sample[key] for sample in samples if sample.get(key) is not None]
                target='cv_constant' if key=='cv' else key
                value=statistics.mean(available) if available else None
                if key=='cv':
                    material['cv_constant']=value*1000 if value is not None else None
                    if material['values'].get('cv'):material['values'].pop('cv_constant',None)
                    elif value is not None:material['values']['cv_constant']=value*1000
                elif value is not None:material['values'][target]=value
                else:material['values'].pop(target,None)
            phi=material['values'].get('phi_cu_effective')
            material['phi_cu_effective']=phi
            if phi is not None:material['values']['strength_m']=math.tan(math.radians(phi))
            else:material['values'].pop('strength_m',None)
            material['sample_count']=len(samples)
            material['statistics_source']='Giá trị trung bình các mẫu được chọn trong Thống kê'
        state['geology_approved']=False;state['pending']=None
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        if workspace is not None:workspace.clear_results();workspace.refresh()

    def category_for_layer(self,uid):
        category=self.store().get('layer_categories',{}).get(uid)
        if category:return category
        soil=next((soil for soil in self.data_project().soils if soil.statistics_id==uid),None)
        code=next((entry['code'] for entry in self.store().get('layer_catalog',{}).values() if entry['id']==uid),soil.name if soil else '')
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        materials=[m for m in workspace.state['materials'] if m['code'].casefold()==code.casefold()] if workspace else []
        if materials:
            return materials[0]['values'].get('category') if len(materials)==1 else None
        return soil.category if soil is not None and soil.category in ('Đất dính','Đất rời') else None

    def set_layer_category(self):
        uid=self.layer_id()
        category={'Đất dính (sét)':'Đất dính','Đất rời (cát)':'Đất rời'}.get(self.soil_category.get())
        if not uid or not category:return
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        if workspace is not None and workspace.busy:return
        soil=next((soil for soil in self.data_project().soils if soil.statistics_id==uid),None)
        code=next((entry['code'] for entry in self.store().get('layer_catalog',{}).values() if entry['id']==uid),soil.name if soil else '')
        materials=[m for m in workspace.state['materials'] if m['code'].casefold()==code.casefold()] if workspace else []
        if len(materials)>1:
            messagebox.showwarning('Loại đất','Mã lớp trùng nguồn; sửa nguồn trước khi chọn loại đất.',parent=self);self.refresh();return
        self.store().setdefault('layer_categories',{})[uid]=category
        for sample in self.store()['samples']:
            if sample['layer']==uid:sample['category']=category
        if soil is not None:soil.category=category
        if materials:materials[0]['values']['category']=category
        self.store()['revision']+=1;self.data_project().calculation_results={}
        for attr in ('_choice_group_results','_cdm_reports','_step_result_cache'):getattr(self.app,attr,{}).clear()
        if workspace is not None:workspace.invalidate('geology_approved')
        if hasattr(self.app,'invalidate_assessment'):self.app.invalidate_assessment()
        self.refresh();self.status.set('Đã chọn '+category+' cho lớp '+code+'. Cần xác nhận và tính lại; chỉ tiêu riêng lỗ khoan giữ nguyên.')

    def layer_id(self):return self.layer_map.get(self.layer.get())
    def change_table_selection(self,variable):
        if variable is self.indicator and self.indicator.get() in ('e-logP','Cv-logP'):
            self.pressure_kind.set(self.indicator.get())
        self.refresh()

    def selected_param(self):
        key=next((key for key,value in PARAMS.items() if value[0]==self.indicator.get()),None)
        if key is not None:self._last_depth_param=key
        return key or getattr(self,'_last_depth_param','gamma')
    def curve_cells(self,sample):
        values=sample.get('original',{}).get('values',{})
        return tuple('; '.join(f'{fmt(p)}:{fmt(v)}' for p,v in zip(values.get(grid,[]),values.get(parameter,[]))) or '—'
                     for grid,parameter in (('ep','e'),('cvp','cv')))

    def refresh(self):
        self.reset_input_button.configure(text=('Xóa dữ liệu toàn tuyến' if self.app.design_mode.get() == 'TÍNH TOÀN TUYẾN' else 'Xóa chỉ tiêu đất'))
        state=getattr(self.app,'_ai_analysis_state',{}) if self.app.design_mode.get()=='TÍNH TOÀN TUYẾN' else {}
        self.borehole_import_notice.set(state.get('last_borehole_import_summary',''))
        self.layer_map={}
        for i,soil in enumerate(self.data_project().soils):
            if not soil.statistics_id:soil.statistics_id=uuid.uuid4().hex
            self.layer_map[f'{i+1}. {soil.name or "Lớp đất"}']=soil.statistics_id
        for entry in self.store().get('layer_catalog',{}).values():
            if entry['id'] not in self.layer_map.values():
                self.layer_map[f"{len(self.layer_map)+1}. {entry['code']}"]=entry['id']
        self.layer_combo.configure(values=tuple(self.layer_map))
        if self.layer.get() not in self.layer_map:self.layer.set(next(iter(self.layer_map),''))
        data=self.store();uid=self.layer_id();samples=[s for s in data['samples'] if s['layer']==uid]
        holes=('Tất cả',*sorted({s.get('hole','') for s in samples}));self.hole_combo.configure(values=holes)
        if self.hole.get() not in holes:self.hole.set('Tất cả')
        self.visible=[s for s in samples if (self.hole.get()=='Tất cả' or s.get('hole')==self.hole.get())]
        self.sample_tree.delete(*self.sample_tree.get_children());self.value_tree.delete(*self.value_tree.get_children())
        for sample in self.visible:
            self.sample_tree.insert('','end',iid=sample['id'],values=('☑' if sample.get('use',True) else '☐',sample.get('hole'),sample.get('sample'),fmt(sample['from']),fmt(sample['to']),*(parameter_text(sample.get(key),key) or "—" for key in PARAMS),*self.curve_cells(sample),sample.get('source','')),tags=('excluded' if not sample.get('use',True) else '',))
        category=self.category_for_layer(uid)
        self.soil_category.set({'Đất dính':'Đất dính (sét)','Đất rời':'Đất rời (cát)'}.get(category,''))
        for sample in self.visible:self.value_tree.insert('','end',iid=sample['id'],values=({'Đất dính':'Sét / đất dính','Đất rời':'Cát / đất rời'}.get(sample.get('category') or category,'Chưa xác định'),*(parameter_text(sample.get(key),key) or "—" for key in PARAMS),*self.curve_cells(sample),sample.get('source','')),tags=('excluded' if not sample.get('use',True) else '',))
        used=sum(s.get('use',True) for s in self.visible);key=self.selected_param();missing=sum(s.get(key) is None for s in self.visible)
        self.summary.set(f'{len(self.visible)} mẫu trong bộ lọc / {len(samples)} mẫu của lớp · Dùng: {used} · Loại: {len(self.visible)-used} · Thiếu {PARAMS[key][0]}: {missing}')
        fallback_count=sum(bool(sample.get('original',{}).get('cs_fallback')) for sample in self.visible)
        if fallback_count:self.summary.set(self.summary.get()+f' · Cs dùng Cr: {fallback_count} mẫu')
        self.recalculate(quiet=True)
    def fingerprint(self):return json.dumps([(s['id'],s.get('use',True),*(s.get(k) for k in ('from','to','ground','top','category',*PARAMS))) for s in self.visible],ensure_ascii=False,sort_keys=True)
    def recalculate(self,quiet=False):
        self.computed={};rows=[('n','Số mẫu hợp lệ'),('min','Nhỏ nhất'),('max','Lớn nhất'),('mean','Trung bình'),('sd','Độ lệch chuẩn mẫu'),('cv','Hệ số biến thiên'),('selected','Giá trị chọn')]
        choices=self.store()['choices'].get(self.layer_id(),{})
        for key in PARAMS:
            values=[s[key] for s in self.visible if s.get('use',True) and s.get(key) is not None]
            n=len(values);mean=statistics.mean(values) if values else None;sd=statistics.stdev(values) if n>1 else None
            saved=choices.get(key,{})
            self.computed[key]={'n':n,'min':min(values) if values else None,'max':max(values) if values else None,'mean':mean,'sd':sd,'cv':sd/abs(mean) if sd is not None and mean else None,'selected':saved.get('value')}
        curve_groups={}
        for sample in self.visible:
            if not sample.get('use',True):continue
            values=sample.get('original',{}).get('values',{})
            for grid,parameter in (('ep','e'),('cvp','cv')):
                pressures,observations=values.get(grid,[]),values.get(parameter,[])
                if len(pressures)!=len(observations):continue
                own={}
                for pressure,value in zip(pressures,observations):
                    if not isinstance(pressure,(int,float)) or not isinstance(value,(int,float)):continue
                    if not math.isfinite(pressure) or not math.isfinite(value) or pressure<0:continue
                    own.setdefault(pressure,[]).append(value)
                for pressure,observations in own.items():
                    if len(set(observations))==1:curve_groups.setdefault((parameter,pressure),[]).append(observations[0])
        self._stat_curve_columns=sorted(curve_groups,key=lambda item:(0 if item[0]=='e' else 1,item[1]))
        curve_computed={}
        for curve_key in self._stat_curve_columns:
            values=curve_groups[curve_key];n=len(values);mean=statistics.mean(values);sd=statistics.stdev(values) if n>1 else None
            curve_computed[curve_key]={'n':n,'min':min(values),'max':max(values),'mean':mean,'sd':sd,'cv':sd/abs(mean) if sd is not None and mean else None,'selected':None}
        columns=('metric',*PARAMS,*(f'curve_{index}' for index in range(len(self._stat_curve_columns))))
        self.stats_tree.configure(columns=columns)
        # Reconfiguring Treeview columns resets heading text and column options.
        # Restore scalar headings as well as dynamically added curve headings.
        self.stats_tree.heading('metric',text='Đại lượng')
        for index,(parameter,(label,unit)) in enumerate(PARAMS.items()):
            self.stats_tree.heading(parameter,text=label)
            self.stats_tree.column(parameter,width=125 if index<2 else 90,
                                   minwidth=65,anchor='center',stretch=False)
        for index,(parameter,pressure) in enumerate(self._stat_curve_columns):
            column=f'curve_{index}'
            self.stats_tree.heading(column,text=('e-logP' if parameter=='e' else 'Cv-logP')+f' · P={pressure:g}')
            self.stats_tree.column(column,width=170,minwidth=120,anchor='center',stretch=False)
        self.stats_tree.column('metric',width=165,minwidth=145,stretch=False)
        self.stats_tree.delete(*self.stats_tree.get_children())
        for key,label in rows:
            scalar_values=[fmt(self.computed[p][key]) if key=='n' else parameter_text(self.computed[p][key],p) or '—' for p in PARAMS]
            curve_values=[fmt(curve_computed[p][key]) if key=='n' else (f'{curve_computed[p][key]:.3f}' if curve_computed[p][key] is not None else '—') for p in self._stat_curve_columns]
            self.stats_tree.insert('','end',values=(label,*scalar_values,*curve_values))
        key=self.selected_param();saved=choices.get(key,{})
        if saved:
            self.mode.set(saved['mode']);self.kind.set(saved.get('kind',KINDS[0]));self.manual.set(fmt(saved.get('value')) if saved['mode']==MODES[0] else '');self.at.set(str(saved.get('at','')))
        self.draw_chart()
        stale=any(c.get('mode')!=MODES[0] and c.get('fingerprint')!=self.fingerprint() for c in choices.values())
        if stale:self.status.set('Dữ liệu hoặc bộ lọc đã thay đổi. Cần lưu lại lựa chọn trước khi áp dụng; giá trị trong bộ tính chưa tự đổi.')
        elif not quiet:self.status.set('Đã thống kê. HSBĐ = ĐLC mẫu / |trung bình|; không xác định khi trung bình bằng 0. Không tự loại mẫu.')
    def select_indicator(self,event):
        col=self.stats_tree.identify_column(event.x)
        index=int(col[1:])-2 if col else -1
        row=self.stats_tree.identify_row(event.y)
        values=self.stats_tree.item(row,'values') if row else ()
        if 0<=index<len(PARAMS):
            self.indicator.set(PARAMS[list(PARAMS)[index]][0]);self.recalculate(quiet=True)
            if values and values[0]=='Giá trị chọn':self.open_value_choice()
        elif 0<=index-len(PARAMS)<len(getattr(self,'_stat_curve_columns',[])):
            parameter,pressure=self._stat_curve_columns[index-len(PARAMS)]
            self.pressure_kind.set('e-logP' if parameter=='e' else 'Cv-logP')
            self.draw_pressure_chart()
            if values and values[0]=='Giá trị chọn':self.open_lab_statistics()
    def open_value_choice(self,event=None):
        if event is None and self.indicator.get() in ('e-logP','Cv-logP'):
            self.open_lab_statistics();return 'break'
        existing=getattr(self,'_value_choice_popup',None)
        if existing is not None and existing.winfo_exists():
            existing.lift();existing.focus_set();return 'break'
        if event is not None:
            col=self.stats_tree.identify_column(event.x)
            index=int(col[1:])-2 if col else -1
            if 0<=index<len(PARAMS):self.indicator.set(PARAMS[list(PARAMS)[index]][0])
            elif 0<=index-len(PARAMS)<len(getattr(self,'_stat_curve_columns',[])):
                self.open_lab_statistics();return 'break'
        if not self.layer_id():
            messagebox.showinfo('Chọn giá trị','Chọn lớp đất trước khi chọn chỉ tiêu.',parent=self);return
        popup=tk.Toplevel(self);popup.title('Chọn giá trị dùng tính toán');popup.transient(self.winfo_toplevel())
        self._value_choice_popup=popup
        body=ttk.Frame(popup,padding=14);body.pack(fill='both',expand=True)
        allowed=tuple(v[0] for k,v in PARAMS.items() if k in soil_parameter_keys(self.data_project().method))
        selected=tk.StringVar(value=self.indicator.get() if self.indicator.get() in allowed else allowed[0]);unit=tk.StringVar()
        ttk.Label(body,text='Chỉ tiêu').grid(row=0,column=0,sticky='w',pady=4)
        selector=ttk.Combobox(body,textvariable=selected,values=allowed,state='readonly',width=24)
        selector.grid(row=0,column=1,sticky='ew',pady=4)
        ttk.Label(body,text='Cách chọn').grid(row=1,column=0,sticky='w',pady=4)
        mode=ttk.Combobox(body,textvariable=self.mode,values=MODES,state='readonly',width=28);mode.grid(row=1,column=1,sticky='ew',pady=4)
        ttk.Label(body,textvariable=unit).grid(row=2,column=0,sticky='w',pady=4)
        manual=ttk.Entry(body,textvariable=self.manual);manual.grid(row=2,column=1,sticky='ew',pady=4)
        ttk.Label(body,text='Dạng hàm').grid(row=3,column=0,sticky='w',pady=4)
        kind=ttk.Combobox(body,textvariable=self.kind,values=KINDS,state='readonly');kind.grid(row=3,column=1,sticky='ew',pady=4)
        ttk.Label(body,text='Mốc độ sâu / cao độ').grid(row=4,column=0,sticky='w',pady=4)
        reference=ttk.Combobox(body,textvariable=self.reference,values=REFS,state='readonly');reference.grid(row=4,column=1,sticky='ew',pady=4)
        ttk.Label(body,text='Độ sâu / cao độ đại diện (m)').grid(row=5,column=0,sticky='w',pady=4)
        at=ttk.Entry(body,textvariable=self.at);at.grid(row=5,column=1,sticky='ew',pady=4)
        def update(*_):
            manual.configure(state='normal' if self.mode.get()==MODES[0] else 'disabled')
            fitted=self.mode.get()==MODES[2]
            kind.configure(state='readonly' if fitted else 'disabled')
            reference.configure(state='readonly' if fitted else 'disabled')
            at.configure(state='normal' if fitted else 'disabled')
        def load(*_):
            self.indicator.set(selected.get());self.recalculate(quiet=True)
            key=self.selected_param();saved=self.store()['choices'].get(self.layer_id(),{}).get(key)
            if not saved:self.mode.set(MODES[1]);self.manual.set('');self.at.set('')
            elif saved.get('reference') in REFS:self.reference.set(saved['reference'])
            feedback.set('')
            unit.set('Nhập giá trị ('+PARAMS[key][1]+')');update()
        selector.bind('<<ComboboxSelected>>',load);mode.bind('<<ComboboxSelected>>',update)
        ttk.Label(body,text='Trung bình: dùng các mẫu được đánh dấu Dùng. Hàm: lấy giá trị tại độ sâu đại diện, không ngoại suy.',wraplength=480).grid(row=6,column=0,columnspan=2,sticky='w',pady=8)
        actions=ttk.Frame(body);actions.grid(row=7,column=0,columnspan=2,sticky='e')
        feedback=tk.StringVar()
        ttk.Label(body,textvariable=feedback,wraplength=480).grid(row=8,column=0,columnspan=2,sticky='w',pady=(8,0))
        def save():
            if not self.choose_value():return
            saved=self.store()['choices'].get(self.layer_id(),{}).get(self.selected_param(),{})
            value=saved.get('value')
            if value is not None:feedback.set('Giá trị đã lưu: '+parameter_text(value,self.selected_param())+'. Đóng cửa sổ rồi bấm Áp dụng chỉ tiêu.')
        ttk.Button(actions,text='Lưu giá trị',command=save).pack(side='left',padx=4)
        ttk.Button(actions,text='TB tất cả',command=lambda:feedback.set(self.choose_all_means())).pack(side='left',padx=4)
        ttk.Button(actions,text='Đóng',command=popup.destroy).pack(side='left',padx=4)
        load();popup.update_idletasks();popup.grab_set();popup.lift();selector.focus_set()
        return 'break'

    def choose_all_means(self):
        if not self.layer_id():
            messagebox.showinfo('Chọn giá trị','Chọn lớp đất trước.',parent=self)
            return 'Chưa chọn lớp đất.'
        choices,skipped=mean_value_choices(self.visible,soil_parameter_keys(self.data_project().method),
            self.fingerprint(),self.kind.get(),self.reference.get())
        self.store()['choices'].setdefault(self.layer_id(),{}).update(choices)
        self.recalculate(quiet=True)
        text='Đã lưu trung bình '+str(len(choices))+' chỉ tiêu của lớp đang chọn. Bấm Áp dụng chỉ tiêu.'
        if skipped:text+=' Chưa lưu: '+', '.join(PARAMS[key][0]+' ('+reason+')' for key,reason in skipped.items())
        self.status.set(text)
        return text

    def choose_value(self):
        try:
            if not self.layer_id():raise ValueError('Thêm lớp đất trong thẻ Thông số tính toán trước.')
            key=self.selected_param();curve=None;x=None
            if key not in soil_parameter_keys(self.data_project().method):
                raise ValueError('Chỉ tiêu này không thuộc phương pháp đã chọn trong thiết lập tính toán.')
            if self.mode.get()==MODES[0]:value=numeric(self.manual.get())
            elif self.mode.get()==MODES[1]:value=self.computed[key]['mean']
            else:
                points=[(depth(s,self.reference.get()),s[key]) for s in self.visible if s.get('use',True) and s.get(key) is not None and depth(s,self.reference.get()) is not None]
                curve=fit_curve(points,self.kind.get());x=numeric(self.at.get())
                if x is None:raise ValueError('Nhập x đại diện theo trục đang chọn.')
                if not curve['min']<=x<=curve['max']:raise ValueError('x ngoài phạm vi mẫu. Chưa cho phép ngoại suy; chọn x trong khoảng '+fmt(curve['min'])+' đến '+fmt(curve['max'])+'.')
                value=evaluate(curve,x)
            value=validate_statistical_value(key,value)
            config={'mode':self.mode.get(),'kind':self.kind.get(),'reference':self.reference.get(),'at':x,'value':value,'curve':curve,'fingerprint':self.fingerprint(),'sample_ids':[s['id'] for s in self.visible if s.get('use',True) and s.get(key) is not None]}
            self.store()['choices'].setdefault(self.layer_id(),{})[key]=config;self.recalculate(quiet=True)
            self.status.set(f'Đã chọn {PARAMS[key][0]} = {fmt(value)} {PARAMS[key][1]}. Hàm chỉ lấy tại x đại diện để nhập giá trị cố định; chưa áp dụng vào lớp.')
            return True
        except (ValueError,KeyError) as exc:
            popup=getattr(self,'_value_choice_popup',None)
            messagebox.showerror('Chọn giá trị',str(exc),parent=popup if popup is not None and popup.winfo_exists() else self)
            return False
    def edit_chart_axes(self,pressure=False):
        pressure_chart=pressure
        key=(self.pressure_kind.get(),'logP') if pressure_chart else (self.selected_param(),self.reference.get())
        settings=self._chart_axes.get(key,{})
        window=tk.Toplevel(self);window.title('Chỉnh sửa biểu đồ');window.transient(self.winfo_toplevel())
        body=ttk.Frame(window,padding=14);body.pack(fill='both',expand=True)
        variables={name:tk.StringVar(value='' if settings.get(name) is None else str(settings[name]))
                   for name in ('xmin','xmax','zmin','zmax','xstep','zstep')}
        vertical=tk.BooleanVar(value=settings.get('vertical',True));horizontal=tk.BooleanVar(value=settings.get('horizontal',True))
        auto_x=tk.BooleanVar(value=settings.get('xstep') is None);auto_z=tk.BooleanVar(value=settings.get('zstep') is None)
        unit='kgf/cm²' if pressure_chart else '10⁻³ cm²/s' if key[0]=='cv' else PARAMS[key[0]][1]
        yunit=('không thứ nguyên' if key[0]=='e-logP' else '10⁻³ cm²/s') if pressure_chart else 'm'
        entries={}
        labels=(('xmin',f'Trục ngang nhỏ nhất ({unit})'),('xmax',f'Trục ngang lớn nhất ({unit})'),
                ('zmin',f'Trục dọc nhỏ nhất ({yunit})'),('zmax',f'Trục dọc lớn nhất ({yunit})'),
                ('xstep','Bước log₁₀(P)' if pressure_chart else f'Bước trục ngang ({unit})'),('zstep','Bước log₁₀(Cv)' if pressure_chart and key[0]=='Cv-logP' else f'Bước trục dọc ({yunit})'))
        for row,(name,label) in enumerate(labels):
            ttk.Label(body,text=label).grid(row=row,column=0,sticky='w',pady=4)
            entry=ttk.Entry(body,textvariable=variables[name],width=16);entry.grid(row=row,column=1,padx=(12,0),pady=4);entries[name]=entry
        def update_auto():
            entries['xstep'].configure(state='disabled' if auto_x.get() else 'normal')
            entries['zstep'].configure(state='disabled' if auto_z.get() else 'normal')
        ttk.Checkbutton(body,text='Tự chia lưới trục ngang',variable=auto_x,command=update_auto).grid(row=6,column=0,columnspan=2,sticky='w')
        ttk.Checkbutton(body,text='Tự chia lưới trục dọc',variable=auto_z,command=update_auto).grid(row=7,column=0,columnspan=2,sticky='w')
        ttk.Checkbutton(body,text='Hiện đường gióng dọc',variable=vertical).grid(row=8,column=0,columnspan=2,sticky='w')
        ttk.Checkbutton(body,text='Hiện đường gióng ngang',variable=horizontal).grid(row=9,column=0,columnspan=2,sticky='w')
        ttk.Label(body,text='Giới hạn để trống: lấy theo số liệu.').grid(row=10,column=0,columnspan=2,sticky='w',pady=5)
        def apply():
            try:
                values={}
                for name,var in variables.items():
                    if (name=='xstep' and auto_x.get()) or (name=='zstep' and auto_z.get()):values[name]=None;continue
                    text=var.get().strip();value=float(text.replace(',','.')) if text else None
                    if value is not None and not math.isfinite(value):raise ValueError('Giá trị trục phải là số hữu hạn.')
                    if name.endswith('step') and (value is None or value<=0):raise ValueError('Bước chia phải là số dương.')
                    values[name]=value
                limits=getattr(self,'_pressure_data_limits' if pressure_chart else '_chart_data_limits',None)
                if limits:
                    scale=1000 if key[0]=='cv' else 1
                    fallback=(limits[0]*scale,limits[1]*scale,limits[2],limits[3])
                    for lo,hi,step,i in (('xmin','xmax','xstep',0),('zmin','zmax','zstep',2)):
                        low=values[lo] if values[lo] is not None else fallback[i]
                        high=values[hi] if values[hi] is not None else fallback[i+1]
                        if low>=high:raise ValueError('Giá trị lớn nhất phải lớn hơn giá trị nhỏ nhất.')
                        if pressure_chart and (lo=='xmin' or (lo=='zmin' and key[0]=='Cv-logP')):
                            if low<=0:raise ValueError('P và Cv phải lớn hơn 0 trên trục logarit.')
                            low,high=math.log10(low),math.log10(high)
                        if values[step] and (high-low)/values[step]>200:raise ValueError('Bước chia quá nhỏ; tối đa 200 khoảng mỗi trục.')
            except ValueError as exc:messagebox.showerror('Chỉnh sửa biểu đồ',str(exc),parent=window);return
            self._chart_axes[key]={**values,'vertical':vertical.get(),'horizontal':horizontal.get()}
            self.draw_chart();window.destroy()
        def reset():
            for var in variables.values():var.set('')
            auto_x.set(True);auto_z.set(True);vertical.set(True);horizontal.set(True);update_auto()
        buttons=ttk.Frame(body);buttons.grid(row=11,column=0,columnspan=2,sticky='e',pady=(12,0))
        for label,command in (('Đặt lại',reset),('Áp dụng',apply),('Đóng',window.destroy)):
            ttk.Button(buttons,text=label,command=command).pack(side='left',padx=4)
        update_auto();window.grab_set()

    def draw_chart(self):
        if not hasattr(self,'canvas'):return
        if hasattr(self,'pressure_canvas'):
            self.draw_pressure_chart()
        c=self.canvas;c.delete('all');key=self.selected_param();points=[]
        for sample in self.visible:
            z=depth(sample,self.reference.get())
            if sample.get(key) is not None and z is not None:points.append((sample[key],z,sample))
        self.drawn=[]
        if not points:c.create_text(12,20,anchor='nw',text='Chưa có mẫu có giá trị và độ sâu.',fill='#64748B');return
        w=c.winfo_width();h=c.winfo_height()
        if w<150 or h<180:
            c.create_text(w/2,h/2,text='Kéo rộng khung để xem biểu đồ.',width=max(100,w-16),fill='#64748B');return
        fitted_display=self.plot_kind.get()!='Hằng số'
        left,right,top,bottom=62,w-18,(78 if fitted_display else 60),h-(96 if fitted_display else 70)
        values=[x for x,z,s in points];zs=[z for x,z,s in points];xmin,xmax=min(values),max(values);zmin,zmax=min(zs),max(zs)
        dx=max((xmax-xmin)*.15,abs(xmax)*.05,.00001 if key=='cv' else .01);xmin-=dx;xmax+=dx
        dz=max((zmax-zmin)*.08,.2);zmin-=dz;zmax+=dz
        self._chart_data_limits=(xmin,xmax,zmin,zmax)
        settings=self._chart_axes.get((key,self.reference.get()),{})
        scale=1000 if key=='cv' else 1
        xmin=settings['xmin']/scale if settings.get('xmin') is not None else xmin
        xmax=settings['xmax']/scale if settings.get('xmax') is not None else xmax
        zmin=settings['zmin'] if settings.get('zmin') is not None else zmin
        zmax=settings['zmax'] if settings.get('zmax') is not None else zmax
        if xmin>=xmax or zmin>=zmax:
            c.create_text(12,20,anchor='nw',text='Giới hạn trục không phù hợp; mở Chỉnh sửa biểu đồ để đặt lại.');return
        def px(x):return left+(x-xmin)/(xmax-xmin)*(right-left)
        def py(z):
            t=(z-zmin)/(zmax-zmin)
            return bottom-t*(bottom-top) if self.reference.get()==REFS[2] else top+t*(bottom-top)
        names={'gamma':'Trọng lượng thể tích tự nhiên γ','e0':'Hệ số rỗng e₀','cc':'Chỉ số nén Cc','cs':'Chỉ số nở Cs','pc':'Áp lực tiền cố kết Pc','cv':'Hệ số cố kết Cv','co':'Sức kháng cắt không thoát nước c₀','cohesion_c':'Lực dính c','phi_cu_effective':'Góc ma sát hữu hiệu φ′ CU','friction_phi':'Góc ma sát φ cắt','spt_n':'Chỉ số xuyên tiêu chuẩn Nₛₚₜ'}
        unit='10⁻³ cm²/s' if key=='cv' else PARAMS[key][1]
        title=names[key]+(' ('+unit+')' if unit!='—' else ' (không thứ nguyên)')
        c.create_text(w/2,8,anchor='n',text=PARAMS[key][0]+(' ('+unit+')' if unit!='—' else ' (không thứ nguyên)'),width=w-16,font=(UI_FONT,12,'bold'),fill='#0F172A')
        if self.plot_kind.get()!='Hằng số':
            c.create_text(w/2,40,text='y = az + b',font=(UI_FONT,10,'bold'),fill='#15803D')
        c.create_line(left,top,right,top,fill='#334155',width=1.5,arrow='last')
        if self.reference.get()==REFS[2]:
            c.create_line(left,bottom,left,top,fill='#334155',width=1.5,arrow='last')
        else:
            c.create_line(left,top,left,bottom,fill='#334155',width=1.5,arrow='last')
        vertical='H (m)' if self.reference.get()==REFS[2] else 'z (m)'
        c.create_text(12,(top+bottom)/2,text=vertical,angle=90,font=(UI_FONT,10,'bold'))
        ticks=min(5,max(3,int((right-left)/65)+1))
        settings=self._chart_axes.get((key,self.reference.get()),{})
        self._chart_limits=(xmin,xmax,zmin,zmax)
        def axis_ticks(lo,hi,step,count):
            if step is None:
                raw=(hi-lo)/max(1,count-1);base=10**math.floor(math.log10(raw))
                step=next(factor*base for factor in (1,2,5,10) if factor*base>=raw)
            first=math.ceil(lo/step-1e-10);last=math.floor(hi/step+1e-10)
            if last-first>200:return [lo+(hi-lo)*i/(count-1) for i in range(count)]
            return [i*step for i in range(first,last+1)]
        for x in axis_ticks(xmin,xmax,settings['xstep']/scale if settings.get('xstep') else None,ticks):
            if settings.get('vertical',True):c.create_line(px(x),top,px(x),bottom,fill='#E2E8F0')
            c.create_text(px(x),top-15,text=chart_number(x,key),font=(UI_FONT,10))
        for z in axis_ticks(zmin,zmax,settings.get('zstep'),min(12,max(3,int((bottom-top)/90)+1))):
            if settings.get('horizontal',True):c.create_line(left,py(z),right,py(z),fill='#E2E8F0')
            c.create_text(left-5,py(z),text=chart_number(z),anchor='e',font=(UI_FONT,10))
        selected=set(self.sample_tree.selection())
        for x,z,sample in points:
            if not (xmin<=x<=xmax and zmin<=z<=zmax):continue
            a,b=px(x),py(z);color='#0284C7' if sample.get('use',True) else '#CBD5E1';size=5 if sample['id'] in selected else 3
            c.create_oval(a-size,b-size,a+size,b+size,fill=color,outline='#0F172A' if size==5 else color);self.drawn.append((a,b,sample))
        used=[(z,value) for value,z,sample in points if sample.get('use',True)]
        mean=self.computed.get(key,{}).get('mean')
        detail=self.reference.get()+' (m).'
        equation_values=None
        def endpoint_label(value,x,y):
            c.create_oval(x-3,y-3,x+3,y+3,fill='#16A34A',outline='#16A34A')
            label=c.create_text(x+7,min(bottom-10,max(top+10,y)),anchor='w',
                                text=chart_number(value,key),font=(UI_FONT,10,'bold'),fill='#15803D')
            bounds=c.bbox(label)
            if bounds:
                if bounds[2]>right-3:c.move(label,right-3-bounds[2],0)
                bounds=c.bbox(label)
                if bounds[0]<left+3:c.move(label,left+3-bounds[0],0)
                bounds=c.bbox(label)
                background=c.create_rectangle(bounds[0]-2,bounds[1]-1,bounds[2]+2,bounds[3]+1,fill='white',outline='')
                c.tag_lower(background,label)
        if self.plot_kind.get()=='Hằng số':
            if mean is not None:
                if xmin<=mean<=xmax:
                    c.create_line(px(mean),top,px(mean),bottom,fill='#16A34A',width=2)
                    endpoint_label(mean,px(mean),top)
                    endpoint_label(mean,px(mean),bottom)
                detail+=' Trung bình của các mẫu được dùng.'
            else:
                detail+='\nChưa có mẫu hợp lệ được dùng để tính trung bình.'
        else:
            try:
                curve=fit_curve(used,KINDS[0])
                coords=[]
                for i in range(61):
                    z=curve['min']+(curve['max']-curve['min'])*i/60
                    value=evaluate(curve,z)
                    coords.extend((px(value),py(z)))
                visible_endpoints=[]
                # Clip each fitted segment to the selected plotting limits.
                for i in range(0,len(coords)-2,2):
                    x1,y1,x2,y2=coords[i:i+4];dx,dy=x2-x1,y2-y1;enter,leave=0.0,1.0;valid=True
                    for direction,distance in ((-dx,x1-left),(dx,right-x1),(-dy,y1-top),(dy,bottom-y1)):
                        if direction==0:
                            if distance<0:valid=False;break
                        else:
                            ratio=distance/direction
                            if direction<0:enter=max(enter,ratio)
                            else:leave=min(leave,ratio)
                            if enter>leave:valid=False;break
                    if valid:
                        first=(x1+enter*dx,y1+enter*dy);last=(x1+leave*dx,y1+leave*dy)
                        c.create_line(*first,*last,fill='#16A34A',width=2)
                        visible_endpoints.extend((first,last))
                if visible_endpoints:
                    for x,y in (visible_endpoints[0],visible_endpoints[-1]):
                        fraction=(bottom-y)/(bottom-top) if self.reference.get()==REFS[2] else (y-top)/(bottom-top)
                        z=zmin+fraction*(zmax-zmin)
                        endpoint_label(evaluate(curve,z),x,y)
                scale=1000 if key=='cv' else 1
                a=curve['coefficients'][1]/curve['scale']*scale
                b=(curve['coefficients'][0]-curve['coefficients'][1]*curve['center']/curve['scale'])*scale
                equation_values=(a,b,curve.get('r2'))
                detail+='\ny = az + b; '
                detail+=f'a = {a:.6g}; b = {b:.6g}'
                detail+='; R² = '+chart_number(curve.get('r2'))
            except ValueError as exc:
                detail+='\n'+str(exc)
        self.detail.set(detail)
        # Keep the two statistical values in their own band, also in copied images.
        table_left,table_right,table_top=8,w-8,h-(82 if fitted_display else 56)
        divider=table_left+(table_right-table_left)*.63
        c.create_rectangle(table_left,table_top,table_right,h-4,fill='white',outline='#94A3B8')
        c.create_line(divider,table_top,divider,table_top+52,fill='#94A3B8')
        c.create_line(table_left,table_top+26,table_right,table_top+26,fill='#94A3B8')
        coefficient=self.computed.get(key,{}).get('cv')
        c.create_text(table_left+6,table_top+13,anchor='w',text='Hệ số biến thiên',font=(UI_FONT,10))
        c.create_text((divider+table_right)/2,table_top+13,text=f'{coefficient:.3f}' if coefficient is not None else '—',font=(UI_FONT,10))
        c.create_text(table_left+6,table_top+39,anchor='w',text='Giá trị trung bình',font=(UI_FONT,10,'bold'))
        c.create_text((divider+table_right)/2,table_top+39,text=chart_number(mean,key) if mean is not None else '—',font=(UI_FONT,10,'bold'),fill='#15803D')
        if fitted_display:
            c.create_line(table_left,table_top+52,table_right,table_top+52,fill='#94A3B8')
            if equation_values is not None:
                a,b,r2=equation_values
                text=f'a = {a:.4g}; b = {b:.4g}; R² = {r2:.3f}' if r2 is not None else f'a = {a:.4g}; b = {b:.4g}'
            else:text='Chưa đủ dữ liệu lập hàm y = az + b.'
            c.create_text((table_left+table_right)/2,table_top+65,text=text,font=(UI_FONT,10),fill='#15803D')
    def draw_pressure_chart(self):
        """Display laboratory curves; averaging here never changes soil inputs."""
        if not hasattr(self,'pressure_canvas'):return
        c=self.pressure_canvas;c.delete('all');self._pressure_drawn=[]
        kind=self.pressure_kind.get();grid,parameter=('ep','e') if kind=='e-logP' else ('cvp','cv')
        log_y=parameter=='cv'
        curves=[];ignored=0
        for sample in self.visible:
            values=sample.get('original',{}).get('values',{})
            pressures=values.get(grid,[]);observations=values.get(parameter,[])
            if len(pressures)!=len(observations) or len(pressures)<2:continue
            pairs=[]
            for pressure,value in zip(pressures,observations):
                if not isinstance(pressure,(int,float)) or not isinstance(value,(int,float)):continue
                if not math.isfinite(pressure) or not math.isfinite(value):continue
                if pressure<=0 or (log_y and value<=0):ignored+=1;continue
                pairs.append((float(pressure),float(value)))
            if len(pairs)<2 or any(a[0]>=b[0] for a,b in zip(pairs,pairs[1:])):continue
            curves.append((sample,pairs))
        # The pressure table uses measured values at identical pressure stages.
        # It does not fill missing observations by interpolation.
        stages={}
        for sample in self.visible:
            if not sample.get('use',True):continue
            values=sample.get('original',{}).get('values',{})
            pressures=values.get(grid,[]);observations=values.get(parameter,[])
            if len(pressures)!=len(observations):continue
            own={}
            for pressure,value in zip(pressures,observations):
                if not isinstance(pressure,(int,float)) or not isinstance(value,(int,float)):continue
                if not math.isfinite(pressure) or not math.isfinite(value) or pressure<0:continue
                own.setdefault(pressure,[]).append(value)
            for pressure,values_at_stage in own.items():
                if len(set(values_at_stage))==1:stages.setdefault(pressure,[]).append(values_at_stage[0])
        stage_mean=[(pressure,statistics.mean(values)) for pressure,values in sorted(stages.items())]
        if not curves:
            c.create_text(12,20,anchor='nw',text='Chưa có đường cong '+kind+' với ít nhất hai cấp P > 0.',width=max(120,c.winfo_width()-24),fill='#64748B')
            self.pressure_detail.set('P (kgf/cm²), thang logarit. Không dùng độ sâu hoặc Cv trung bình để tạo đường cong.');return
        w,h=c.winfo_width(),c.winfo_height()
        if w<150 or h<140:return
        columns=max(1,int((w-64)/48))
        blocks=max(1,math.ceil(len(stage_mean)/columns))
        table_height=20+blocks*54
        h=max(h,64+110+52+table_height+8)
        c.configure(scrollregion=(0,0,w,h))
        left,right,top,bottom=58,w-16,64,h-52-table_height-8
        all_points=[pair for _,pairs in curves for pair in pairs]
        xmin,xmax=min(p for p,v in all_points),max(p for p,v in all_points)
        ymin,ymax=min(v for p,v in all_points),max(v for p,v in all_points)
        settings=self._chart_axes.get((kind,'logP'),{})
        xmin=settings['xmin'] if settings.get('xmin') is not None else 10.**math.floor(math.log10(xmin))
        xmax=settings['xmax'] if settings.get('xmax') is not None else 10.**math.ceil(math.log10(xmax))
        if log_y:
            ymin=settings['zmin'] if settings.get('zmin') is not None else 10.**math.floor(math.log10(ymin))
            ymax=settings['zmax'] if settings.get('zmax') is not None else 10.**math.ceil(math.log10(ymax))
            if ymin==ymax and settings.get('zmax') is None:ymax=ymin*10
            ystep=settings.get('zstep') or 1.
        else:
            pad=max((ymax-ymin)*.08,abs(ymax)*.03,.001);ymin-=pad;ymax+=pad
            target=max(2,min(6,int((bottom-top)/28)))
            raw=(ymax-ymin)/target;base=10.**math.floor(math.log10(raw))
            ystep=settings.get('zstep') or next(factor*base for factor in (1,2,5,10) if factor*base>=raw)
            ymin=settings['zmin'] if settings.get('zmin') is not None else math.floor(ymin/ystep)*ystep
            ymax=settings['zmax'] if settings.get('zmax') is not None else math.ceil(ymax/ystep)*ystep
        self._pressure_data_limits=(xmin,xmax,ymin,ymax)
        if xmin<=0 or xmin>=xmax or ymin>=ymax or (log_y and ymin<=0):
            c.create_text(12,20,anchor='nw',text='Giới hạn trục chưa hợp lệ; P và Cv trên trục logarit phải lớn hơn 0.',width=max(100,w-24));return
        logmin,logmax=math.log10(xmin),math.log10(xmax)
        ylo,yhi=(math.log10(ymin),math.log10(ymax)) if log_y else (ymin,ymax)
        start=math.ceil(ylo/ystep-1e-9);end=math.floor(yhi/ystep+1e-9)
        yticks=[10.**(i*ystep) if log_y else i*ystep for i in range(start,min(end,start+200)+1)]
        def tick_text(value):
            if abs(value)<ystep*1e-10:value=0.
            return format(value,'.8g')
        tick_font=tkfont.Font(family=UI_FONT,size=9)
        left=max(left,max((tick_font.measure(tick_text(v)) for v in yticks),default=0)+16)
        if right-left<50 or bottom-top<35:
            c.create_text(w/2,h/2,text='Kéo rộng khung để xem biểu đồ.',width=max(100,w-20),fill='#64748B');return
        def px(p):return left+(math.log10(p)-logmin)/(logmax-logmin)*(right-left)
        def py(v):
            coordinate=math.log10(v) if log_y else v
            return bottom-(coordinate-ylo)/(yhi-ylo)*(bottom-top)
        c.create_text(w/2,8,anchor='n',text=kind,font=(UI_FONT,12,'bold'),fill='#0F172A')
        c.create_line(left,33,left+22,33,fill='#DC2626',width=3)
        c.create_text(left+28,33,anchor='w',text='Trung bình',font=(UI_FONT,10))
        c.create_text(left,top-13,anchor='w',text='e (không thứ nguyên)' if parameter=='e' else 'Cv (10⁻³ cm²/s) · logarit',font=(UI_FONT,10))
        c.create_text((left+right)/2,bottom+38,text='P (kgf/cm²) · logarit',font=(UI_FONT,10))
        xstep=settings.get('xstep')
        if xstep:
            start=math.ceil(logmin/xstep);end=math.floor(logmax/xstep)
            ticks=[10**(i*xstep) for i in range(start,min(end,start+200)+1)]
        else:
            ticks=[factor*10.**power for power in range(math.floor(logmin)-1,math.ceil(logmax)+1) for factor in (1,) if xmin*(1-1e-10)<=factor*10.**power<=xmax*(1+1e-10)]
        if settings.get('vertical',True):
            for power in range(math.floor(logmin),math.ceil(logmax)):
                for factor in range(2,10):
                    pressure=factor*10.**power
                    if xmin<pressure<xmax:c.create_line(px(pressure),top,px(pressure),bottom,fill='#CBD5E1',width=.5)
        if settings.get('horizontal',True):
            if log_y:
                for power in range(math.floor(ylo),math.ceil(yhi)):
                    for factor in range(2,10):
                        value=factor*10.**power
                        if ymin<value<ymax:c.create_line(left,py(value),right,py(value),fill='#CBD5E1',width=.5)
            else:
                minor_step=ystep/2
                first=math.ceil(ymin/minor_step);last=math.floor(ymax/minor_step)
                for index in range(first,min(last,first+400)+1):
                    if index%2:c.create_line(left,py(index*minor_step),right,py(index*minor_step),fill='#E2E8F0',width=.5)
        last_label_right=-math.inf
        for pressure in sorted(ticks):
            x=px(pressure)
            if settings.get('vertical',True):c.create_line(x,top,x,bottom,fill='#64748B')
            label=format(pressure,'.8g');half=tick_font.measure(label)/2
            if x-half>=last_label_right+8:
                c.create_text(x,bottom+16,text=label,font=tick_font)
                last_label_right=x+half
        last_label_y=math.inf
        for value in yticks:
            y=py(value)
            if settings.get('horizontal',True):c.create_line(left,y,right,y,fill='#64748B')
            if last_label_y-y>=tick_font.metrics('linespace')+5:
                c.create_text(left-6,y,text=tick_text(value),anchor='e',font=tick_font)
                last_label_y=y
        c.create_rectangle(left,top,right,bottom,outline='#64748B')
        c.create_line(left,bottom,right,bottom,fill='#334155',width=1.5,arrow='last')
        c.create_line(left,bottom,left,top,fill='#334155',width=1.5,arrow='last')
        def line(pairs,color,width):
            # Smooth only the displayed shape; retain original points and tables.
            displayed=smooth_pressure_pairs(pairs,log_y=log_y)
            for first,second in zip(displayed,displayed[1:]):
                x1,y1=px(first[0]),py(first[1]);x2,y2=px(second[0]),py(second[1])
                dx,dy=x2-x1,y2-y1;enter,leave=0.,1.;valid=True
                for direction,distance in ((-dx,x1-left),(dx,right-x1),(-dy,y1-top),(dy,bottom-y1)):
                    if direction==0:
                        if distance<0:valid=False;break
                    else:
                        ratio=distance/direction
                        if direction<0:enter=max(enter,ratio)
                        else:leave=min(leave,ratio)
                        if enter>leave:valid=False;break
                if valid:c.create_line(x1+enter*dx,y1+enter*dy,x1+leave*dx,y1+leave*dy,fill=color,width=width)
        palette=('#0284C7','#7C3AED','#059669','#D97706','#DB2777','#475569')
        for index,(sample,pairs) in enumerate(curves):
            color='#1D4ED8' if sample.get('use',True) else '#CBD5E1'
            line(pairs,color,1)
            for p,v in pairs:
                if xmin<=p<=xmax and ymin<=v<=ymax:
                    x,y=px(p),py(v);c.create_oval(x-2,y-2,x+2,y+2,fill=color,outline=color)
                    self._pressure_drawn.append((x,y,f"{sample.get('hole','')} / {sample.get('sample','')}: P = {fmt(p)} kgf/cm²; {parameter} = {fmt(v)}"))
        used=[pairs for sample,pairs in curves if sample.get('use',True)]
        mean=[(p,v) for p,v in stage_mean if p>0 and (not log_y or v>0)]
        if mean:
            line(mean,'#DC2626',3)
            for p,v in mean:
                if xmin<=p<=xmax and ymin<=v<=ymax:
                    x,y=px(p),py(v);c.create_oval(x-3,y-3,x+3,y+3,fill='#DC2626',outline='#DC2626')
                    self._pressure_drawn.append((x,y,f'Trung bình {len(stages[p])} mẫu tại P = {fmt(p)} kgf/cm²; {parameter} = {fmt(v)}'))
        table_top=bottom+52
        c.create_text(w/2,table_top+9,text='P–'+parameter+' · Trung bình theo cấp',font=(UI_FONT,10,'bold'))
        table_top+=20
        for block in range(blocks):
            points_at_stage=stage_mean[block*columns:(block+1)*columns]
            row_top=table_top+block*54
            c.create_rectangle(8,row_top,w-8,row_top+52,fill='white',outline='#64748B')
            label_width=48;start_x=8+label_width
            c.create_line(8,row_top+26,w-8,row_top+26,fill='#64748B')
            c.create_line(start_x,row_top,start_x,row_top+52,fill='#64748B')
            c.create_text(8+label_width/2,row_top+13,text='P',font=(UI_FONT,10,'bold'))
            c.create_text(8+label_width/2,row_top+39,text=parameter+' TB',font=(UI_FONT,10,'bold'))
            if points_at_stage:
                cell_width=(w-8-start_x)/len(points_at_stage)
                for index,(pressure,value) in enumerate(points_at_stage):
                    x=start_x+index*cell_width
                    if index:c.create_line(x,row_top,x,row_top+52,fill='#64748B')
                    c.create_text(x+cell_width/2,row_top+13,text=f'{pressure:.5g}',font=(UI_FONT,10))
                    c.create_text(x+cell_width/2,row_top+39,text=f'{value:.3f}',font=(UI_FONT,10,'bold'))
            else:c.create_text((start_x+w-8)/2,row_top+39,text='—')
        detail=f'{len(curves)} đường cong mẫu; {len(used)} mẫu được dùng. Đường đỏ đậm: trung bình.'
        detail+='\nTrung bình theo cấp áp lực thực đo; không nội suy số liệu thiếu.'
        if used and not mean:detail+='\nChưa có cấp áp lực hợp lệ để vẽ đường trung bình.'
        if ignored:detail+=f'\n{ignored} điểm P/Cv không dương giữ trong bảng, không vẽ trên trục logarit.'
        self.pressure_detail.set(detail)

    def copy_pressure_chart(self):
        from lab_statistics import copy_canvas_image
        try:
            copy_canvas_image(self.pressure_canvas)
            self.status.set('Đã sao chép biểu đồ '+self.pressure_kind.get()+'. Dán bằng Ctrl+V.')
        except Exception as exc:messagebox.showerror('Sao chép biểu đồ',str(exc),parent=self)

    def hover_pressure_point(self,event):
        points=getattr(self,'_pressure_drawn',[])
        if points:
            mouse_x,mouse_y=self.pressure_canvas.canvasx(event.x),self.pressure_canvas.canvasy(event.y)
            x,y,text=min(points,key=lambda point:(point[0]-mouse_x)**2+(point[1]-mouse_y)**2)
            if (x-mouse_x)**2+(y-mouse_y)**2<100:self.status.set(text)

    def hover_point(self,event):
        if not self.drawn:return
        x,y,s=min(self.drawn,key=lambda p:(p[0]-event.x)**2+(p[1]-event.y)**2)
        if (x-event.x)**2+(y-event.y)**2<100:
            key=self.selected_param();unit='10⁻³ cm²/s' if key=='cv' else PARAMS[key][1]
            text=f'{s.get("hole","")} / {s.get("sample","")}: {chart_number(s.get("from"))}–{chart_number(s.get("to"))} m · {PARAMS[key][0]} = {chart_number(s.get(key),key)} {unit}'
            if self.status.get()!=text:self.status.set(text)

    def touch(self):self.store()['revision']+=1;self.update_ai_summary();self.refresh()
    def add_sample(self):self.sample_dialog(None)
    def edit_selected(self,event=None):
        selected=self.sample_tree.selection()
        if selected:self.sample_dialog(next(s for s in self.store()['samples'] if s['id']==selected[0]))
    def sample_dialog(self,sample):
        if not self.layer_id():messagebox.showinfo('Lớp đất','Thêm lớp đất trong Thông số tính toán trước.',parent=self);return
        popup=tk.Toplevel(self);popup.title('Mẫu địa chất');popup.geometry('660x620')
        canvas=tk.Canvas(popup,highlightthickness=0);bar=ttk.Scrollbar(popup,command=canvas.yview);canvas.configure(yscrollcommand=bar.set);bar.pack(side='right',fill='y');canvas.pack(fill='both',expand=True)
        body=ttk.Frame(canvas,padding=12);item=canvas.create_window(0,0,window=body,anchor='nw');body.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')));canvas.bind('<Configure>',lambda e:canvas.itemconfigure(item,width=e.width))
        fields=[('hole','Hố khoan'),('sample','Mã mẫu'),('from','Độ sâu từ (m)'),('to','Độ sâu đến (m)'),('ground','Cao độ mặt đất (m)'),('top','Độ sâu đỉnh lớp (m)'),*((k,v[0]+' ('+v[1]+')') for k,v in PARAMS.items() if k in soil_parameter_keys(self.data_project().method)),('source','Nguồn'),('reason','Lý do sửa / loại')];vars={}
        for i,(key,label) in enumerate(fields):
            ttk.Label(body,text=label).grid(row=i,column=0,sticky='w',pady=3);var=tk.StringVar(value='' if sample is None or sample.get(key) is None else parameter_text(sample[key],key));vars[key]=var;ttk.Entry(body,textvariable=var,width=42).grid(row=i,column=1,sticky='ew',padx=8)
        initial={key:var.get() for key,var in vars.items()}
        use=tk.BooleanVar(value=sample.get('use',True) if sample else True);ttk.Checkbutton(body,text='Dùng thống kê',variable=use).grid(row=len(fields),column=1,sticky='w')
        def save():
            try:
                data={key:((sample.get(key) if sample and var.get()==initial[key] else numeric(var.get())) if key in ('from','to','ground','top',*PARAMS) else var.get().strip()) for key,var in vars.items()}
                if not data['hole'] or data['from'] is None or data['to'] is None or data['from']<0 or data['to']<data['from']:raise ValueError('Cần hố khoan và khoảng độ sâu hợp lệ.')
                if sample and not data['reason']:raise ValueError('Ghi lý do khi sửa hoặc loại mẫu.')
                data.update(id=sample['id'] if sample else uuid.uuid4().hex,layer=sample['layer'] if sample else self.layer_id(),use=use.get())
                if sample:
                    data={**sample,**data}
                    from geotech_memory import GeotechMemory,project_scope
                    try:
                        scope=project_scope(self.data_project(),self.app.current_username or '')
                        GeotechMemory().record_correction(scope,'sample_edit',
                            {k:sample.get(k) for k in data if k not in ('history','original')},
                            {k:data.get(k) for k in data if k not in ('history','original')},
                            source=sample.get('source',''),
                            context={'layer_code':sample.get('layer_code'),'sample':sample.get('sample'),
                                     'hole':sample.get('hole'),'from':sample.get('from'),'to':sample.get('to'),
                                     'initial_detection':sample.get('original',{}),
                                     'cell_evidence':sample.get('source','')},
                            actor=self.app.current_username or 'Người dùng',confirmed=True)
                    except (OSError,ValueError) as exc:
                        raise ValueError('Chưa lưu được lịch sử sửa: '+str(exc)) from exc
                    except Exception as exc:
                        import sqlite3
                        if isinstance(exc,sqlite3.Error):raise ValueError('Lỗi bộ nhớ SQLite: '+str(exc)) from exc
                        raise
                    data['history']=sample.get('history',[])+[deepcopy({k:v for k,v in sample.items() if k!='history'})];sample.clear();sample.update(data)
                else:self.store()['samples'].append(data)
                self.touch();popup.destroy()
            except ValueError as exc:messagebox.showerror('Mẫu địa chất',str(exc),parent=popup)
        ttk.Button(body,text='Lưu mẫu',command=save).grid(row=len(fields)+1,column=1,pady=10)
    def toggle_sample(self):
        self.edit_selected()
    def apply_values(self):
        target_project=self.data_project()
        uid=self.layer_id();choices={k:v for k,v in self.store()['choices'].get(uid,{}).items() if k in PARAMS and k in soil_parameter_keys(self.data_project().method)}
        soil=next((s for s in self.data_project().soils if s.statistics_id==uid),None)
        code=next((entry['code'] for entry in self.store().get('layer_catalog',{}).values() if entry['id']==uid),soil.name if soil is not None else '')
        workspace=getattr(self.app,'_ai_analysis_workspace',None)
        materials=[m for m in workspace.state['materials'] if m['code'].casefold()==code.casefold()] if workspace is not None else []
        if len(materials)>1:
            messagebox.showwarning('Chọn đúng nguồn','Mã lớp trùng giữa các nguồn. Sửa tại Bảng tổng hợp chỉ tiêu trước khi áp dụng.',parent=self);return
        if not choices or (soil is None and not materials):messagebox.showinfo('Áp dụng','Lưu lựa chọn và khai báo lớp đất trong Bảng tổng hợp chỉ tiêu trước.',parent=self);return
        for key,config in choices.items():
            if config['mode']!=MODES[0] and config['fingerprint']!=self.fingerprint():messagebox.showwarning('Cần thống kê lại','Dữ liệu / bộ lọc khác lúc chọn '+PARAMS[key][0]+'. Lưu lựa chọn lại trước khi áp dụng.',parent=self);return
        popup=tk.Toplevel(self);popup.title('Chọn giá trị tính toán');popup.geometry('780x430')
        ttk.Label(popup,text='Nhấp đúp để chọn từng chỉ tiêu. Giá trị hàm tại x đại diện được áp dụng như một giá trị cố định cho lớp.',wraplength=740,padding=8).pack(fill='x')
        tree=self.make_tree(popup,('use','key','old','new','mode'),('Áp dụng','Chỉ tiêu','Đang dùng','Giá trị mới','Cách lấy'));selected=set()
        for key,config in choices.items():
            target='cv_constant' if key=='cv' else key
            old=getattr(soil,target,None) if soil is not None else materials[0]['values'].get(target)
            if key=='cv' and old is not None:old*=.001
            tree.insert('','end',iid=key,values=('☐',PARAMS[key][0],fmt(old),fmt(config['value']),config['mode']))
        def toggle(event):
            key=tree.identify_row(event.y)
            if not key:return
            if key in selected:selected.remove(key)
            else:selected.add(key)
            tree.set(key,'use','☑' if key in selected else '☐')
        tree.bind('<Double-1>',toggle)
        selection_bar=ttk.Frame(popup);selection_bar.pack(fill='x',padx=8,pady=4)
        def select_all(enabled):
            selected.clear()
            if enabled:selected.update(choices)
            for key in choices:tree.set(key,'use','☑' if key in selected else '☐')
        ttk.Button(selection_bar,text='Chọn tất cả',command=lambda:select_all(True)).pack(side='left',padx=4)
        ttk.Button(selection_bar,text='Bỏ chọn tất cả',command=lambda:select_all(False)).pack(side='left',padx=4)
        def commit():
            if not selected:
                messagebox.showinfo('Áp dụng','Chọn ít nhất một chỉ tiêu hoặc bấm Chọn tất cả.',parent=popup);return
            current=self.store()['choices'].get(uid,{})
            if self.data_project() is not target_project or self.layer_id()!=uid or any(current.get(key)!=choices[key] or (choices[key]['mode']!=MODES[0] and choices[key]['fingerprint']!=self.fingerprint()) for key in selected):
                messagebox.showwarning('Cần thống kê lại','Lớp, mẫu hoặc lựa chọn đã thay đổi. Đóng cửa sổ và chọn lại giá trị.',parent=popup);return
            workspace=getattr(self.app,'_ai_analysis_workspace',None)
            materials=[m for m in workspace.state['materials'] if m['code'].casefold()==code.casefold()] if workspace is not None else []
            if len(materials)>1:
                messagebox.showwarning('Chọn đúng nguồn','Mã lớp trùng giữa các nguồn. Sửa tại Bảng tổng hợp chỉ tiêu trước khi áp dụng.',parent=popup);return
            if 'cv' in selected and not messagebox.askyesno('Cv dùng tính toán','Lưu Cv trung bình? Đường cong Cv theo áp lực hợp lệ vẫn được ưu tiên khi tính.',parent=popup):return
            self.app._geology_statistics_backup=deepcopy(self.data_project())
            if workspace is not None:self.app._geology_statistics_materials_backup=deepcopy(workspace.state['materials'])
            for key in selected:
                config=choices[key];value=config['value']
                if soil is not None:apply_scalar(soil,key,value)
                if soil is not None:soil.parameter_sources[key]={'source':'Nhập trực tiếp' if config['mode']==MODES[0] else 'Thống kê','mode':config['mode'],'value':value,'reference':config.get('reference'),'at':config.get('at'),'revision':self.store()['revision'],'sample_ids':config.get('sample_ids',[])}
            if materials:
                material=materials[0]
                for key in selected:
                    target='cv_constant' if key=='cv' else key
                    material['values'][target]=choices[key]['value']*(1000 if key=='cv' else 1)
                    if key=='cv':material['cv_constant']=material['values'][target]
                    if key=='phi_cu_effective':
                        material['values']['strength_m']=math.tan(math.radians(choices[key]['value']))
                        material['phi_cu_effective']=choices[key]['value']
                material['average_edited']=True
                material['statistics_source']='Giá trị đã chọn trong Thống kê'
                workspace.invalidate('geology_approved')
            self.data_project().calculation_results={}
            for attr in ('_choice_group_results','_cdm_reports','_step_result_cache'):getattr(self.app,attr,{}).clear()
            self.app.chart_data={};self.app.populate()
            if hasattr(self.app,'render_chart'):self.app.render_chart()
            self.status.set('Đã áp dụng '+str(len(selected))+' chỉ tiêu vào '+self.layer.get()+'. Các chỉ tiêu khác giữ nguyên; cần tính lại kết quả.');popup.destroy()
        ttk.Button(popup,text='Áp dụng',command=commit).pack(pady=10)
    def import_file(self):
        if not self.layer_id():messagebox.showinfo('Import','Thêm lớp đất trước khi gán mẫu.',parent=self);return
        filenames=filedialog.askopenfilenames(parent=self,title='Chọn các file Excel/CSV để gộp mẫu',filetypes=[('Excel / CSV','*.xlsx *.xlsm *.csv')])
        if not filenames:return
        from ai_analysis_data import unique_data_sources
        filenames,notes=unique_data_sources(filenames,known_sources=self.store().get('imported_sources'),scope='manual')
        if notes:messagebox.showinfo('Bỏ qua file trùng dữ liệu','\n'.join(notes),parent=self)
        for filename in filenames:
            self.import_one_file(filename)

    def import_one_file(self,filename):
        try:
            path=Path(filename);tables={}
            if path.suffix.lower()=='.csv':
                with path.open(encoding='utf-8-sig',newline='') as f:
                    sample=f.read(4096);f.seek(0)
                    try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
                    except csv.Error:dialect=csv.excel
                    tables[path.stem]=list(csv.reader(f,dialect))
            else:
                from openpyxl import load_workbook
                book=load_workbook(path,read_only=True,data_only=True)
                try:
                    for sheet in book:
                        rows=[]
                        for row in sheet.iter_rows(values_only=True):
                            if len(rows)>=10000:raise ValueError('Mỗi sheet tối đa 10.000 hàng. Chưa import file này.')
                            rows.append(list(row))
                        tables[sheet.title]=rows
                finally:book.close()
            popup=self.map_import(path,tables)
            self.wait_window(popup)
        except Exception as exc:messagebox.showerror('Import mẫu',str(exc),parent=self)
    def map_import(self,path,tables):
        popup=tk.Toplevel(self);popup.title('Ánh xạ cột mẫu địa chất');popup.geometry('850x650')
        top=ttk.Frame(popup,padding=8);top.pack(fill='x');sheet=tk.StringVar(value=next(iter(tables)));header=tk.StringVar(value='1');forced=tk.BooleanVar(value=False)
        ttk.Label(top,text='Sheet').pack(side='left');ttk.Combobox(top,textvariable=sheet,values=tuple(tables),state='readonly',width=22).pack(side='left',padx=5)
        ttk.Label(top,text='Hàng tiêu đề').pack(side='left');ttk.Entry(top,textvariable=header,width=5).pack(side='left')
        ttk.Checkbutton(top,text='Gán mẫu vào lớp',variable=forced).pack(side='left',padx=10)
        ttk.Label(popup,text='Chỉ chọn cột có đơn vị đúng như nhãn. Ô thiếu để trống. Không tự suy đoán hay đổi đơn vị.',padding=8).pack(fill='x')
        body=ttk.Frame(popup,padding=8);body.pack(fill='both',expand=True);variables={};menus={};column_ids={}
        labels={'layer':'Lớp đất','hole':'Hố khoan','sample':'Mã mẫu','from':'Từ (m)','to':'Đến (m)','ground':'Cao độ mặt đất (m)','top':'Đỉnh lớp sâu (m)',**{k:v[0]+' ('+v[1]+')' for k,v in PARAMS.items()}}
        for i,(key,label) in enumerate(labels.items()):
            r=i%10;c=(i//10)*2;ttk.Label(body,text=label).grid(row=r,column=c,sticky='w',pady=4);var=tk.StringVar();variables[key]=var;menu=ttk.Combobox(body,textvariable=var,state='readonly',width=25);menu.grid(row=r,column=c+1,padx=(5,15));menus[key]=menu
        aliases={'layer':('layer','lop','lopdat','tenlop'),'hole':('hole','boring','hokhoan','lo khoan'),'sample':('sample','mau','mamau'),'from':('from','tu','zfrom'),'to':('to','den','zto'),'gamma':('gamma','dtrongtn','dungtrong'),'e0':('e0',),'cc':('cc',),'cs':('cs',),'pc':('pc',),'cv':('cv',),'co':('co',),'cohesion_c':('c',),'phi_cu_effective':('phicueffective','phicu','phiprimecu','phicuuhieu'),'spt_n':('sptn','spt')}
        def load_headers():
            try:
                row=tables[sheet.get()][int(header.get())-1];column_ids.clear()
                for i,value in enumerate(row):column_ids[f'{i+1}: {value or "(trống)"}']=i
                for key,menu in menus.items():
                    menu.configure(values=('',*column_ids));variables[key].set('')
                    for text,i in column_ids.items():
                        value=normal(row[i])
                        if any(value==normal(a) or (len(normal(a))>=2 and value.startswith(normal(a))) for a in aliases.get(key,(key,))):variables[key].set(text);break
            except (ValueError,IndexError,KeyError) as exc:messagebox.showerror('Tiêu đề',str(exc),parent=popup)
        ttk.Button(top,text='Đọc tiêu đề',command=load_headers).pack(side='right');load_headers()
        def commit():
            try:
                for key in ('hole','from','to'):
                    if not variables[key].get():raise ValueError('Cần chọn cột '+labels[key])
                if not forced.get() and not variables['layer'].get():raise ValueError('Chọn cột lớp đất hoặc Gán mọi mẫu vào lớp đang chọn.')
                layer_names={}
                for i,soil in enumerate(self.data_project().soils):
                    for name in (str(i+1),soil.name,str(i+1)+'. '+soil.name):
                        if normal(name) in layer_names and layer_names[normal(name)]!=soil.statistics_id:layer_names[normal(name)]=None
                        else:layer_names[normal(name)]=soil.statistics_id
                pending=[];failures=[];existing={(s['layer'],s.get('hole'),s.get('sample'),s['from'],s['to']) for s in self.store()['samples']};duplicates=0
                for n,row in enumerate(tables[sheet.get()][int(header.get()):],int(header.get())+1):
                    if not any(v is not None and str(v).strip() for v in row):continue
                    values={key:(row[column_ids[var.get()]] if var.get() and column_ids[var.get()]<len(row) else None) for key,var in variables.items()}
                    if not values['hole']:continue
                    try:
                        uid=self.layer_id() if forced.get() else layer_names.get(normal(values['layer']))
                        if not uid:raise ValueError('Chưa khớp lớp đất: '+str(values['layer']))
                        data={key:(numeric(value) if key in ('from','to','ground','top',*PARAMS) else str(value or '').strip()) for key,value in values.items()};data['layer']=uid
                        if data['from'] is None or data['to'] is None or data['from']<0 or data['to']<data['from']:raise ValueError('Khoảng độ sâu chưa hợp lệ.')
                        ident=(uid,data['hole'],data['sample'],data['from'],data['to'])
                        data.update(id=uuid.uuid4().hex,use=True,source=f'{path.name} / {sheet.get()} / hàng {n}',original={k:(v if isinstance(v,(str,int,float,bool,type(None))) else str(v)) for k,v in values.items()});pending.append(data)
                    except ValueError as exc:failures.append(f'Hàng {n}: {exc}')
                if failures:raise ValueError('Chưa import vì có hàng lỗi:\n'+'\n'.join(failures[:12]))
                if not pending:raise ValueError('Chưa có mẫu hợp lệ.')
                merged=deepcopy(self.store()['samples']);report=merge_sample_observations(merged,pending)
                if not report['added_values']:
                    messagebox.showinfo('Dữ liệu đã có',import_report_text(report)+'\n'+'\n'.join(report['conflicts'][:5]),parent=popup);return
                if not messagebox.askyesno('Xác nhận nhập dữ liệu',import_report_text(report)+'\nKiểm tra ánh xạ lớp và đơn vị trước khi nhập.\n'+'\n'.join(report['conflicts'][:5]),parent=popup):return
                self.store()['samples']=merged
                from ai_analysis_data import data_source_fingerprint
                if len(tables)==1:self.store().setdefault('imported_sources',{})['manual:'+data_source_fingerprint(path)]=path.name
                self.touch();self.status.set(import_report_text(report));popup.destroy()
            except (ValueError,KeyError,IndexError) as exc:messagebox.showerror('Import mẫu',str(exc),parent=popup)
        ttk.Button(popup,text='Xác nhận nhập',command=commit).pack(pady=10)
        return popup
    def ai_support(self,calculate):
        try:
            self.recalculate(quiet=True)
            key=self.selected_param()
            if calculate and self.mode.get()==MODES[2]:
                self.choose_value()
            values=[{k:sample.get(k) for k in ('hole','sample','from','to','ground','top','use',key,'source')} for sample in self.visible]
            context={'layer':self.layer.get(),'indicator':PARAMS[key],'reference':self.reference.get(),'samples':values,'computed':self.computed,'choices':{k:v for k,v in self.store()['choices'].get(self.layer_id(),{}).items() if k in PARAMS},
                     'rule':'Dữ liệu mẫu gốc chỉ đọc. Các kết quả thống kê do SOILFIRM PRO tính. Không tự loại mẫu hoặc sửa công thức. Không gán số liệu thiếu bằng 0. Hàm chỉ được lấy tại x đại diện để áp dụng giá trị cố định vào lớp.'}
            text=json.dumps(context,ensure_ascii=False)
            if len(text)>22000:raise ValueError('Phạm vi đang chọn quá lớn để gửi AI trong một lần. Lọc hố khoan hoặc lý trình; mẫu gốc vẫn giữ đầy đủ.')
            self.app._ai_statistics_context=text
            from chat_dialog import show_ai_dialog
            window=show_ai_dialog(self.app)
            question=('Giải thích kết quả thống kê đã tính cho lớp này, so sánh lấy trung bình và hàm tại độ sâu đại diện; nêu dữ liệu còn thiếu.' if calculate else 'Phân tích các mẫu địa chất này: chỉ ra giá trị thiếu, mẫu có khả năng trùng hoặc bất thường, xu hướng theo độ sâu và đề xuất cách xử lý. Không tự loại mẫu; nêu hố khoan, mã mẫu và nguồn của đề xuất.')
            window._soilfirm_ai_draft.set(question);window.lift();self.status.set('Đã chuyển dữ liệu bộ lọc sang AI đang chọn. Bấm Hỏi AI; dữ liệu gốc chưa thay đổi.')
        except Exception as exc:messagebox.showerror('AI thống kê',str(exc),parent=self)
    def export_calculation_values(self):
        """Export all sample observations by layer, borehole and depth."""
        self.export_all_statistics(initialfile='Tong_hop_chi_tieu_tung_mau.xlsx')

    def export_all_statistics(self, initialfile='Thong_ke_toan_bo_chi_tieu.xlsx'):
        """Export every layer and parameter; Excel owns the editable statistics."""
        data=self.store();samples=data.get('samples',[])
        if not samples:
            messagebox.showinfo('Xuất dữ liệu','Chưa có mẫu để xuất thống kê.',parent=self);return
        filename=filedialog.asksaveasfilename(parent=self,defaultextension='.xlsx',
            initialfile=initialfile,filetypes=[('Excel','*.xlsx')])
        if not filename:return
        try:
            import re
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill,Border,Side,Alignment
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table,TableStyleInfo
            from openpyxl.workbook.properties import CalcProperties
            book=Workbook();book.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
            overview=book.active;overview.title='Tong_hop'
            catalog=data.get('layer_catalog',{})
            names={entry['id']:entry['code'] for entry in catalog.values()}
            for soil in self.data_project().soils:
                if soil.statistics_id:names.setdefault(soil.statistics_id,soil.name)
            groups={}
            for sample in samples:groups.setdefault(sample.get('layer'),[]).append(sample)
            project=self.data_project()
            project_name=str(getattr(project,'project_name','') or getattr(project,'name','') or '')
            symbols={'gamma':'γ','e0':'e₀','cc':'Cc','cs':'Cs','pc':'Pc','cv':'Cv',
                     'co':'c₀ (Su)','cohesion_c':'c','phi_cu_effective':'φ′ CU',
                     'friction_phi':'φ cắt','spt_n':'Nₛₚₜ'}
            units={key:info[1] for key,info in PARAMS.items()}
            detail=book.create_sheet('Chi_tiet_mau')
            detail.append(['STT','Lớp đất','Lỗ khoan','Mã mẫu','Từ (m)','Đến (m)','Loại đất','Dùng',
                *(symbols[key]+' ('+units[key]+')' for key in PARAMS),'Lý do','Nguồn'])
            for index,sample in enumerate(samples,1):
                code=names.get(sample.get('layer'),sample.get('layer_code') or 'Lớp chưa xác định')
                category=sample.get('category') or data.get('layer_categories',{}).get(sample.get('layer'))
                detail.append([index,code,sample.get('hole',''),sample.get('sample',''),sample.get('from'),sample.get('to'),
                    {'Đất dính':'Sét / đất dính','Đất rời':'Cát / đất rời'}.get(category,'Chưa xác định'),
                    1 if sample.get('use',True) else 0,*(sample.get(key) for key in PARAMS),sample.get('reason',''),sample.get('source','')])
            detail.freeze_panes='E2';detail.auto_filter.ref=detail.dimensions
            for cell in detail[1]:
                cell.font=Font(name=UI_FONT,size=11,bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True)
            for row in detail.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,str) and cell.value.startswith(('=','+','@')):cell.data_type='s'
                    if isinstance(cell.value,(int,float)):
                        cell.number_format='0' if cell.column in (1,8,19) else '0.000E+00' if cell.column==14 else '0.000'
            for column in detail.columns:
                detail.column_dimensions[column[0].column_letter].width=min(65,max(14,max(len(str(cell.value or '')) for cell in column)+2))
            metrics=(('n','Số mẫu hợp lệ'),('max','Lớn nhất'),('min','Nhỏ nhất'),
                     ('mean','Trung bình'),('sd','Độ lệch chuẩn mẫu'),('cv','Hệ số biến thiên'))
            thin=Side(style='thin',color='CBD5E1');border=Border(left=thin,right=thin,top=thin,bottom=thin)
            header_fill=PatternFill('solid',fgColor='17365D')
            def title(sheet,text,last):
                sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last)
                sheet.cell(1,1,text);sheet.row_dimensions[1].height=30
                sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=last)
                sheet.cell(2,1,project_name or 'SOILFIRM PRO')
                sheet.merge_cells(start_row=3,start_column=1,end_row=3,end_column=last)
                sheet.cell(3,1,'Dùng = 1: tham gia thống kê; Dùng = 0: giữ số liệu, không tính. Ô thiếu để trống.')
                sheet.row_dimensions[3].height=32
            def section(sheet,row,text,last):
                sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last)
                cell=sheet.cell(row,1,text);cell.fill=PatternFill('solid',fgColor='DCE6F1')
                cell.font=Font(name=UI_FONT,size=12,bold=True,color='17365D')
                sheet.row_dimensions[row].height=26
            def formula_rows(sheet,first,last,columns,label_end):
                # Hidden per-observation helpers exclude omitted/blank cells without
                # interpreting missing data as zero; MIN/MAX therefore remain valid.
                bottom=last+2;result={}
                for offset,(metric,label) in enumerate(metrics):
                    row=bottom+offset
                    sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=label_end)
                    sheet.cell(row,1,label);result[metric]=row
                    for col,helper in columns:
                        rng=f'{get_column_letter(helper)}{first}:{get_column_letter(helper)}{last}'
                        expr={'n':f'COUNT({rng})','max':f'IF(COUNT({rng})=0,"",MAX({rng}))',
                              'min':f'IF(COUNT({rng})=0,"",MIN({rng}))',
                              'mean':f'IF(COUNT({rng})=0,"",AVERAGE({rng}))',
                              'sd':f'IF(COUNT({rng})<2,"",STDEV({rng}))',
                              'cv':f'IF(COUNT({rng})<2,"",IF(AVERAGE({rng})=0,"",STDEV({rng})/ABS(AVERAGE({rng}))))'}[metric]
                        cell=sheet.cell(row,col,'='+expr)
                        cell.number_format='0' if metric=='n' else '0.000'
                        cell.font=Font(name=UI_FONT,size=11,bold=metric=='mean',color='008000')
                        cell.fill=PatternFill('solid',fgColor='EAF2F8' if metric=='mean' else 'F4F7FA')
                return result,bottom+len(metrics)
            overview_headers=['Lớp đất','Chỉ tiêu','Đơn vị','Số mẫu','Lớn nhất','Nhỏ nhất','Trung bình','ĐLC mẫu','Hệ số biến thiên','Giá trị chọn']
            title(overview,'THỐNG KÊ TOÀN BỘ CHỈ TIÊU ĐẤT',len(overview_headers))
            overview.append([])
            for col,value in enumerate(overview_headers,1):overview.cell(5,col,value)
            overview_row=6
            for index,(uid,rows) in enumerate(groups.items(),1):
                code=str(names.get(uid) or next((r.get('layer_code') for r in rows if r.get('layer_code')),None) or 'Lớp chưa xác định')
                sheet_name=re.sub(r'[\\/*?:\[\]]','_',f'{index:02d}_{code}')[:31]
                sheet=book.create_sheet(sheet_name)
                widest=max((len({pressure for sample in rows for pressure in sample.get('original',{}).get('values',{}).get(grid,[]) if isinstance(pressure,(int,float))}) for grid in ('ep','cvp','mvp')),default=0)
                visible_last=max(19,6+widest)
                helper_base=visible_last+2
                title(sheet,'THỐNG KÊ CHỈ TIÊU · LỚP '+code,19)
                headers=['STT','Lỗ khoan','Mã mẫu','Từ (m)','Đến (m)','Dùng',
                         *(symbols[key]+(' ('+units[key]+')' if units[key]!='—' else ' (–)') for key in PARAMS),'Lý do','Nguồn']
                for col,value in enumerate(headers,1):sheet.cell(5,col,value)
                main_rows={}
                for offset,sample in enumerate(rows,6):
                    main_rows[id(sample)]=offset
                    values=[offset-5,sample.get('hole',''),sample.get('sample',''),sample.get('from'),sample.get('to'),
                            1 if sample.get('use',True) else 0,*(sample.get(key) for key in PARAMS),sample.get('reason',''),sample.get('source','')]
                    for col,value in enumerate(values,1):
                        cell=sheet.cell(offset,col,value)
                        if isinstance(value,(int,float)):cell.number_format='0' if col in (1,6,17) else '0.000' if col in (8,9,10) else '0.000E+00' if col==12 else '0.00'
                        if isinstance(value,str) and value.startswith(('=','+','@')):cell.data_type='s'
                        if not sample.get('use',True):cell.fill=PatternFill('solid',fgColor='F1F1F1')
                last=5+len(rows);helpers=[]
                for param_index,key in enumerate(PARAMS):
                    col=7+param_index;helper=helper_base+param_index;helpers.append((col,helper))
                    letter=get_column_letter(col);helper_letter=get_column_letter(helper)
                    sheet.column_dimensions[helper_letter].hidden=True
                    for row in range(6,last+1):sheet.cell(row,helper,f'=IF(AND($F{row}=1,ISNUMBER({letter}{row})),{letter}{row},"")')
                results,cursor=formula_rows(sheet,6,last,helpers,6)
                table=Table(displayName=f'MauLop{index}',ref=f'A5:S{last}')
                table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
                sheet.add_table(table)
                ref="'"+sheet.title.replace("'","''")+"'!"
                for param_index,key in enumerate(PARAMS):
                    col=get_column_letter(7+param_index)
                    values=[code,symbols[key],units[key],*[f'={ref}{col}{results[metric]}' for metric in ('n','max','min','mean','sd','cv')],
                            data.get('choices',{}).get(uid,{}).get(key,{}).get('value')]
                    for c,value in enumerate(values,1):overview.cell(overview_row,c,value)
                    overview_row+=1
                # Laboratory curves have independent pressure stages and Cv units.
                for grid,parameter,label,unit in (('ep','e','e–P','không thứ nguyên'),('cvp','cv','Cv–P','10⁻³ cm²/s'),('mvp','mv','mv–P','m²/T')):
                    observations=[];pressures=set()
                    for sample in rows:
                        original=sample.get('original',{}).get('values',{})
                        ps,vs=original.get(grid,[]),original.get(parameter,[])
                        if len(ps)!=len(vs) or not ps:continue
                        pairs={pressure:value for pressure,value in zip(ps,vs) if isinstance(pressure,(int,float)) and isinstance(value,(int,float))}
                        if not pairs:continue
                        pressures.update(pairs);observations.append((sample,pairs))
                    if not observations:continue
                    pressures=sorted(pressures);cursor+=2
                    section(sheet,cursor,label+' · P (kgf/cm²) · '+parameter+' ('+unit+')',max(19,6+len(pressures)))
                    cursor+=1;header_row=cursor
                    curve_headers=['STT','Lỗ khoan','Mã mẫu','Từ (m)','Đến (m)','Dùng',*[f'P = {pressure:g}' for pressure in pressures]]
                    for col,value in enumerate(curve_headers,1):sheet.cell(cursor,col,value)
                    first=cursor+1
                    for offset,(sample,pairs) in enumerate(observations,first):
                        vals=[offset-first+1,sample.get('hole',''),sample.get('sample',''),sample.get('from'),sample.get('to'),f'=$F{main_rows[id(sample)]}',*[pairs.get(pressure) for pressure in pressures]]
                        for col,value in enumerate(vals,1):sheet.cell(offset,col,value)
                    end=first+len(observations)-1
                    curve_helpers=[];helper_start=helper_base+len(PARAMS)+1
                    for k,pressure in enumerate(pressures):
                        col=7+k;helper=helper_start+k;curve_helpers.append((col,helper))
                        letter=get_column_letter(col);sheet.column_dimensions[get_column_letter(helper)].hidden=True
                        for row in range(first,end+1):sheet.cell(row,helper,f'=IF(AND($F{row}=1,ISNUMBER({letter}{row})),{letter}{row},"")')
                    curve_results,cursor=formula_rows(sheet,first,end,curve_helpers,6)
                    for k,pressure in enumerate(pressures):
                        col=get_column_letter(7+k)
                        values=[code,label+f' · P={pressure:g}',unit,*[f'={ref}{col}{curve_results[metric]}' for metric in ('n','max','min','mean','sd','cv')],None]
                        for c,value in enumerate(values,1):overview.cell(overview_row,c,value)
                        overview_row+=1
                    for cell in sheet[header_row]:
                        if cell.column<=len(curve_headers):cell.fill=header_fill;cell.font=Font(name=UI_FONT,size=11,bold=True,color='FFFFFF')
                sheet.freeze_panes='G6';sheet.print_title_rows='1:5';sheet.print_area=f'A1:{get_column_letter(visible_last)}{cursor}'
            for sheet in book:
                sheet.sheet_view.showGridLines=False
                from openpyxl.worksheet.properties import PageSetupProperties
                if sheet.sheet_properties.pageSetUpPr is None:sheet.sheet_properties.pageSetUpPr=PageSetupProperties()
                sheet.sheet_properties.pageSetUpPr.fitToPage=True
                sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
                sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
                if sheet is detail:
                    sheet.print_title_rows='1:1';sheet.print_area=sheet.dimensions;sheet.row_dimensions[1].height=36
                    for row in sheet.iter_rows(min_row=2):
                        for cell in row:
                            cell.font=Font(name=UI_FONT,size=11);cell.border=border
                    continue
                for row in sheet:
                    for cell in row:
                        if cell.value is None:continue
                        if cell.font.name!=UI_FONT:cell.font=Font(name=UI_FONT,size=11,color='000000')
                        cell.alignment=Alignment(vertical='center',horizontal='center' if cell.column>3 else 'left',wrap_text=True)
                        if cell.row>=5:cell.border=border
                        if cell.data_type=='f':
                            cell.font=Font(name=UI_FONT,size=11,bold='AVERAGE' in cell.value,color='008000')
                            if sheet is overview:cell.number_format='0' if cell.column==4 else '0.000'
                    sheet.row_dimensions[row[0].row].height=max(sheet.row_dimensions[row[0].row].height or 0,24)
                sheet.cell(1,1).font=Font(name=UI_FONT,size=16,bold=True,color='17365D')
                for cell in sheet[5]:
                    if cell.value is not None:cell.font=Font(name=UI_FONT,size=11,bold=True,color='FFFFFF');cell.fill=header_fill
                sheet.row_dimensions[5].height=42
                for col in range(1,sheet.max_column+1):
                    dimension=sheet.column_dimensions[get_column_letter(col)]
                    if not dimension.hidden:dimension.width=15 if col>6 else (7 if col in (1,6) else 18)
                sheet.column_dimensions['R'].width=25;sheet.column_dimensions['S'].width=55
                sheet.sheet_properties.outlinePr.summaryRight=False
            overview.freeze_panes='D6';overview.auto_filter.ref=f'A5:J{overview_row-1}'
            overview.column_dimensions['B'].width=28;overview.column_dimensions['C'].width=20
            overview.print_area=f'A1:J{overview_row-1}'
            overview.cell(3,1,'Thống kê toàn bộ lớp, không phụ thuộc bộ lọc màn hình. Hệ số biến thiên = STDEV(mẫu dùng) / |AVERAGE(mẫu dùng)|; cần ≥ 2 mẫu và TB ≠ 0.')
            book.save(filename)
            self.status.set('Đã xuất bảng tổng hợp từng mẫu, từng độ sâu: '+str(len(groups))+' lớp, '+str(len(samples))+' mẫu; kèm thống kê theo lớp và công thức Excel.')
        except Exception as exc:messagebox.showerror('Xuất dữ liệu',str(exc),parent=self)

    def export_excel(self):
        filename=filedialog.asksaveasfilename(parent=self,defaultextension='.xlsx',initialfile='Thong_ke_dia_chat.xlsx',filetypes=[('Excel','*.xlsx')])
        if not filename:return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill
            book=Workbook();ws=book.active;ws.title='Thong_ke';ws.append(['Lớp đất',self.layer.get(),'Hố khoan',self.hole.get()]);ws.append([self.stats_tree.heading(column,'text') for column in self.stats_tree['columns']])
            for item in self.stats_tree.get_children():
                values=self.stats_tree.item(item,'values');ws.append([values[0],*(numeric(v) for v in values[1:])])
            ws.append(['ĐLC mẫu: STDEV.S; HSBĐ = ĐLC / |TB|. Giá trị chọn chưa tự áp dụng vào bộ tính.'])
            samples=book.create_sheet('Mau_goc');keys=('layer','hole','sample','from','to','ground','top',*PARAMS,'use','reason','source');samples.append(list(keys))
            for sample in self.store()['samples']:samples.append([sample.get(k) for k in keys])
            cfg=book.create_sheet('Gia_tri_chon');cfg.append(['Lớp ID','Chỉ tiêu','Phương pháp','Giá trị','Trục x','x đại diện','Phương trình'])
            for uid,choices in self.store()['choices'].items():
                for key,value in choices.items():
                    if key not in PARAMS:continue
                    cfg.append([uid,PARAMS[key][0],value['mode'],value['value'],value.get('reference'),value.get('at'),(value.get('curve') or {}).get('formula')])
            for sheet in book:
                sheet.freeze_panes='B3' if sheet is ws else 'D2';sheet.auto_filter.ref=sheet.dimensions
                for cell in sheet[2 if sheet is ws else 1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163047')
                for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(cell.value or '')) for cell in col)+2))
            book.save(filename);self.status.set('Đã xuất: '+filename)
        except Exception as exc:messagebox.showerror('Xuất Excel',str(exc),parent=self)
    def export_pdf(self):
        filename=filedialog.asksaveasfilename(parent=self,defaultextension='.pdf',initialfile='Thong_ke_dia_chat.pdf',filetypes=[('PDF','*.pdf')])
        if not filename:return
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4,landscape
            from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
            from reportlab.lib.styles import ParagraphStyle
            from xml.sax.saxutils import escape
            fonts=[Path('C:/Windows/Fonts/arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
            font_path=next((p for p in fonts if p.is_file()),None)
            if font_path is None:raise ValueError('Chưa có font Unicode để xuất PDF tiếng Việt.')
            pdfmetrics.registerFont(TTFont('SoilFirmStats',str(font_path)));style=ParagraphStyle('stats',fontName='SoilFirmStats',fontSize=8,leading=11)
            story=[Paragraph('THỐNG KÊ ĐỊA CHẤT – '+escape(self.layer.get()),style),Spacer(1,10),Paragraph(escape(self.summary.get()),style),Spacer(1,10)]
            data=[['Đại lượng',*(v[0] for v in PARAMS.values())]]+[list(self.stats_tree.item(i,'values')) for i in self.stats_tree.get_children()]
            table=Table(data,repeatRows=1);table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'SoilFirmStats'),('FONTSIZE',(0,0),(-1,-1),7),('GRID',(0,0),(-1,-1),.4,colors.lightgrey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E2E8F0'))]));story.append(table)
            for key,value in self.store()['choices'].get(self.layer_id(),{}).items():
                if key not in PARAMS:continue
                story.extend([Spacer(1,6),Paragraph(escape(PARAMS[key][0]+': '+value['mode']+' = '+fmt(value['value'])+' '+PARAMS[key][1]+'; '+str((value.get('curve') or {}).get('formula',''))),style)])
            story.extend([Spacer(1,10),Paragraph('Độ lệch chuẩn mẫu; HSBĐ = ĐLC / |TB|. Hàm được lấy tại x đại diện để áp dụng giá trị cố định. Không tự suy diễn giá trị tiêu chuẩn hay giá trị thiết kế theo mức tin cậy.',style)])
            SimpleDocTemplate(filename,pagesize=landscape(A4),leftMargin=22,rightMargin=22).build(story);self.status.set('Đã xuất: '+filename)
        except Exception as exc:messagebox.showerror('Xuất PDF',str(exc),parent=self)
