"""Xuất bảng khối lượng ngang theo mẫu KL.xlsx; mọi khối lượng có công thức Excel."""
from __future__ import annotations

import math

from choice import parse_km_to_meters
from model import expansion_geometry, treatment_ranges


LABELS = {
    6: ('1', 'Cự ly', 'm'),
    7: ('2', 'Bề rộng đường B', 'm'),
    8: ('3', 'Bề rộng xử lý Bxl', 'm'),
    9: ('4', 'Chiều cao đắp Htk', 'm'),
    10: ('5', 'Chiều cao bù lún Hbl', 'm'),
    11: ('6', 'Cao độ đỉnh cọc CDM', 'm'),
    12: ('7', 'Cao độ đáy cọc CDM', 'm'),
    14: ('1', 'Chiều dài một cọc CDM', 'm'),
    15: ('2', 'Tổng số cọc CDM', 'cọc'),
    16: ('3', 'Tổng chiều dài cọc CDM', 'm'),
    17: ('4', 'Số cọc đại trà khoan lấy mẫu (2%)', 'cọc'),
    18: ('5', 'Tổng chiều dài lấy mẫu cọc đại trà', 'm'),
    19: ('6', 'Lấp hố khoan lõi bằng vữa M100 (Ø100)', 'm³'),
    20: ('7', 'Nén một trục cọc đại trà (2 m/mẫu)', 'mẫu'),
    21: ('8', 'Nén tĩnh cọc đại trà', 'test'),
    22: ('9', 'Số cọc thử CDM', 'cọc'),
    23: ('10', 'Chiều dài cọc thử CDM hàm lượng 230 kg/m³', 'm'),
    24: ('11', 'Chiều dài cọc thử CDM hàm lượng 260 kg/m³', 'm'),
    25: ('12', 'Chiều dài cọc thử CDM hàm lượng 290 kg/m³', 'm'),
    26: ('13', 'Số cọc thử khoan lấy mẫu', 'cọc'),
    27: ('14', 'Tổng chiều dài lấy mẫu cọc thử', 'm'),
    28: ('15', 'Nén một trục cọc thử', 'mẫu'),
    29: ('16', 'Nén tĩnh cọc thử', 'test'),
    30: ('17', 'Lấp hố khoan lõi cọc thử bằng M100', 'm³'),
    31: ('18', 'Số lớp vải gia cường CDM', 'lớp'),
    32: ('19', 'Vải địa kỹ thuật gia cường CDM', 'm²'),
    33: ('20', 'Cọc quan trắc chuyển vị ngang', 'cái'),
    34: ('21', 'Bàn đo lún mặt', 'bàn'),
    36: ('1', 'Chiều sâu đào thay đất', 'm'),
    37: ('2', 'Chiều sâu cọc tre', 'm'),
    38: ('3', 'Khối lượng đào lớp đất yếu', 'm³'),
    39: ('4', 'Khối lượng lấp lại bằng K95', 'm³'),
    40: ('5', 'Tổng chiều dài cọc tre', 'm'),
    41: ('6', 'Số lớp vải gia cường cơ học', 'lớp'),
    42: ('7', 'Vải địa kỹ thuật gia cường', 'm²'),
    43: ('8', 'Cọc quan trắc chuyển vị ngang', 'cái'),
    44: ('9', 'Bàn đo lún mặt', 'bàn'),
    46: ('1', 'Số lớp vải gia cường đắp trực tiếp', 'lớp'),
    47: ('2', 'Vải địa kỹ thuật gia cường', 'm²'),
    48: ('3', 'Số lớp vải ngăn cách', 'lớp'),
    49: ('4', 'Vải địa kỹ thuật ngăn cách', 'm²'),
    50: ('5', 'Cọc quan trắc chuyển vị ngang', 'cái'),
    51: ('6', 'Bàn đo lún mặt', 'bàn'),
    53: ('1', 'Chiều sâu cừ tràm', 'm'),
    54: ('2', 'Mật độ cừ tràm (nhập theo thiết kế)', 'cọc/m²'),
    55: ('3', 'Tổng chiều dài cừ tràm', 'm'),
    57: ('1', 'Thời gian chờ lún sau đắp', 'ngày'),
    58: ('2', 'Bàn đo lún theo chiều dài đoạn', 'bàn'),
    60: ('1', 'Chiều cao gia tải', 'm'),
    61: ('2', 'Bề rộng đỉnh gia tải', 'm'),
    62: ('3', 'Hệ số mái gia tải', '-'),
    63: ('4', 'Khối lượng đắp gia tải', 'm³'),
    64: ('5', 'Khối lượng dỡ gia tải', 'm³'),
    66: ('1', 'Khoảng cách PVD', 'm'),
    67: ('2', 'Chiều dài PVD', 'm'),
    68: ('3', 'Hệ số ô lưới PVD (tam giác √3/2; vuông 1)', '-'),
    69: ('4', 'Số bấc thấm PVD', 'bấc'),
    70: ('5', 'Tổng chiều dài PVD', 'm'),
    71: ('6', 'Khoảng cách SD', 'm'),
    72: ('7', 'Chiều dài cọc cát SD', 'm'),
    73: ('8', 'Hệ số ô lưới SD', '-'),
    74: ('9', 'Số cọc cát SD', 'cọc'),
    75: ('10', 'Tổng chiều dài cọc cát SD', 'm'),
    76: ('11', 'Chiều dày đệm cát thoát nước', 'm'),
    77: ('12', 'Thể tích đệm cát thoát nước', 'm³'),
    79: ('1', 'Khoảng cách cọc CDM', 'm'),
    80: ('2', 'Hệ số ô lưới CDM', '-'),
    81: ('3', 'Diện tích một ô cọc CDM', 'm²'),
    82: ('4', 'Đường kính cọc CDM', 'm'),
    83: ('5', 'Mật độ cọc tre', 'cọc/m²'),
    84: ('6', 'Hệ số mái đào', '-'),
    85: ('7', 'Chiều rộng vải neo mỗi bên', 'm'),
    86: ('8', 'Tỷ lệ lấy mẫu cọc đại trà', '-'),
    87: ('9', 'Đường kính lỗ khoan lõi', 'm'),
    88: ('10', 'Cao độ tự nhiên Ztn', 'm'),
}

GROUPS = {
    13: 'I. PHƯƠNG ÁN XỬ LÝ BẰNG CỌC CDM',
    35: 'II. PHƯƠNG ÁN ĐÀO THAY ĐẤT + CỌC TRE',
    45: 'III. PHƯƠNG ÁN ĐẮP TRỰC TIẾP',
    52: 'IV. CỌC CỪ TRÀM',
    56: 'V. PHƯƠNG ÁN CHỜ LÚN',
    59: 'VI. PHƯƠNG ÁN GIA TẢI',
    65: 'VII. PHƯƠNG ÁN PVD / SD',
    78: 'THÔNG SỐ ĐẦU VÀO ĐỂ KIỂM SOÁT CÔNG THỨC',
}

INPUT_ROWS = {4, 5, 7, 8, 9, 10, 14, 22, 31, 36, 37, 41, 46, 48, 53, 54,
              57, 60, 61, 62, 66, 67, 68, 71, 72, 73, 76, 79, 80, 82,
              83, 84, 85, 86, 87, 88}
SUM_ROWS = set(LABELS)-INPUT_ROWS-{11, 12, 81}


def _source(record):
    p = record['project_snapshot']
    payload = record.get('payload', {})
    group = payload.get('treatment_group', '')
    toe = p.crest_half_width+p.slope_width
    if p.expansion_width > 0:
        width = sum(right-left for left,right in treatment_ranges(p))
        top_width = p.expansion_width*(2 if p.expansion_side == 'Hai bên' else 1)
    else:
        width = 2*toe
        top_width = 2*p.crest_half_width
    mode = payload.get('treatment', '')
    is_pvd = group == 'drainage' and mode.startswith('PVD')
    is_sd = group == 'drainage' and mode.startswith('SD')
    factor = math.sqrt(3)/2 if payload.get('drain_pattern') == 'Tam giác' else 1.0
    cdm_factor = payload.get('cdm_cell_area', 0.0)/max(payload.get('cdm_s', 0.0)**2, 1e-8)
    return {
        7: 2*p.crest_half_width, 8: width, 9: p.h_design, 10: p.h_bl,
        14: payload.get('cdm_lc', 0.0) if group == 'cdm' else None,
        22: None, 31: payload.get('cdm_n_layer', 0) if group == 'cdm' else None,
        36: payload.get('replacement_depth', 0.0) if group == 'mechanical' else None,
        37: payload.get('bamboo_depth', 0.0) if group == 'mechanical' else None,
        41: None, 46: None, 48: None,
        53: payload.get('cajuput_depth', 0.0) if group == 'mechanical' else None,
        54: 25.0 if group == 'mechanical' else None,
        57: (payload.get('wait_days', 0.0) if group == 'drainage' else
             p.assessment_days if group == 'natural' else None),
        60: payload.get('surcharge_height', 0.0) if group in ('mechanical', 'drainage') else None,
        61: top_width, 62: 1.2,
        66: payload.get('drain_spacing', 0) if is_pvd else None,
        67: payload.get('drain_length', 0) if is_pvd else None,
        68: factor if is_pvd else None,
        71: payload.get('drain_spacing', 0) if is_sd else None,
        72: payload.get('drain_length', 0) if is_sd else None,
        73: factor if is_sd else None,
        76: p.h_sand_cushion if (is_pvd or is_sd) else None,
        79: payload.get('cdm_s', 0.0) if group == 'cdm' else None,
        80: cdm_factor if group == 'cdm' else None,
        82: payload.get('cdm_d', 0.0) if group == 'cdm' else None,
        83: 25.0 if group == 'mechanical' else None,
        84: 0.0 if p.expansion_width > 0 else 1.0, 85: 2.3, 86: 0.02, 87: 0.1,
        88: p.ground_elevation,
    }


def _formulas(col):
    c = col
    guard = lambda condition, expression: f'=IF({condition},{expression},"")'
    return {
        6: f'={c}5-{c}4',
        11: guard(f'{c}14>0', f'{c}88+{c}9+{c}10'),
        12: guard(f'{c}14>0', f'{c}11-{c}14'),
        15: guard(f'{c}14>0', f'ROUNDUP({c}8*{c}6/{c}81,0)'),
        16: guard(f'{c}15>0', f'{c}15*{c}14'),
        17: guard(f'{c}15>0', f'ROUNDUP({c}15*{c}86,0)'),
        18: guard(f'{c}17>0', f'{c}17*{c}14'),
        19: guard(f'{c}18>0', f'PI()*{c}87^2/4*{c}18'),
        20: guard(f'{c}18>0', f'ROUNDUP({c}18/2,0)'),
        21: guard(f'{c}17>0', f'{c}17'),
        23: guard(f'{c}22>0', f'{c}22*{c}14/3'),
        24: guard(f'{c}22>0', f'{c}23'),
        25: guard(f'{c}22>0', f'{c}23'),
        26: guard(f'{c}22>0', f'{c}22'),
        27: guard(f'{c}26>0', f'{c}26*{c}14'),
        28: guard(f'{c}27>0', f'ROUNDUP({c}27/2,0)'),
        29: guard(f'{c}26>0', f'{c}26'),
        30: guard(f'{c}27>0', f'PI()*{c}87^2/4*{c}27'),
        32: guard(f'{c}31>0', f'({c}8+2*{c}85)*{c}6*{c}31'),
        33: guard(f'{c}34>0', f'{c}34*2'),
        34: guard(f'{c}15>0', f'MAX(3,3*ROUND({c}6/100,0))'),
        38: guard(f'{c}36>0', f'({c}8+{c}84*{c}36)*{c}36*{c}6'),
        39: guard(f'{c}38>0', f'{c}38'),
        40: guard(f'{c}37>0', f'ROUNDUP(({c}8+2*{c}84*{c}36)*{c}6*{c}83,0)*{c}37'),
        42: guard(f'{c}41>0', f'({c}8+2*{c}85)*{c}6*{c}41'),
        43: guard(f'{c}44>0', f'{c}44*2'),
        44: guard(f'{c}36>0', f'MAX(3,3*ROUND({c}6/100,0))'),
        47: guard(f'{c}46>0', f'({c}8+2*{c}85)*{c}6*{c}46'),
        49: guard(f'{c}48>0', f'({c}8+2*{c}85)*{c}6*{c}48'),
        50: guard(f'{c}51>0', f'{c}51*2'),
        51: guard(f'OR({c}46>0,{c}48>0)', f'MAX(3,3*ROUND({c}6/100,0))'),
        55: guard(f'AND({c}53>0,{c}54>0)',
                  f'ROUNDUP(({c}8+2*{c}84*{c}36)*{c}6*{c}54,0)*{c}53'),
        58: guard(f'{c}57>0', f'MAX(3,3*ROUND({c}6/100,0))'),
        63: guard(f'{c}60>0', f'({c}61*{c}60+{c}62*{c}60^2)*{c}6'),
        64: guard(f'{c}63>0', f'{c}63'),
        69: guard(f'AND({c}66>0,{c}67>0)', f'ROUNDUP({c}8*{c}6/({c}68*{c}66^2),0)'),
        70: guard(f'{c}69>0', f'{c}69*{c}67'),
        74: guard(f'AND({c}71>0,{c}72>0)', f'ROUNDUP({c}8*{c}6/({c}73*{c}71^2),0)'),
        75: guard(f'{c}74>0', f'{c}74*{c}72'),
        77: guard(f'{c}76>0', f'{c}8*{c}6*{c}76'),
        81: guard(f'{c}79>0', f'{c}79^2*{c}80'),
    }


def export_kl(path, records):
    """Tạo KL cùng bố cục mẫu, một cột mỗi đoạn và tổng SUM theo từng hàng."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ValueError('Cần cài openpyxl để xuất bảng KL Excel.') from exc
    records = [row for row in records if row.get('project_snapshot') is not None]
    if not records:
        raise ValueError('Chưa có phân đoạn đã chọn để xuất khối lượng.')
    book = Workbook()
    sheet = book.active
    sheet.title = 'KL'
    last = get_column_letter(3+len(records))
    total = get_column_letter(4+len(records))
    note = get_column_letter(5+len(records))
    sheet.merge_cells(f'A1:{note}1')
    sheet['A1'] = 'BẢNG KHỐI LƯỢNG XỬ LÝ ĐẤT YẾU'
    for key, value in {'A2': 'TT', 'B2': 'Hạng mục', 'C2': 'Đơn vị',
                       f'{total}2': 'Tổng cộng', f'{note}2': 'Ghi chú'}.items():
        sheet[key] = value
    for row, (seq, label, unit) in LABELS.items():
        sheet.cell(row, 1, seq)
        sheet.cell(row, 2, label)
        sheet.cell(row, 3, unit)
    for row, label in GROUPS.items():
        sheet.merge_cells(start_row=row, start_column=1, end_row=row,
                          end_column=3)
        sheet.cell(row, 1, label)
    for idx, rec in enumerate(records, 4):
        col = get_column_letter(idx)
        p = rec['project_snapshot']
        start = parse_km_to_meters(p.station_from)
        end = parse_km_to_meters(p.station_to)
        if start is None and end is not None:
            start = end-float(rec['length'])
        if end is None and start is not None:
            end = start+float(rec['length'])
        if start is None:
            start = 0.0
            end = float(rec['length'])
        sheet[f'{col}2'] = rec.get('section_no') or idx-3
        sheet[f'{col}3'] = f'{p.station_from} – {p.station_to}  |  {rec.get("opt_name", "")}'
        sheet[f'{col}4'] = start
        sheet[f'{col}5'] = end
        for row, value in _source(rec).items():
            if value is not None and (value != 0 or row in (4, 5, 7, 8, 9, 10, 88)):
                sheet[f'{col}{row}'] = value
        for row, formula in _formulas(col).items():
            sheet[f'{col}{row}'] = formula
        if (rec.get('payload', {}).get('treatment_group') == 'cdm' and
                rec.get('cdm_count_mode', 'cad' if rec.get('cad_cdm_count') is not None else 'formula') == 'cad' and
                rec.get('cad_cdm_count') is not None):
            sheet[f'{col}15'] = int(rec['cad_cdm_count'])
            sheet[f'{note}15'] = 'Số CIRCLE layer CDM<STT> từ DXF; hàng 16 trở đi tính bằng công thức.'
        for row in LABELS:
            if row not in INPUT_ROWS and row not in (11, 12, 81):
                sheet[f'{total}{row}'] = f'=SUM(D{row}:{last}{row})'
    sheet[f'{note}4'] = 'Lý trình tính bằng mét; số 0 là mốc nội bộ nếu chưa có lý trình.'
    sheet[f'{note}38'] = 'Mái đào 1:1 mặc định; sửa hệ số tại hàng 84.'
    sheet[f'{note}54'] = 'Nhập mật độ cừ tràm theo hồ sơ thiết kế.'
    sheet[f'{note}22'] = 'Số cọc thử CDM do thiết kế thí nghiệm khai báo.'
    sheet[f'{note}63'] = 'Tính theo mặt cắt hình thang gia tải.'
    sheet[f'{note}69'] = 'Số bấc/cọc làm tròn lên cho toàn phân đoạn.'
    sheet[f'{note}86'] = 'Tỷ lệ 2% theo mẫu; có thể sửa.'
    sheet.column_dimensions['A'].width = 6
    sheet.column_dimensions['B'].width = 64
    sheet.column_dimensions['C'].width = 13
    for idx in range(4, 5+len(records)):
        sheet.column_dimensions[get_column_letter(idx)].width = 22
    sheet.column_dimensions[note].width = 70
    navy = PatternFill('solid', fgColor='1E3A5F')
    pale = PatternFill('solid', fgColor='EAF2FA')
    input_fill = PatternFill('solid', fgColor='FFF2CE')
    thin = Side(style='hair', color='C9D4E3')
    for row in sheet.iter_rows(min_row=1, max_row=88, max_col=5+len(records)):
        for cell in row:
            cell.font = Font(name='Times New Roman', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=cell.column in (2, 5+len(records)))
            if cell.row >= 2 and cell.value is not None:
                cell.border = Border(bottom=thin)
            if cell.column >= 4 and cell.column <= 3+len(records) and cell.row in INPUT_ROWS:
                cell.fill = input_fill
            if cell.data_type == 'f':
                cell.number_format = '#,##0.00;[Red](#,##0.00);–'
        if row[0].row in GROUPS:
            sheet.row_dimensions[row[0].row].height = 25
            for cell in row:
                cell.fill = pale
                cell.font = Font(name='Times New Roman', size=10, bold=True, color='1E3A5F')
    for cell in sheet[1]+sheet[2]:
        cell.fill = navy
        cell.font = Font(name='Times New Roman', size=11, bold=True, color='FFFFFF')
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[3].height = 38
    sheet.freeze_panes = 'D4'
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:3'
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.save(path)
