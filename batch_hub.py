"""Section summary, reviewed batch calculation and combined output; quantities only."""
from copy import deepcopy
from pathlib import Path
import json
import math
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from ui_theme import ProcessingBar


def fmt(value):
    return '—' if value is None else f'{value:.3f}' if isinstance(value,(float,int)) else str(value)


def consult_batch_ai(app,provider,label,context,cancel):
    """Ask the selected service; accept data-only JSON, never source or formulas."""
    import requests
    from ai_analysis_data import parse_ai_json
    if cancel.is_set():raise InterruptedError('Đã dừng tính AI.')
    providers={'Cloudflare AI':'cloudflare','Gemini':'gemini','DeepSeek':'deepseek','DeepSeek (g4f)':'deepseek_free','DeepSeek (miễn phí)':'ollama','Qwen (miễn phí)':'ollama_qwen','Groq':'groq','Grok (xAI)':'grok','ChatGPT':'openai','NVIDIA AI':'nvidia','Kimi AI':'kimi'}
    payload={'username':app.current_username,'key':app.current_login_key,'provider':providers[provider],
             'text':label,'context':json.dumps(context,ensure_ascii=False,default=str),
             'history':[],'tools':False}
    if len(payload['context'])>28000:raise ValueError('Ngữ cảnh đoạn quá lớn để gửi AI; giảm số lớp/phương án trong lần tính.')
    timeout=(5,135) if provider=='NVIDIA AI' else (5,105)
    try:
        if provider in ('DeepSeek (g4f)', 'DeepSeek (miễn phí)', 'Qwen (miễn phí)'):
            from local_ai_engine import make_local_post
            response = make_local_post(cancel, engine=providers[provider] if providers[provider] in ('ollama','ollama_qwen') else 'g4f')(app.API_BASE_URL+'/api/chat/ai',json=payload,timeout=timeout)
        else:
            response=requests.post(app.API_BASE_URL+'/api/chat/ai',json=payload,timeout=timeout)
        data=response.json()
    except requests.Timeout as exc:raise ValueError(provider+' quá thời gian chờ khi tính hàng loạt.') from exc
    except (requests.RequestException,ValueError) as exc:raise ValueError(provider+' không trả phản hồi hợp lệ: '+str(exc)) from exc
    if cancel.is_set():raise InterruptedError('Đã dừng tính AI.')
    if not response.ok or not data.get('success'):raise ValueError(provider+': '+str(data.get('message') or 'HTTP '+str(response.status_code)))
    actual_provider=str(data.get('source') or '').strip().casefold()
    if actual_provider != providers[provider]:
        raise ValueError('Phản hồi không xác nhận đúng AI đã chọn: '+provider+'. Đã dừng, không dùng kết quả từ AI khác.')
    answer=str(data.get('answer') or '').strip()
    if not answer or data.get('truncated'):raise ValueError(provider+' trả phản hồi rỗng hoặc bị cắt.')
    try:result=parse_ai_json(answer)
    except Exception as exc:raise ValueError(provider+' trả sai cấu trúc JSON: '+answer[:800]) from exc
    if not isinstance(result,dict):raise ValueError('AI phải trả một đối tượng JSON.')
    return result


def build_view(parent,app):
    toolbar=ttk.Frame(parent,padding=8);toolbar.pack(fill='x')
    selected=set(getattr(app,'_summary_selected_numbers',[]));filter_var=tk.StringVar(value='Tất cả')
    note=tk.StringVar(value='Chọn phân đoạn, tính toán và chọn phương án trong cùng bảng.');app.summary_progress_text=tk.StringVar()
    controls=ttk.Frame(parent,padding=(8,0));controls.pack(fill='x')
    ttk.Label(controls,text='Lọc trạng thái:').pack(side='left')
    ttk.Combobox(controls,textvariable=filter_var,state='readonly',width=18,values=('Tất cả','Chưa tính','Đạt','Chưa đạt','Chờ chọn phương án','Đã chọn phương án','Thiếu dữ liệu / Lỗi')).pack(side='left',padx=5)
    ttk.Label(controls,textvariable=note).pack(side='left',padx=8)
    ttk.Label(parent,textvariable=app.summary_progress_text,padding=8,foreground='#087780').pack(fill='x')
    app._summary_progressbar=ProcessingBar(parent)  # Compatibility; active progress is centered on the app.
    area=ttk.Frame(parent);area.pack(fill='both',expand=True,padx=8,pady=6);area.rowconfigure(0,weight=1);area.columnconfigure(1,weight=1)
    fixed=ttk.Treeview(area,columns=('check','number','station'),show='headings',selectmode='extended')
    moving=ttk.Treeview(area,columns=('hole','before','option','after','limit','params','check','review'),show='headings',selectmode='extended')
    for tree,columns,headers,widths in [(fixed,('check','number','station'),('Chọn','STT','Lý trình'),(45,65,210)),(moving,tuple(moving['columns']),('Lỗ khoan','Sc dư TXL (cm)','Phương án','Sc dư SXL (cm)','Cho phép (cm)','Thông số xử lý','Kiểm toán','Trạng thái'),(135,120,200,120,115,300,130,160))]:
        for key,label,width in zip(columns,headers,widths):tree.heading(key,text=label);tree.column(key,width=width,minwidth=width,stretch=False)
        tree.tag_configure('fail',foreground='#B91C1C');tree.tag_configure('review',foreground='#B45309');tree.tag_configure('pass',foreground='#166534')
    scroll_lock=[False]
    def moved(other,a,b):
        sy.set(a,b)
        if not scroll_lock[0]:
            scroll_lock[0]=True
            try:other.yview_moveto(a)
            finally:scroll_lock[0]=False
    sy=ttk.Scrollbar(area,command=lambda *a:(fixed.yview(*a),moving.yview(*a)))
    sx=ttk.Scrollbar(area,orient='horizontal',command=moving.xview)
    fixed.configure(yscrollcommand=lambda a,b:moved(moving,a,b));moving.configure(yscrollcommand=lambda a,b:moved(fixed,a,b),xscrollcommand=sx.set)
    fixed.grid(row=0,column=0,sticky='ns');moving.grid(row=0,column=1,sticky='nsew');sy.grid(row=0,column=2,sticky='ns');sx.grid(row=1,column=1,sticky='ew')
    def selection(origin,other):
        ids=origin.selection()
        if other.selection()!=ids:other.selection_set(ids)
    fixed.bind('<<TreeviewSelect>>',lambda e:selection(fixed,moving));moving.bind('<<TreeviewSelect>>',lambda e:selection(moving,fixed))
    rows={};records={}
    def refresh():
        rows.clear();records.clear()
        session=getattr(app,'_ai_analysis_state',{}) or {}
        for section in session.get('sections',[]):rows[int(section['section_no'])]=section
        for rec in getattr(app,'_saved_sections_data',[]):
            if rec.get('section_no') is not None:records[int(rec['section_no'])]=rec;rows.setdefault(int(rec['section_no']),{})
        for n,rec in getattr(app,'_batch_candidates',{}).items():
            records.setdefault(int(n),rec);rows.setdefault(int(n),{})
        for failure in getattr(app,'_batch_failures',[]):
            n=failure.get('section_no')
            if n is not None and int(n) not in records:records[int(n)]={'status':'LỖI','notes':failure.get('reason')};rows.setdefault(int(n),{})
        for tree in (fixed,moving):tree.delete(*tree.get_children())
        count={'Đã chọn phương án':0,'Chờ chọn phương án':0,'Chưa tính':0,'Thiếu dữ liệu / Lỗi':0,'Chưa đạt':0}
        for n in sorted(rows):
            section=rows[n];rec=records.get(n);passed=rec and rec.get('status')=='ĐẠT'
            state=('Chưa tính' if rec is None else 'Thiếu dữ liệu / Lỗi' if rec.get('status') in ('LỖI','THIẾU DỮ LIỆU') else 'Chờ chọn phương án' if rec.get('needs_review') and passed else 'Đã chọn phương án' if passed else 'Chưa đạt')
            count[state]+=1
            wanted=filter_var.get()
            if wanted!='Tất cả' and not (wanted==state or wanted=='Đạt' and passed or wanted=='Chưa đạt' and rec and not passed):continue
            station=(rec or {}).get('station') or (section.get('station_from','')+' – '+section.get('station_to','')).strip(' –')
            project=(rec or {}).get('project_snapshot');before=(rec or {}).get('before') or {}
            tag='pass' if state=='Đã chọn phương án' else 'review' if state=='Chờ chọn phương án' else 'fail' if rec and not passed else ''
            fixed.insert('','end',iid=str(n),values=('☑' if n in selected else '☐',n,station),tags=(tag,))
            moving.insert('','end',iid=str(n),values=(getattr(project,'borehole_name',None) or section.get('borehole_selected','—'),fmt(before.get('residual_cm')),(rec or {}).get('opt_name','—'),fmt((rec or {}).get('residual')),fmt((rec or {}).get('limit',section.get('limit_cm'))),(rec or {}).get('params','—'),(rec or {}).get('status','Chưa tính'),state),tags=(tag,))
        note.set(f'{len(rows)} đoạn · Đã chọn {count["Đã chọn phương án"]} · Chờ chọn phương án {count["Chờ chọn phương án"]} · Chưa tính {count["Chưa tính"]} · Chưa đạt/lỗi {count["Chưa đạt"]+count["Thiếu dữ liệu / Lỗi"]}')
        app._summary_selected_numbers=sorted(selected)
    def toggle(event):
        item=fixed.identify_row(event.y)
        if not item:return
        if fixed.identify_column(event.x)=='#1':
            n=int(item)
            if n in selected:selected.remove(n)
            else:selected.add(n)
            fixed.set(item,'check','☑' if n in selected else '☐');app._summary_selected_numbers=sorted(selected)
    fixed.bind('<ButtonRelease-1>',toggle)
    def select_all():selected.update(rows);refresh()
    def clear_selection():selected.clear();refresh()
    def run_batch(numbers=None, use_ai=True):
        if app.design_mode.get() != 'TÍNH TOÀN TUYẾN':
            messagebox.showinfo('Luồng tính', 'Chức năng này sử dụng dữ liệu TÍNH TOÀN TUYẾN.', parent=app);return
        if not app.require_full_license():return
        workspace=app._ai_analysis_workspace
        if workspace.busy:return
        if not app.calculation_settings()['after']:
            app.batch_before_view.calculate(use_ai, selected if numbers is None else numbers);return
        session=workspace.state
        if not all(session.get(k) for k in ('geology_approved','boreholes_approved','sections_approved')):
            messagebox.showinfo('Dữ liệu đầu vào','Xác nhận chỉ tiêu và lỗ khoan ở mục 2, dữ liệu phân đoạn ở mục 1 trước khi tính hàng loạt.',parent=app);return
        targets = selected if numbers is None else set(numbers)
        sections=deepcopy([s for s in session['sections'] if not targets or int(s['section_no']) in targets])
        if not sections:messagebox.showinfo('Tính hàng loạt','Chưa có phân đoạn được chọn.',parent=app);return
        if use_ai and (not app.current_username or not app.current_login_key):
            messagebox.showinfo('AI hàng loạt','Đăng nhập trước khi gọi AI.',parent=app);return
        provider=app.ai_provider_var.get()
        settings=deepcopy(session['settings']);options=deepcopy(session.get('options',[]));template=deepcopy(session['template']);captured=deepcopy(session)
        allowed_options=app.calculation_options()
        options=[name for name in options if name in allowed_options] or list(allowed_options)
        def action(progress):
            from ai_analysis_data import project_for_section,make_record,validate_analysis_project
            from soilfirm_ai_engine import before_treatment,execute_tool,optimization_instructions
            from batch_calculation import OPTIONS
            for current,section in enumerate(sections,1):
                n=int(section['section_no']);pending=None
                try:
                    progress(f'Đang tính đoạn {current}/{len(sections)} · STT {n}')
                    hole=next((h for h in captured['boreholes'] if h['name']==section['borehole_selected']),None)
                    if hole is None:raise ValueError('Chọn lỗ khoan cho phân đoạn trước khi tính.')
                    p=project_for_section(captured['template'],section,captured['materials'],hole);validate_analysis_project(p);before=before_treatment(p)
                    pending={'before':before,'natural':bool(before['pass_check']),'project':p,'section':section}
                    allowed=options or list(OPTIONS)
                    if use_ai:
                        progress(f'{provider} đang lập kế hoạch · STT {n}')
                        ai_plan=consult_batch_ai(app,provider,
                            'Lập kế hoạch kiểm toán đoạn SOILFIRM PRO. Chỉ trả JSON {"options":[tên phương án],"criterion":"priority hoặc wait","reason":"giải thích ngắn"}. Không tính nhẩm hoặc sửa công thức. Khi đã đạt trước xử lý giữ options theo danh sách. Nếu user_priority=true giữ đúng thứ tự tất cả phương án đã chọn; nếu false xếp thứ tự toàn bộ allowed_options dựa vào dữ liệu và quy tắc. Không bịa đầu vào.',
                            {'section':section,'before':before,'soils':[vars(x) for x in p.soils],
                             'allowed_options':allowed,'user_priority':bool(options),'settings':settings,
                             'rules':optimization_instructions()},workspace.cancel_event)
                    else:
                        ai_plan={'options':list(allowed),'criterion':'priority','reason':'Tính thủ công theo phạm vi và thứ tự người dùng cấu hình.'}
                    planned=ai_plan.get('options')
                    if not isinstance(planned,list) or len(planned)!=len(allowed) or len(set(planned))!=len(planned) or set(planned)!=set(allowed):
                        raise ValueError('AI chưa lập đủ danh sách phương án hợp lệ: '+str(planned))
                    if options and planned!=allowed:raise ValueError('AI thay đổi thứ tự ưu tiên người dùng đã chọn.')
                    criterion=ai_plan.get('criterion','priority')
                    if criterion not in ('priority','wait'):raise ValueError('Tiêu chí AI không hợp lệ.')
                    pending['ai_plan']=ai_plan
                    if pending['natural']:record=make_record(captured,section,p,before)
                    else:
                        configured=deepcopy(settings)
                        for key,value in (('pvd_spacing',template.drain_spacing),('pvd_diameter',template.drain_diameter),('surcharge_height',template.surcharge_height)):
                            if key not in configured and value>0:configured[key]=value
                        result=execute_tool({'tool':'optimize','params':{'options':planned,'criterion':criterion}},p,configured,length=section['length'],progress=progress)
                        pending['result']=result
                        if result.get('input_requirements'):raise ValueError('Thiếu đầu vào phương án: '+', '.join(r['label'] for r in result['input_requirements']))
                        choices=result.get('options',{})
                        choice=next((choices[a['option']] for a in result.get('summary',{}).get('attempts',[]) if a.get('pass_check') and a.get('option') in choices),None)
                        if choice is None:record={'section_no':n,'station':section['station_from']+' – '+section['station_to'],'status':'CHƯA ĐẠT','before':before,'limit':p.residual_limit_cm,'notes':'Chưa tìm được phương án đạt trong phạm vi đã chọn.'}
                        else:record=make_record(captured,section,p,before,choice)
                    if use_ai:
                        progress(f'{provider} đang đối chiếu kết quả · STT {n}')
                        result_summary=pending.get('result',{}).get('summary',{})
                        attempts=[{k:a.get(k) for k in ('option','pass_check','residual_cm','limit_cm','design')} for a in result_summary.get('attempts',[])]
                        ai_review=consult_batch_ai(app,provider,
                            'Đối chiếu kết quả kiểm toán SOILFIRM PRO. Chỉ trả JSON {"recommendation":"đề xuất ngắn","comparison":"so sánh ngắn các phương án theo số thực","missing":"dữ liệu còn thiếu hoặc chuỗi rỗng"}. Tối đa 600 ký tự mỗi trường. Không tự đổi kết quả ĐẠT/CHƯA ĐẠT, không bịa phương án chưa tính và không sửa công thức. Phương án phải được người dùng xác nhận chọn.',
                            {'section_no':n,'before':before,'plan':ai_plan,'attempts':attempts,
                             'selected_option':record.get('opt_name'),'status':record.get('status'),
                             'rules':optimization_instructions()},workspace.cancel_event)
                        for key in ('recommendation','comparison','missing'):
                            if not isinstance(ai_review.get(key,''),str):raise ValueError('AI trả diễn giải không hợp lệ.')
                        record['ai_provider']=provider
                        record['ai_plan']=ai_plan
                        record['ai_review']={key:ai_review.get(key,'')[:600] for key in ('recommendation','comparison','missing')}
                        record['notes']=' · '.join(filter(None,[record.get('notes',''),provider+': '+record['ai_review']['recommendation'],record['ai_review']['comparison'],record['ai_review']['missing']]))
                    else:
                        record['calculation_mode']='manual_route'
                        record['notes']=' · '.join(filter(None,[record.get('notes',''),'Tính thủ công từ dữ liệu toàn tuyến đã xác nhận; chờ xác nhận chọn phương án.']))
                    record['needs_review']=True
                    record['before_pass']=bool(before['pass_check'])
                except InterruptedError:raise
                except Exception as exc:record={'section_no':n,'station':section['station_from']+' – '+section['station_to'],'status':'LỖI','needs_review':True,'notes':str(exc)}
                workspace.events.put(('batch_record',(n,record,pending)))
            return len(sections)
        def done(count):
            app.summary_progress_text.set(f'Đã tính {count} đoạn. Xác nhận chọn phương án trước khi tổng hợp hồ sơ.')
            refresh()
            app.switch_step(16)
        workspace.run_job('AI lập kế hoạch → SOILFIRM PRO kiểm toán → AI đối chiếu…' if use_ai else 'Tính hàng loạt thủ công từ dữ liệu tuyến…',action,done)
    def run_ai(numbers=None):
        if not app.require_full_license():return
        targets = sorted(selected) if numbers is None else list(numbers)
        app._ai_analysis_workspace.edit_settings(on_save=lambda:run_batch(targets, use_ai=True))
    def run_manual(numbers=None):return run_batch(numbers, use_ai=False)
    def configure():app._ai_analysis_workspace.edit_settings()
    def show_detail(event=None):
        tree=moving if event and event.widget is moving else fixed
        item=tree.identify_row(event.y) if event else (fixed.selection()[0] if fixed.selection() else '')
        if not item:return
        n=int(item);pending=getattr(app,'_batch_runs',{}).get(n)
        if pending and (pending.get('natural') or pending.get('result')):
            workspace=app._ai_analysis_workspace
            section_index=next((i for i,s in enumerate(workspace.state['sections']) if int(s['section_no'])==n),None)
            if section_index is not None:workspace.state['index']=section_index
            workspace.state['pending']=deepcopy(pending);workspace.show_pending();workspace.tabs.select(workspace.pages[3])
        else:
            rec=records.get(n,{})
            messagebox.showinfo(f'Chi tiết STT {n}',rec.get('notes') or str(rec.get('params') or 'Đoạn chưa có kết quả. Bổ sung dữ liệu hoặc chạy tính.'),parent=app)
    fixed.bind('<Double-1>',show_detail);moving.bind('<Double-1>',show_detail)
    def approve():
        chosen=sorted(selected) or [int(i) for i in fixed.selection()]
        candidates=getattr(app,'_batch_candidates',{});eligible=[n for n in chosen if n in candidates and candidates[n].get('status')=='ĐẠT']
        if not eligible:messagebox.showinfo('Xác nhận phương án','Chọn các đoạn có phương án ĐẠT chờ chọn.',parent=app);return
        if not messagebox.askyesno('Xác nhận phương án',f'Xác nhận phương án đề xuất của {len(eligible)} đoạn đã chọn?',parent=app):return
        session=app._ai_analysis_workspace.state
        from design_workflow import section_signature
        for n in eligible:
            section=next((item for item in session['sections'] if int(item['section_no'])==n),None)
            pending=getattr(app,'_batch_runs',{}).get(n) or {}
            if section is None or not pending.get('project') or section_signature(app._ai_analysis_workspace.build_project(section)) != section_signature(pending['project']):
                messagebox.showinfo('Cần tính lại',f'STT {n}: đầu vào đã thay đổi; chưa xác nhận phương án.',parent=app);continue
            rec=deepcopy(candidates.pop(n));rec.pop('needs_review',None)
            app._saved_sections_data=[r for r in getattr(app,'_saved_sections_data',[]) if r.get('section_no')!=n]+[rec]
            session['records']=[r for r in session['records'] if r.get('section_no')!=n]+[deepcopy(rec)]
        app.refresh_treatment_boq();refresh()
    for text,command,style in [('Chọn tất cả đoạn',select_all,'TButton'),('Bỏ chọn tất cả',clear_selection,'TButton'),('Tính bằng AI',run_ai,'Accent.TButton'),('Xác nhận phương án',approve,'Accent.TButton'),('Xem kết quả đoạn',show_detail,'TButton')]:ttk.Button(toolbar,text=text,command=command,style=style).pack(side='left',padx=2)
    filter_var.trace_add('write',lambda *_:refresh())
    app._summary_run_manual=run_manual;app._summary_run_ai=run_ai;app.refresh_data_preview=lambda:None
    refresh();return refresh


def build_export_view(parent,app):
    body=ttk.Frame(parent,padding=12);body.pack(fill='both',expand=True)
    scope=tk.StringVar(value='Tất cả đoạn đã chọn');note=tk.StringVar(value='Xuất một Excel dữ liệu, một JSON và một PDF chung gồm trước và sau xử lý.')
    ttk.Label(body,text='Phạm vi xuất').pack(anchor='w')
    ttk.Combobox(body,textvariable=scope,state='readonly',values=('Tất cả đoạn đã chọn','Các đoạn được chọn ở bảng tổng hợp'),width=48).pack(anchor='w',pady=8)
    ttk.Label(body,textvariable=note,wraplength=950).pack(anchor='w',pady=8)
    def export():
        selected=set(getattr(app,'_summary_selected_numbers',[]))
        records=[deepcopy(r) for r in getattr(app,'_saved_sections_data',[]) if r.get('status')=='ĐẠT' and not r.get('needs_review') and (scope.get()=='Tất cả đoạn đã chọn' or r.get('section_no') in selected)]
        if not records:messagebox.showinfo('Xuất hồ sơ','Chưa có đoạn đã chọn trong phạm vi xuất.',parent=app);return
        source=getattr(app,'_section_excel_path',None) or getattr(app,'_ai_analysis_state',{}).get('source')
        if not source or not Path(source).is_file():
            source=filedialog.askopenfilename(parent=app,title='Chọn lại Data nguồn',filetypes=[('Excel','*.xlsx *.xlsm')])
        if not source:return
        target=filedialog.asksaveasfilename(parent=app,title='Tên bộ hồ sơ chung',initialfile='SoilFirm_2026.11',defaultextension='.json',filetypes=[('Bộ hồ sơ','*.json')])
        if not target:return
        try:
            from batch_bundle import export_bundle
            files=export_bundle(source,Path(target).with_suffix(''),records,individual=False)
            folders=getattr(app,'_results_folders',{})
            folders['TÍNH TOÀN TUYẾN']=str(Path(target).parent)
            app._results_folders=folders
            note.set('Đã xuất '+str(len(records))+' đoạn:\n'+'\n'.join(files))
        except Exception as exc:messagebox.showerror('Xuất hồ sơ',str(exc),parent=app)
    def preview_data():
        from tempfile import TemporaryDirectory
        from openpyxl import load_workbook
        from section_excel import export_data_results
        records=[r for r in getattr(app,'_saved_sections_data',[]) if r.get('status')=='ĐẠT' and not r.get('needs_review')]
        source=getattr(app,'_section_excel_path',None)
        if not records or not source or not Path(source).is_file():
            messagebox.showinfo('Xem Data xuất','Xác nhận phương án và chọn Data nguồn trước khi xem dữ liệu xuất.',parent=app);return
        try:
            with TemporaryDirectory(prefix='soilfirm_preview_') as folder:
                target=Path(folder)/'preview.xlsx';export_data_results(source,target,records,[r['section_no'] for r in records])
                book=load_workbook(target,data_only=True)
                popup=tk.Toplevel(app);popup.title('Data xuất · '+Path(source).name);popup.geometry('1100x650')
                tabs=ttk.Notebook(popup);tabs.pack(fill='both',expand=True)
                for sheet in book.worksheets:
                    if sheet.sheet_state!='visible':continue
                    from openpyxl.utils import get_column_letter
                    columns=[c for c in range(1,sheet.max_column+1) if not sheet.column_dimensions[get_column_letter(c)].hidden]
                    frame=ttk.Frame(tabs);tabs.add(frame,text=sheet.title)
                    tree=ttk.Treeview(frame,columns=[str(c) for c in columns],show='headings')
                    for c in columns:tree.heading(str(c),text=get_column_letter(c));tree.column(str(c),width=145,stretch=False)
                    for r in range(1,sheet.max_row+1):
                        if sheet.row_dimensions[r].hidden:continue
                        values=[sheet.cell(r,c).value for c in columns]
                        if any(v is not None for v in values):tree.insert('','end',values=['' if v is None else str(v) for v in values])
                    sy=ttk.Scrollbar(frame,command=tree.yview);sx=ttk.Scrollbar(frame,orient='horizontal',command=tree.xview)
                    tree.configure(yscrollcommand=sy.set,xscrollcommand=sx.set);sy.pack(side='right',fill='y');sx.pack(side='bottom',fill='x');tree.pack(fill='both',expand=True)
                book.close()
        except Exception as exc:messagebox.showerror('Xem Data xuất',str(exc),parent=app)
    ttk.Button(body,text='Xem trước dữ liệu xuất',command=preview_data).pack(anchor='w',pady=8)
    ttk.Button(body,text='Xuất hồ sơ' ,command=export).pack(anchor='w',pady=8)
    def open_results_folder():
        import os, subprocess
        folder=getattr(app,'_results_folders',{}).get('TÍNH TOÀN TUYẾN')
        if not folder or not Path(folder).is_dir():
            messagebox.showinfo('Thư mục kết quả','Xuất bộ hồ sơ trước khi mở thư mục kết quả.',parent=app);return
        try:
            if hasattr(os,'startfile'):os.startfile(folder)
            else:subprocess.Popen(['xdg-open',folder])
        except Exception as exc:messagebox.showerror('Mở thư mục kết quả',str(exc),parent=app)
    ttk.Button(body,text='Mở thư mục kết quả',command=open_results_folder).pack(anchor='w',pady=8)
    ttk.Button(body,text='Lưu dự án',command=app.save_file).pack(anchor='w',pady=8)
