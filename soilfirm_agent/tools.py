from __future__ import annotations
import copy,hashlib,json,zipfile,math
from pathlib import Path
from datetime import datetime,timezone
from collections.abc import Callable
import jsonschema
from .calculator import execute_numeric
from .documents import inspect_document,read_document,suggest_mapping,mapped_records
from .table_python import run_table_function
from .code_repair import CodeRepair
import uuid
from .web import search_web_duckduckgo,read_web_source,enrich_search_sources

def schema(properties,required):return {'type':'object','properties':properties,'required':required,'additionalProperties':False}
def field(kind='string',description=''):return {'type':kind,'description':description}
DEFINITIONS=[
 ('web_search','Search the web without API key; retrieve selected HTML/PDF passages if requested.',schema({'query':field(),'read_sources':field('boolean')},['query','read_sources'])),
 ('read_source','Read a public HTML page or text PDF and select passages related to the question.',schema({'url':field(),'question':field()},['url','question'])),
 ('python_calculate','Run restricted numerical Python assignments; assign result. Supports arithmetic and math functions, no imports, loops, filesystem, network or attributes. Include units and approved formula basis.',schema({'code':field(),'variables_json':field(),'basis':field(),'units_json':field()},['code','variables_json','basis','units_json'])),
 ('solve_equation','Solve a scalar nonlinear equation f(x)=0 in a bracket using bisection. Variables JSON and explicit formula basis are required.',schema({'expression':field(),'lower':field('number'),'upper':field('number'),'tolerance':field('number'),'variables_json':field(),'basis':field()},['expression','lower','upper','tolerance','variables_json','basis'])),
 ('list_files','List project files only.',schema({},[])),
 ('read_file','Read txt/json/csv or a bounded Excel range. Does not execute Excel macros/formulas. Data and formula strings are returned separately.',schema({'path':field(),'sheet':field(),'cell_range':field()},['path','sheet','cell_range'])),
 ('write_file','Create a new text/JSON/CSV output file inside project outputs; never overwrite an existing file.',schema({'path':field(),'content':field()},['path','content'])),
 ('write_excel','Create a new XLSX from JSON rows inside outputs. Formula-looking strings are stored as text.',schema({'path':field(),'sheet':field(),'rows_json':field()},['path','sheet','rows_json'])),
 ('memory_search','Search local soilfirm_memory.json and an explicitly configured trusted memory server.',schema({'query':field(),'limit':{'type':'integer','minimum':1,'maximum':20}},['query','limit'])),
 ('soilfirm_action','Invoke a registered internal calculation/export/drawing-preparation callback, never an arbitrary method or shell command.',schema({'action':field(),'params_json':field()},['action','params_json']))]

DOCUMENT_TOOLS=[
 ('inspect_document','Inspect all sheets/merged ranges/header previews or PDF page count.',schema({'path':field()},['path'])),
 ('read_document','Read paginated Excel rows (max200) or PDF pages (max5) with source cells, merged origin, cached formula values and table bounding boxes. OCR optional.',schema({'path':field(),'sheet':field(),'start':field('integer'),'end':field('integer'),'table_mode':{'enum':['lines','text']},'ocr':field('boolean')},['path','sheet','start','end','table_mode','ocr'])),
 ('suggest_mapping','Read header rows/pages and suggest exact column aliases; never treat suggestions as verified units/mapping.',schema({'path':field(),'sheet':field(),'start':field('integer'),'end':field('integer'),'kind':{'enum':['summary','borehole','vane','consolidation']}},['path','sheet','start','end','kind'])),
 ('extract_geotech','Extract canonical records with explicit mapping_json {field:{column,unit}}. Excel columns A/B; PDF columns 1/2. cv_main may be Cv1_2 only when selected by user/project. Saves dataset in outputs.',schema({'path':field(),'sheet':field(),'start':field('integer'),'end':field('integer'),'kind':{'enum':['summary','borehole','vane','consolidation']},'mapping_json':field(),'cv_main':field()},['path','sheet','start','end','kind','mapping_json','cv_main'])),
 ('run_table_python','Admin: generate def parse_rows(rows), using for/if, cell(row,"A"), number, text, lower, split, replace, findall, record(row,values), list.append. No imports/files/shell. Return records list; preserves source.',schema({'path':field(),'sheet':field(),'start':field('integer'),'end':field('integer'),'code':field()},['path','sheet','start','end','code'])),
 ('lookup_records','Read a saved extracted dataset by exact filters_json {borehole:"HK1",depth:2.5}; fields_json array or [] for all fields.',schema({'path':field(),'filters_json':field(),'fields_json':field()},['path','filters_json','fields_json'])),
 ('memory_update','Admin: automatically persist a sourced dictionary/mapping record locally and to configured server; report pending sync explicitly.',schema({'record_json':field()},['record_json'])),
 ('memory_sync','Admin: retry queued memory updates using configured authenticated server.',schema({},[])),
 ('code_list','Admin: list source files under configured code root.',schema({},[])),
 ('code_read','Admin: read source from configured code root before a repair.',schema({'path':field()},['path'])),
 ('code_patch','Admin: exact source replacement with expected SHA256, syntax validation and backup. Does not deploy or execute modified source.',schema({'path':field(),'expected_sha256':field(),'old':field(),'new':field(),'reason':field()},['path','expected_sha256','old','new','reason']))]
READ_TOOLS={'web_search','read_source','list_files','read_file','memory_search','inspect_document','read_document','suggest_mapping','extract_geotech','lookup_records'}

class ToolSet:
    def __init__(self,root,*,actions=None,memory_fetch=None,memory_write=None,role="reader",code_root=None,cancel=None,progress=None):
        self.root=Path(root).resolve();self.root.mkdir(parents=True,exist_ok=True)
        self.role=role;self.memory_write=memory_write;self.code_repair=CodeRepair(code_root) if code_root else None
        self.actions=dict(actions or {});self.memory_fetch=memory_fetch;self.cancel=cancel;self.progress=progress
        self.outputs=self.path('outputs');self.outputs.mkdir(exist_ok=True)
        self.definitions=[{'name':n,'description':d,'parameters':copy.deepcopy(s)} for n,d,s in DEFINITIONS+DOCUMENT_TOOLS if role=='admin' or n in READ_TOOLS]
        for d in self.definitions:
            if d['name']=='soilfirm_action':d['parameters']['properties']['action']['enum']=sorted(self.actions) or ['not_configured']
    def check_cancel(self):
        if self.cancel and self.cancel.is_set():raise InterruptedError('Agent cancelled')
    def path(self,name,write=False):
        p=(self.root/name).resolve()
        if not p.is_relative_to(self.root):raise ValueError('Path outside project')
        if write and not p.is_relative_to(self.outputs.resolve()):raise ValueError('Writes must be inside outputs/')
        return p
    def audit(self,name,args,result):
        record={'time':datetime.now(timezone.utc).isoformat(),'tool':name,'args_sha256':hashlib.sha256(json.dumps(args,sort_keys=True).encode()).hexdigest(),'ok':result.get('ok',False)}
        with self.path('outputs/agent_audit.jsonl',True).open('a',encoding='utf-8') as f:f.write(json.dumps(record)+'\n')
    def call(self,name,args):
        self.check_cancel()
        definition=next((x for x in self.definitions if x['name']==name),None)
        try:
            if self.role!='admin' and name not in READ_TOOLS:raise PermissionError('Admin required for this tool')
            if not definition:raise ValueError('Unknown tool')
            jsonschema.validate(args,definition['parameters'])
            value=getattr(self,'tool_'+name)(**args)
            result={'ok':True,'result':value}
        except InterruptedError:raise
        except Exception as exc:result={'ok':False,'error':type(exc).__name__,'message':str(exc)[:400]}
        self.audit(name,args,result);return result
    def tool_web_search(self,query,read_sources):
        r=search_web_duckduckgo(query,self.cancel)
        return enrich_search_sources(r,query,self.cancel,self.progress) if read_sources else r
    def tool_read_source(self,url,question):
        r=enrich_search_sources({'sources':[{'title':url,'url':url}],'answer':''},question,self.cancel,self.progress)
        if not any(s.get('read_status')=='read' for s in r['sources']):raise ValueError('No readable source: '+'; '.join(r.get('read_notes',[])))
        return r
    def tool_python_calculate(self,code,variables_json,basis,units_json):
        if not basis.strip():raise ValueError('State formula source/assumption')
        units=json.loads(units_json);variables=json.loads(variables_json)
        if not isinstance(units,dict) or not isinstance(variables,dict):raise ValueError('Units/variables must be objects')
        return {'value':execute_numeric(code=code,variables=variables),'basis':basis,'units':units,'verified':'Arithmetic only; does not validate engineering assumptions'}
    def tool_solve_equation(self,expression,lower,upper,tolerance,variables_json,basis):
        if not basis.strip():raise ValueError('State formula source/assumption')
        return {'solution':execute_numeric(expression=expression,variables=json.loads(variables_json),lower=lower,upper=upper,tolerance=tolerance),'basis':basis}
    def tool_list_files(self):
        return [str(p.relative_to(self.root)) for p in self.root.rglob('*') if p.is_file() and not p.is_symlink()][:300]
    def tool_read_file(self,path,sheet,cell_range):
        p=self.path(path)
        if p.stat().st_size>20*1024*1024:raise ValueError('File exceeds 20 MB')
        if p.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            from openpyxl.utils.cell import range_boundaries,get_column_letter
            bounds=range_boundaries(cell_range or 'A1:AZ100')
            if any(x is None for x in bounds) or (bounds[2]-bounds[0]+1)*(bounds[3]-bounds[1]+1)>5000:raise ValueError('Excel range exceeds 5000 cells')
            with zipfile.ZipFile(p) as z:
                if sum(x.file_size for x in z.infolist())>100*1024*1024:raise ValueError('Expanded Excel too large')
            book=load_workbook(p,data_only=False,read_only=True);cached=load_workbook(p,data_only=True,read_only=True)
            try:
                ws=book[sheet] if sheet else book.worksheets[0];cw=cached[ws.title];rows=[]
                for ri,row in enumerate(ws.iter_rows(min_col=bounds[0],min_row=bounds[1],max_col=bounds[2],max_row=bounds[3]),bounds[1]):
                    self.check_cancel();cells=[]
                    for ci,c in enumerate(row,bounds[0]):
                        address=f'{get_column_letter(ci)}{ri}'
                        cells.append({'cell':address,'value':str(c.value) if isinstance(c.value,datetime) else c.value,'formula':c.value if c.data_type=='f' else None,'cached_value':cw[address].value if c.data_type=='f' else None})
                    rows.append(cells)
                while len(rows)<bounds[3]-bounds[1]+1:
                    ri=bounds[1]+len(rows)
                    rows.append([{'cell':f'{get_column_letter(ci)}{ri}','value':None,'formula':None,'cached_value':None} for ci in range(bounds[0],bounds[2]+1)])
                return {'sheet':ws.title,'sheets':book.sheetnames,'range':cell_range,'rows':rows,'note':'No formula recalculation; missing cached values remain missing.'}
            finally:book.close();cached.close()
        if p.suffix.lower() not in ('.json','.txt','.csv','.md'):raise ValueError('Unsupported file type')
        return {'path':path,'text':p.read_text(encoding='utf-8-sig')[:60000]}
    def create(self,p,content):
        p.parent.mkdir(parents=True,exist_ok=True)
        with p.open('xb') as f:f.write(content)
        return {'path':str(p.relative_to(self.root)),'bytes':len(content),'sha256':hashlib.sha256(content).hexdigest()}
    def tool_write_file(self,path,content):
        p=self.path(path,True)
        if p.suffix.lower() not in ('.json','.txt','.csv','.md'):raise ValueError('Unsupported output extension')
        if len(content.encode())>2*1024*1024:raise ValueError('Output too large')
        if p.suffix.lower()=='.json':json.loads(content)
        return self.create(p,content.encode('utf-8'))
    def tool_write_excel(self,path,sheet,rows_json):
        from io import BytesIO
        from openpyxl import Workbook
        if len(rows_json.encode())>2*1024*1024:raise ValueError("Output too large")
        rows=json.loads(rows_json)
        if not isinstance(rows,list) or len(rows)>5000 or any(not isinstance(r,list) or len(r)>100 for r in rows):raise ValueError('Invalid rows')
        p=self.path(path,True)
        if p.suffix.lower()!='.xlsx':raise ValueError('Expected XLSX output')
        book=Workbook();ws=book.active;ws.title=sheet or 'Results'
        for row in rows:
            if any(not isinstance(v,(str,int,float,bool,type(None))) for v in row):raise ValueError('Cells must be scalar')
            if any(isinstance(v,str) and len(v)>32767 or isinstance(v,float) and not math.isfinite(v) for v in row):raise ValueError("Invalid cell value")
            ws.append(row)
            for c in ws[ws.max_row]:
                if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
        stream=BytesIO();book.save(stream);book.close();return self.create(p,stream.getvalue())
    def tool_memory_search(self,query,limit):
        path=self.path('soilfirm_memory.json')
        if path.exists() and path.stat().st_size>10*1024*1024:raise ValueError('Local memory too large')
        records=json.loads(path.read_text(encoding='utf-8')) if path.exists() else []
        if isinstance(records,dict):records=[{'key':k,'value':v} for k,v in records.items()]
        if not isinstance(records,list):raise ValueError('Memory must be list/object')
        hits=[{'source':'local','record':r} for r in records if any(w in json.dumps(r,ensure_ascii=False).lower() for w in query.lower().split())][:limit]
        if self.memory_fetch:
            remote=self.memory_fetch(query,limit)
            if not isinstance(remote,list):raise ValueError('Memory server must return list')
            hits.extend({'source':'configured_server','record':r} for r in remote[:limit])
        return hits[:limit]
    def tool_soilfirm_action(self,action,params_json):
        if action not in self.actions:raise ValueError('Action not wired to SoilFirm: '+action)
        params=json.loads(params_json)
        if not isinstance(params,dict):raise ValueError('Action params must be object')
        value=self.actions[action](params)
        if not isinstance(value,dict):raise ValueError('Callback must return a structured dict')
        if value.get('success') is False:raise ValueError(value.get('message','Action failed'))
        return value

    def tool_inspect_document(self,path):return inspect_document(self.path(path))
    def tool_read_document(self,path,sheet,start,end,table_mode,ocr):
        return read_document(self.path(path),sheet,start,end,table_mode,ocr,self.cancel)
    def tool_suggest_mapping(self,path,sheet,start,end,kind):
        data=read_document(self.path(path),sheet,start,end,cancel=self.cancel)
        return suggest_mapping(data['rows'],kind)
    def save_dataset(self,records,kind,extra=None):
        name='outputs/'+kind+'_'+uuid.uuid4().hex+'.json'
        result=self.create(self.path(name,True),json.dumps({'kind':kind,'records':records},ensure_ascii=False,default=str,allow_nan=False).encode())
        return {**result,'count':len(records),'preview':records[:10],'status':'found' if records else 'not_found','message':'' if records else 'Chưa tìm được bản ghi theo vùng/ánh xạ hiện tại; kiểm tra tiêu đề, cột, trang/sheet và thử cách đọc khác. Không tự điền số liệu thiếu.',**(extra or {})}
    def tool_extract_geotech(self,path,sheet,start,end,kind,mapping_json,cv_main):
        data=read_document(self.path(path),sheet,start,end,cancel=self.cancel)
        records=mapped_records(data['rows'],json.loads(mapping_json),kind,cv_main)
        result=self.save_dataset(records,kind,{'next_start':data['next_start'],'warnings':data['warnings']})
        if self.role=='admin':
            record={'key':'mapping:'+path+':'+sheet+':'+kind,'value':{'kind':kind,'mapping':json.loads(mapping_json),'Cv_main':cv_main,'status':'observed_mapping'},'source':path+' / '+sheet}
            result['memory_update']=self.tool_memory_update(json.dumps(record,ensure_ascii=False))
        return result
    def tool_run_table_python(self,path,sheet,start,end,code):
        data=read_document(self.path(path),sheet,start,end,cancel=self.cancel)
        records=run_table_function(code,data['rows'])
        result=self.save_dataset(records,'generated',{'next_start':data['next_start']})
        codepath=result['path'].removesuffix('.json')+'.py'
        self.create(self.path(codepath,True),code.encode());result['function_path']=codepath
        return result
    def tool_lookup_records(self,path,filters_json,fields_json):
        p=self.path(path)
        if p.stat().st_size>20*1024*1024:raise ValueError('Dataset exceeds 20 MB')
        data=json.loads(p.read_text(encoding='utf-8'));records=data.get('records')
        filters=json.loads(filters_json);fields=json.loads(fields_json)
        if not isinstance(records,list) or not isinstance(filters,dict) or not isinstance(fields,list) or any(not isinstance(x,str) for x in fields):raise ValueError('Invalid dataset lookup')
        hits=[r for r in records if all(r['values'].get(k)==v for k,v in filters.items())]
        return {'status':'found' if hits else 'not_found','message':'' if hits else 'Không có bản ghi khớp bộ lọc; kiểm tra khóa, mã mẫu, hố khoan và độ sâu trước khi đổi bộ lọc.','available_fields':sorted({k for r in records[:100] for k in r.get('values',{})}),'count':len(hits),'records':[{**r,'values':{k:v for k,v in r['values'].items() if not fields or k in fields}} for r in hits[:100]],'truncated':len(hits)>100}
    def tool_memory_update(self,record_json):
        record=json.loads(record_json)
        if not isinstance(record,dict) or set(record)!={'key','value','source'} or not isinstance(record['key'],str) or not record['key'].strip() or not isinstance(record['source'],str) or not record['source'].strip() or len(record_json)>50000:raise ValueError('Memory requires key,value,source; max50k characters')
        record['key']=record['key'][:200]
        path=self.path('soilfirm_memory.json')
        records=json.loads(path.read_text(encoding='utf-8')) if path.exists() else []
        if isinstance(records,dict):records=[{'key':k,'value':v,'source':'legacy'} for k,v in records.items()]
        if not isinstance(records,list):raise ValueError('Invalid local memory')
        records=[r for r in records if r.get('key')!=record['key']]+[record]
        if len(json.dumps(records))>10*1024*1024:raise ValueError('Memory too large')
        tmp=self.path('outputs/memory_'+uuid.uuid4().hex+'.tmp',True);tmp.write_text(json.dumps(records,ensure_ascii=False),encoding='utf-8');tmp.replace(path)
        pending=self.path('outputs/memory_pending.json',True)
        queue=json.loads(pending.read_text()) if pending.exists() else []
        queue=[r for r in queue if r.get('key')!=record['key']]+[record]
        pending.write_text(json.dumps(queue,ensure_ascii=False),encoding='utf-8')
        return {'local_saved':True,**self.tool_memory_sync()}
    def tool_memory_sync(self):
        p=self.path('outputs/memory_pending.json',True);queue=json.loads(p.read_text()) if p.exists() else []
        if not self.memory_write:return {'server_synced':False,'pending':len(queue),'message':'Server writer not configured; updates remain queued locally'}
        remaining=[];sent=0;errors=[]
        for record in queue:
            self.check_cancel()
            try:
                response=self.memory_write(record)
                if not isinstance(response,dict) or response.get('stored') is not True:raise ValueError('Server did not confirm storage')
                sent+=1
            except Exception as exc:remaining.append(record);errors.append(type(exc).__name__)
        p.write_text(json.dumps(remaining,ensure_ascii=False),encoding='utf-8')
        return {'server_synced':not remaining,'submitted':sent,'pending':len(remaining),'errors':errors}
    def tool_code_read(self,path):
        if not self.code_repair:raise ValueError('Code root not configured')
        return self.code_repair.read(path)
    def tool_code_patch(self,path,expected_sha256,old,new,reason):
        if not self.code_repair:raise ValueError('Code root not configured')
        return self.code_repair.patch(path,expected_sha256,old,new,reason)

    def tool_code_list(self):
        if not self.code_repair:raise ValueError('Code root not configured')
        return [str(p.relative_to(self.code_repair.root)) for p in self.code_repair.root.rglob('*') if p.is_file() and not p.is_symlink() and p.suffix in ('.py','.js','.mjs','.json') and not any(part.startswith('.') or part=='__pycache__' for part in p.relative_to(self.code_repair.root).parts)][:500]
